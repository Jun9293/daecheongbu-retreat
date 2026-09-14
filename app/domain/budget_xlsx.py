"""재정 엑셀 — **화면과 같은 함수의 결과를 놓기만 한다** (7-3 · 5-8과 같은 원칙).

세 시트: 지출 상세내역 · 예산 대비 집행 · 환급 대상자.
숫자는 전부 domain.budget 이 만든 것을 받는다 — 여기서 다시 세면 화면과
파일이 갈리고, 갈린 쪽이 파일이면 그 사실을 아무도 모른 채 카카오톡으로
돌아다닌다.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.budget import BudgetSummary, refund_sheet
from app.models import ExpenseEntry

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
MONEY = "#,##0"


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


계좌머리 = ["은행", "계좌번호", "예금주"]


def _계좌칸(e: ExpenseEntry) -> list:
    """계좌 셋 (7-4) — 부르는 쪽이 보일지를 이미 정했다."""
    return [e.payer_bank, e.payer_account_number, e.payer_account_holder]


def _영수증칸(e: ExpenseEntry) -> tuple[str | None, str | None]:
    """(자동 번호들, 원본 번호들). 원본 번호가 하나도 없으면 칸을 비운다."""
    numbers = ", ".join(str(r.number) for r in e.receipts) or None
    originals = ", ".join(r.original_no for r in e.receipts if r.original_no) or None
    return numbers, originals


def _widths(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def write(
    summary: BudgetSummary,
    entries: list[ExpenseEntry],
    *,
    계좌를_보인다: bool = False,
) -> io.BytesIO:
    """summary(화면과 같은 것)와 지출 목록을 받아 통째로 놓는다.

    **계좌 칸은 빈 칸이 아니라 없는 칸입니다.** 자리만 남기면 「무엇이
    가려져 있다」 가 보이고, 그것도 알 필요 없는 사실입니다 — 화면에서
    같은 판단을 한 그 자리입니다.

    판정은 여기서 하지 않습니다. `permissions.can_see_account` 하나가
    정하고 라우터가 그 답을 넘깁니다 — 화면과 파일이 저마다 판정하면
    같은 표인데 보는 사람이 갈립니다 (5-8).

    **기본이 「안 보임」 입니다.** 부르는 쪽이 빠뜨렸을 때 계좌가 새는
    쪽으로 기울면, 빠뜨린 것을 아무도 눈치채지 못합니다.
    """
    wb = Workbook()

    # 1) 지출 상세내역 — 시트 「Belong 예산」 의 컬럼 순서 (7-1)
    ws = wb.active
    ws.title = "지출 상세내역"
    headers = [
        "구분", "항목", "세부항목-1", "세부항목-2", "세부항목-3", "부서",
        "영수증번호", "원본 영수증번호", "지출일자", "금액", "지원금액", "개인부담액",
        "식사인원", "참석자 명단", "비고", "지급여부", "지급일", "지출자",
    ]
    if 계좌를_보인다:
        headers += 계좌머리
    ws.append(headers)
    for e in entries:
        numbers, originals = _영수증칸(e)
        ws.append(
            [
                e.level1, e.level2, e.level3a, e.level3b, e.level3c,
                e.department.name if e.department else None,
                numbers, originals,
                e.expense_date,
                e.amount,
                # 같은 식을 다시 적지 않는다 — 집행 금액은 모델의 한 곳에서
                e.settlement_amount,
                e.personal_burden_amount if e.is_meal_expense else 0,
                e.meal_headcount,
                " ".join(e.meal_attendee_names or []) or None,
                e.note,
                "지급완료" if e.paid else "미지급",
                e.paid_date,
                e.payer_name,
            ]
            + (_계좌칸(e) if 계좌를_보인다 else [])
        )
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=12):
        for cell in row:
            cell.number_format = MONEY
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    _style_header(ws, len(headers))
    _widths(ws, [12, 16, 14, 14, 14, 10, 10, 12, 12, 12, 12, 12, 9, 30, 24, 10, 12, 10]
            + ([10, 20, 10] if 계좌를_보인다 else []))

    # 2) 예산 대비 집행 — 구분 소계·총계·비율 전부 summary 의 값이다 (7-3)
    ws2 = wb.create_sheet("예산 대비 집행")
    ws2.append(["구분", "항목", "세부항목", "예산금액", "결산금액", "차액", "비율(%)", "집행률(%)"])
    for group in summary.groups:
        for row in group.rows:
            cat = row.category
            ws2.append(
                [cat.level1, cat.level2, cat.level3,
                 row.planned, row.spent, row.remaining, row.ratio_pct, row.progress_pct]
            )
        sub = ws2.max_row + 1
        ws2.append([f"{group.name} 소계", "", "", group.planned, group.spent, group.remaining, "", ""])
        for col in range(1, 9):
            ws2.cell(row=sub, column=col).font = BOLD
    total_row = ws2.max_row + 1
    ws2.append(
        ["총계", "", "", summary.total_planned, summary.total_spent,
         summary.total_remaining, "", summary.progress_pct]
    )
    for col in range(1, 9):
        ws2.cell(row=total_row, column=col).font = BOLD
    if summary.uncategorized_spent:
        ws2.append(["(예산 항목 미지정 지출)", "", "", 0, summary.uncategorized_spent, "", "", ""])
    # 취소된 예산 항목은 줄로 안 싣지만, 거기 걸린 지출의 집행은 총계에 들어 있다 —
    # 줄이 없으면 총계와 줄의 합이 안 맞아 보이므로 따로 한 줄 (7-3)
    if summary.canceled_category_spent:
        ws2.append(["(취소된 예산 항목에 걸린 지출)", "", "", 0,
                    summary.canceled_category_spent, "", "", ""])
    if summary.active_incomes:
        ws2.append([])
        ws2.append(["수입", "", "", "", "", "", "", ""])
        for income in summary.active_incomes:
            ws2.append([income.name, income.note, "", income.amount, "", "", "", ""])
        ws2.append(["총 수입", "", "", summary.total_income, "", "", "", ""])
        ws2.append(["잔액 (총 수입 − 총 지출예산)", "", "", summary.balance, "", "", "", ""])
    for row in ws2.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = MONEY
    _style_header(ws2, 8)
    _widths(ws2, [16, 20, 18, 14, 14, 14, 10, 10])

    # 3) 환급 대상자 — 개인이 대신 낸 것 (7-4). 지급된 것도 함께 싣되 표시한다
    ws3 = wb.create_sheet("환급 대상자")
    # 열 이름은 「환급액」 — 비식대 행에는 전액이 들어가므로 「지원금액」 이라
    # 적으면 열 이름과 값의 뜻이 어긋난다
    # 원본 영수증번호는 지출 상세내역 시트에만 싣는다 — 환급 시트는 사람과 돈을 보는 자리다
    환급열 = ["영수증번호", "항목", "지출일자", "환급액", "지출자"]
    if 계좌를_보인다:
        환급열 += 계좌머리
    환급열.append("지급여부")
    ws3.append(환급열)
    refund = refund_sheet(entries)
    for e in refund.rows:
        ws3.append(
            [
                _영수증칸(e)[0],
                e.budget_category.display_name if e.budget_category else None,
                e.expense_date,
                e.settlement_amount,
                e.payer_name,
            ]
            + (_계좌칸(e) if 계좌를_보인다 else [])
            + ["지급완료" if e.paid else "미지급"]
        )
    ws3.append([])
    # 합계·건수는 domain.budget.refund_sheet 가 센다 — 파일은 놓기만 한다
    ws3.append(["미지급 합계", "", "", refund.unpaid_total, ""]
               + ([""] * len(계좌머리) if 계좌를_보인다 else [])
               + [f"{refund.unpaid_count}건"])
    for row in ws3.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = MONEY
    # 계좌 칸을 안 만들면 열이 줄어든다 — 숫자로 박으면 머리줄만 남는다
    _style_header(ws3, len(환급열))
    # 계좌 칸이 없으면 폭도 빠진다 — 그대로 두면 계좌 폭이 지급여부에 간다
    _widths(ws3, [12, 26, 12, 14, 12] +([10, 20, 10] if 계좌를_보인다 else []) + [10])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
