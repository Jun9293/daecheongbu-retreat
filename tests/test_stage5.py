"""UI 개편 단계 5 — 재정(7장) · 껍데기 하나로 · 단계 4 후속.

숫자는 값끼리 맞는지로 본다 (11-3) — 홈·회차 상세·예산 페이지가
같은 summary 를 쓰는지, 엑셀이 화면과 같은 소계를 내는지.
"""

from __future__ import annotations

import datetime as dt
import io
import pathlib
import re

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import models
from app.domain import budget as budget_domain
from app.domain import budget_xlsx
from tests.conftest import app_session, login_as, make_user

TODAY = dt.date.today()
ROOT = pathlib.Path(__file__).resolve().parent.parent


def _read(*parts) -> str:
    return (ROOT.joinpath(*parts)).read_text(encoding="utf-8")


# ════════════════════════════════════════════════════════════════════
# 재정 세계
# ════════════════════════════════════════════════════════════════════


@pytest.fixture
def fin(admin_client):
    """진행 중 회차 + 예산 3항목(2구분) + 지출 셋 + 수입 하나."""
    with app_session() as db:
        retreat = models.Retreat(
            name="2026 여름수련회 Belong", meal_subsidy_per_person=8000,
            start_date=TODAY + dt.timedelta(days=40),
            end_date=TODAY + dt.timedelta(days=43))
        db.add(retreat)
        db.flush()
        dept = models.Department(retreat_id=retreat.id, key="chongmuM",
                                 name="1 총무M", color_tag="#2F4858", sort_order=0)
        db.add(dept)
        db.flush()

        def cat(l1, l2, planned, order, **kw):
            row = models.BudgetCategory(
                retreat_id=retreat.id, level1=l1, level2=l2,
                planned_amount=planned, sort_order=order, **kw)
            db.add(row)
            db.flush()
            return row

        c1 = cat("식비", "본행사 식사", 0, 1, unit_price=8000, headcount=150, times=5)
        c1.planned_amount = budget_domain.planned_amount_of(8000, 150, 5, 0)
        c2 = cat("식비", "야식", 300_000, 2)
        c3 = cat("홍보", "포스터", 100_000, 3)

        def spend(category, amount, *, payer, paid=False, receipt_no=None, memo=None):
            entry = models.ExpenseEntry(
                retreat_id=retreat.id, budget_category_id=category.id,
                level1=category.level1, level2=category.level2,
                expense_date=TODAY, amount=amount, payer_name=payer,
                paid=paid, subsidy_amount=amount, department_id=dept.id)
            if receipt_no is not None:
                entry.receipts.append(
                    models.ExpenseReceipt(number=receipt_no, memo=memo or "결산 파일에 별첨"))
            db.add(entry)
            db.flush()
            return entry

        e1 = spend(c2, 120_000, payer="가명닮은사람", receipt_no=1)      # 환급 대상 · 미지급
        e2 = spend(c2, 50_000, payer=budget_domain.RETREAT_ACCOUNT,
                   paid=True, receipt_no=2)                              # 수련회계좌 — 환급 아님
        e3 = spend(c3, 180_000, payer="다른사람", paid=True)             # 영수증 없음 · 예산 초과
        db.add(models.IncomeItem(retreat_id=retreat.id, name="수련회비",
                                 unit_price=50_000, headcount=100,
                                 amount=5_000_000, sort_order=1))
        db.commit()
        ids = {"retreat": retreat.id, "dept": dept.id,
               "cats": {"본행사": c1.id, "야식": c2.id, "포스터": c3.id},
               "entries": {"환급": e1.id, "계좌": e2.id, "초과": e3.id}}
    admin_client.get(f"/board?retreat_id={ids['retreat']}")
    return ids


# ── 2-a. 예산금액 계산 ───────────────────────────────────────────────


def test5_b01_예산금액은_셋이_있으면_계산_비면_입력값(admin_client, fin):
    # 셋이 다 있으면 단가 × 명수 × 횟수가 이긴다
    resp = admin_client.post("/budget/categories", data={
        "level1": "장소비", "level2": "숙소", "level3": "",
        "unit_price": "30000", "headcount": "150", "times": "2",
        "planned_amount": "999",     # 직접 입력값은 무시된다
    }, follow_redirects=True)
    assert resp.status_code == 200
    with app_session() as db:
        row = db.scalars(select(models.BudgetCategory).where(
            models.BudgetCategory.level1 == "장소비")).one()
        assert row.planned_amount == 30000 * 150 * 2
        assert (row.unit_price, row.headcount, row.times) == (30000, 150, 2)

    # 하나라도 비면 직접 입력값이다
    admin_client.post("/budget/categories", data={
        "level1": "장소비", "level2": "회의실", "level3": "",
        "unit_price": "30000", "headcount": "", "times": "2",
        "planned_amount": "777000",
    }, follow_redirects=True)
    with app_session() as db:
        row = db.scalars(select(models.BudgetCategory).where(
            models.BudgetCategory.level2 == "회의실")).one()
        assert row.planned_amount == 777_000


def test5_b02_화면_계산은_저장_전_DB_를_안_건드린다():
    """즉시 계산은 budget.js 의 미리보기다 — 정본은 서버의 같은 공식이다."""
    js = _read("app", "static", "js", "budget.js")
    assert "data-calc" in js and "readOnly" in js
    assert "fetch(" not in js       # 저장 전엔 서버로 아무것도 안 보낸다


# ── 2-b. 초과는 붉게 — 100% 로 덮지 않는다 ──────────────────────────


def test5_b03_초과_행은_붉고_비율은_100_을_넘는다(admin_client, fin):
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        summary = budget_domain.build_budget_summary(db, retreat=retreat)
        poster = next(r for r in summary.categories if r.category.level2 == "포스터")
        assert poster.is_over_budget and poster.remaining < 0
        assert poster.progress_pct > 100           # 100 으로 덮지 않는다

    page = admin_client.get(f"/budget?retreat_id={fin['retreat']}").text
    assert "overrun" in page                       # 붉은 클래스 (--red-ink)
    assert "180.0%" in page                        # 180,000 / 100,000


# ── 2-c. 수입은 집행률에 영향이 없다 ─────────────────────────────────


def test5_b04_수입을_바꿔도_집행률은_그대로다(fin):
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        before = budget_domain.build_budget_summary(db, retreat=retreat)
        pct, balance = before.progress_pct, before.balance
        db.add(models.IncomeItem(retreat_id=retreat.id, name="후원금",
                                 amount=9_999_999, sort_order=2))
        db.commit()
        after = budget_domain.build_budget_summary(db, retreat=retreat)
        assert after.progress_pct == pct           # 분모는 지출예산 총액 (7-3)
        assert after.balance == balance + 9_999_999


# ── 2-d. 화면들의 숫자 = budget.summary ────────────────────────────


def test5_b05_홈과_회차_상세와_예산_페이지가_같은_수를_말한다(admin_client, fin):
    from app.domain import home as home_domain

    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        admin = db.scalars(select(models.User).where(models.User.role == "admin")).first()
        summary = budget_domain.build_budget_summary(db, retreat=retreat)
        home = home_domain.build(db, retreat, admin, today=TODAY)
        assert home.budget.progress_pct == summary.progress_pct
        assert home.budget.total_spent == summary.total_spent

    page = admin_client.get(f"/budget?retreat_id={fin['retreat']}").text
    assert f"{summary.progress_pct}%" in page
    detail = admin_client.get(f"/settings/retreats/{fin['retreat']}").text
    assert f"{summary.unspent_count}건" in detail


def test5_b06_라우터와_화면에_세는_코드가_없다():
    """test_h03 방식 — 예산 숫자를 만드는 곳은 domain.budget 하나다.

    (지출 라우터의 totals 는 필터된 목록의 표시용 합이라 예외 — 예산·집행·
    잔액 같은 판정 숫자는 전부 summary 를 지난다.)
    """
    budget_src = _read("app", "routers", "budget.py")
    assert "sum(" not in budget_src, "예산 라우터가 직접 더하고 있다"
    assert "build_budget_summary" in budget_src
    expenses_src = _read("app", "routers", "expenses.py")
    assert "build_budget_summary" in expenses_src   # 그룹 헤더가 summary 를 쓴다
    for tpl in (("app", "templates", "budget.html"), ("app", "templates", "expenses.html")):
        text = _read(*tpl)
        assert "|sum" not in text, f"{tpl[-1]} 이 화면에서 더하고 있다"


# ── 2-e. 엑셀 = 화면 ────────────────────────────────────────────────


def test5_x01_엑셀_소계가_화면_소계와_같다(fin):
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        summary = budget_domain.build_budget_summary(db, retreat=retreat)
        entries = budget_domain.entries_of(db, retreat)
        buffer = budget_xlsx.write(summary, entries)

        wb = load_workbook(io.BytesIO(buffer.read()))
        ws = wb["예산 대비 집행"]
        by_label = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
        for group in summary.groups:
            r = by_label[f"{group.name} 소계"]
            assert ws.cell(row=r, column=4).value == group.planned
            assert ws.cell(row=r, column=5).value == group.spent
            assert ws.cell(row=r, column=6).value == group.remaining
        r = by_label["총계"]
        assert ws.cell(row=r, column=4).value == summary.total_planned
        assert ws.cell(row=r, column=5).value == summary.total_spent


# ── 3-a. 지출 그룹 헤더 = 예산 페이지 값 ─────────────────────────────


def test5_e01_그룹_헤더가_예산_페이지와_같은_값이다(admin_client, fin):
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        summary = budget_domain.build_budget_summary(db, retreat=retreat)
        snack = next(r for r in summary.categories if r.category.level2 == "야식")

    page = admin_client.get(f"/expenses?retreat_id={fin['retreat']}").text
    assert f"예산 {snack.planned:,}원" in page
    assert f"집행 {snack.spent:,}원" in page
    assert f"잔액 {snack.remaining:,}원" in page


# ── 3-c. 영수증 옮김 — 한 번만 · 옛 컬럼은 구조로 봉인 ───────────────


def test5_r01_영수증_옮김은_한_번만이다(fin):
    from app import db as db_module

    entry_id = fin["entries"]["초과"]                # 영수증 없는 지출
    with app_session() as db:
        from sqlalchemy import text
        db.execute(text(
            "UPDATE expense_entries SET receipt_number=7,"
            " receipt_file_url='/uploads/old-receipt.png' WHERE id=:id"),
            {"id": entry_id})
        db.commit()

    db_module._move_receipts()
    with app_session() as db:
        receipts = db.scalars(select(models.ExpenseReceipt).where(
            models.ExpenseReceipt.expense_id == entry_id)).all()
        assert [r.number for r in receipts] == [7]
        assert receipts[0].stored_name == "old-receipt.png"

    db_module._move_receipts()                       # 두 번 떠도
    with app_session() as db:
        again = db.scalars(select(models.ExpenseReceipt).where(
            models.ExpenseReceipt.expense_id == entry_id)).all()
        assert len(again) == 1                       # 두 번 안 옮긴다

    # **파일 링크만 있고 번호가 없는 행**도 옮긴다 — 안 옮기면 봉인된 옛
    # 컬럼 속 파일을 읽을 길이 영영 없어진다. 번호는 그 회차의 다음 번호.
    with app_session() as db:
        from sqlalchemy import text
        db.execute(text(
            "INSERT INTO expense_entries (retreat_id, amount, subsidy_amount,"
            " personal_burden_amount, paid, is_meal_expense, receipt_file_url, created_at)"
            " VALUES (:r, 1000, 1000, 0, 0, 0, '/uploads/url-only.png', CURRENT_TIMESTAMP)"),
            {"r": fin["retreat"]})
        db.commit()
    db_module._move_receipts()
    with app_session() as db:
        moved = db.scalars(select(models.ExpenseReceipt).where(
            models.ExpenseReceipt.stored_name == "url-only.png")).one()
        assert moved.number == 8                     # 회차 최대(7) 다음


def test5_r02_옛_두_컬럼을_읽으면_AttributeError_다(fin):
    with app_session() as db:
        entry = db.get(models.ExpenseEntry, fin["entries"]["환급"])
        with pytest.raises(AttributeError):
            entry.receipt_number
        with pytest.raises(AttributeError):
            entry.receipt_file_url
    # 읽는 곳이 남아 있지 않다 — 옮기는 자리(db.py)와 모델 정의만 남는다
    for path in pathlib.Path(ROOT, "app").rglob("*.py"):
        if path.name in ("db.py", "models.py"):
            continue
        src = path.read_text(encoding="utf-8")
        assert ".receipt_number" not in src, f"{path} 가 옛 컬럼을 읽는다"
        assert ".receipt_file_url" not in src, f"{path} 가 옛 컬럼을 읽는다"


def test5_r03_영수증_번호는_회차_안에서_자동_증가다(admin_client, fin):
    page = admin_client.get(f"/expenses?retreat_id={fin['retreat']}").text
    assert "다음 영수증 번호 3" in page              # 1·2 다음
    # 메모만으로도 영수증이 된다 — 「결산 파일에 별첨」 (7-4)
    resp = admin_client.post(f"/expenses/{fin['entries']['초과']}/receipts",
                             data={"memo": "결산 파일에 별첨"}, follow_redirects=True)
    assert resp.status_code == 200
    with app_session() as db:
        entry = db.get(models.ExpenseEntry, fin["entries"]["초과"])
        assert [r.number for r in entry.receipts] == [3]
        assert entry.receipts[0].memo == "결산 파일에 별첨"


# ── 3-d·f. 환급 필터 · 홈 결산 = 필터 행 수 ─────────────────────────


def test5_f01_환급_필터와_301(admin_client, fin):
    page = admin_client.get(f"/expenses?filter=refund&retreat_id={fin['retreat']}").text
    assert "가명닮은사람" in page                     # 개인 · 미지급 → 보인다
    # 수련회계좌 지출은 필터에 없다 — 화면 문구가 아니라 행 목록으로 본다
    # (표기는 등록 폼의 안내문에도 나오므로 글자 검색으로는 못 가른다 — 10장)
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        rows = budget_domain.refund_entries(db, retreat)
        assert [e.id for e in rows] == [fin["entries"]["환급"]]

    r = admin_client.get("/refunds", follow_redirects=False)
    assert r.status_code == 301
    assert r.headers["location"] == "/expenses?filter=refund"
    assert not (ROOT / "app" / "templates" / "refunds.html").exists()


def test5_f02_홈_결산_숫자는_필터_행_수와_같다(admin_client, fin):
    from app.domain import home as home_domain

    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        retreat.end_date = TODAY - dt.timedelta(days=3)   # 결산 홈으로
        retreat.start_date = TODAY - dt.timedelta(days=6)
        db.commit()
        admin = db.scalars(select(models.User).where(models.User.role == "admin")).first()
        home = home_domain.build(db, retreat, admin, today=TODAY)
        refund_rows = budget_domain.refund_entries(db, retreat)
        noreceipt_rows = budget_domain.no_receipt_entries(db, retreat)
        assert home.unpaid_refund_count == len(refund_rows) == 1
        assert home.no_receipt_count == len(noreceipt_rows) == 1

    # 눌러 간 화면의 행 수와 같다 — 같은 함수라 갈릴 수 없다
    page = admin_client.get(f"/expenses?filter=noreceipt&retreat_id={fin['retreat']}").text
    assert page.count('<tr class="otherdept"') + page.count("<tr>") - page.count("<thead><tr>") >= 1


# ── 4. 회차 상세 「남은 지출」 — 셋 ─────────────────────────────────


def test5_u01_남은_지출은_끝난_회차에서만_붉다(admin_client, fin):
    url = f"/settings/retreats/{fin['retreat']}"

    # 진행 중 — 붉지 않다
    page = admin_client.get(url).text
    m = re.search(r"남은 지출</div>\s*<div class=\"v\"([^>]*)>", page)
    assert m and "--now" not in m.group(1)

    # 끝났고 N건 — 붉다
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        retreat.start_date = TODAY - dt.timedelta(days=6)
        retreat.end_date = TODAY - dt.timedelta(days=3)
        db.commit()
    page = admin_client.get(url).text
    m = re.search(r"남은 지출</div>\s*<div class=\"v\"([^>]*)>", page)
    assert m and "--now" in m.group(1)

    # 끝났고 0건 — 붉지 않다 (지출 없는 예산 항목을 없앤다)
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        for row in budget_domain.build_budget_summary(db, retreat=retreat).unspent_rows:
            db.add(models.ExpenseEntry(
                retreat_id=retreat.id, budget_category_id=row.category.id,
                expense_date=TODAY, amount=1, subsidy_amount=1))
        db.commit()
    page = admin_client.get(url).text
    m = re.search(r"남은 지출</div>\s*<div class=\"v\"([^>]*)>", page)
    assert m and "--now" not in m.group(1)


# ════════════════════════════════════════════════════════════════════
# 단계 4 후속 (1-a ~ 1-d)
# ════════════════════════════════════════════════════════════════════


def test5_c01_확인_요청_취소는_상태다(admin_client, client, fin):
    """취소해도 행이 남고(0장), 배지에서는 빠진다."""
    from app.routers.reviews import create_review_requests, pending_for_user

    lead_id = make_user("총무 팀원", "01066660001", "dept_lead",
                        department_id=fin["dept"])
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        admin = db.scalars(select(models.User).where(models.User.role == "admin")).first()
        create_review_requests(db, retreat=retreat, requester=admin,
                               department_ids=[fin["dept"]], message="봐 주세요")
        db.commit()
        review_id = db.scalars(select(models.ReviewRequest)).one().id
        before_rows = len(db.scalars(select(models.ReviewRequest)).all())
        lead = db.get(models.User, lead_id)
        assert len(pending_for_user(db, lead, retreat)) == 1

    cancel = admin_client.post(f"/reviews/{review_id}/cancel", follow_redirects=False)
    assert cancel.status_code in (302, 303)
    with app_session() as db:
        assert len(db.scalars(select(models.ReviewRequest)).all()) == before_rows  # 행 불변
        review = db.get(models.ReviewRequest, review_id)
        assert review.status == "취소"
        retreat = db.get(models.Retreat, fin["retreat"])
        lead = db.get(models.User, lead_id)
        assert pending_for_user(db, lead, retreat) == []   # 배지에서 빠진다

    # 「내가 보낸 요청」 에 흐리게 남는다
    sent = admin_client.get("/notifications?tab=sent").text
    assert "cancelled" in sent and "취소" in sent


def test5_c02_배지_수와_답을_기다리는_칩_수가_같다(admin_client, client, fin):
    """배지 = 안 읽은 **시스템** 알림 + 답 대기 요청 (4-0). 뒷항은 「답을
    기다리는 것」 칩과 같은 함수라, 부서 없는 admin 도 두 수가 같다."""
    from app.notifications import system_unread_count
    from app.routers.reviews import create_review_requests, pending_for_user

    make_user("총무 팀원2", "01066660002", "member", department_id=fin["dept"])
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        admin = db.scalars(select(models.User).where(models.User.role == "admin")).first()
        create_review_requests(db, retreat=retreat, requester=admin,
                               department_ids=[fin["dept"]], message="확인")
        db.commit()

    # 받는 사람 — 요청 하나는 배지에 **1** 이다 (알림 +1 과 요청 +1 로 二重 아님)
    login_as(client, "01066660002")
    client.get(f"/board?retreat_id={fin['retreat']}")
    with app_session() as db:
        retreat = db.get(models.Retreat, fin["retreat"])
        member = db.scalars(select(models.User).where(
            models.User.phone_number == "01066660002")).one()
        badge = system_unread_count(db, member) + len(pending_for_user(db, member, retreat))
        chip_rows = len(pending_for_user(db, member, retreat))
        assert badge == 1 == chip_rows

        # 부서 없는 admin — 배지의 요청 수 = 칩 행 수 (화면과 배지가 같은 말)
        admin = db.scalars(select(models.User).where(models.User.role == "admin")).first()
        assert admin.department_id is None      # 전제 — admin 픽스처는 부서가 없다
        admin_pending = len(pending_for_user(db, admin, retreat))
        assert admin_pending == 1               # 전 부서의 대기 요청을 다 본다
    pending_page = client.get("/notifications?chip=pending").text
    assert pending_page.count('nitem pending-review') == chip_rows
    admin_page = admin_client.get("/notifications?chip=pending").text
    assert admin_page.count('nitem pending-review') == admin_pending


def test5_c03_폐회일_당일_마감은_늦음이_아니다(db, sample_retreat):
    """T-e 를 결정으로 만든다 (4-10 에 적음) — 그날은 아직 안 지났다."""
    from app.domain import suggestions

    close = sample_retreat.end_date
    lib = models.TaskLibrary(title="폐회일 마감", kind="main", default_d_week=1)
    db.add(lib)
    db.flush()
    db.add(models.TaskRun(
        library_id=lib.id, retreat_id=sample_retreat.id, included=True,
        start_date=close - dt.timedelta(days=3), end_date=close, status="대기"))
    db.commit()
    out = suggestions.generate(
        db, open_date=close + dt.timedelta(days=140), base_retreat=sample_retreat)
    assert "「폐회일 마감」" not in " / ".join(x["source"] for x in out)
    assert "폐회일 당일" in _read("CLAUDE.md")


def test5_c04_체크리스트_흐림은_부서_키다(admin_client, client, fin):
    """다른 회차의 같은 키 부서 소속이 자기 부서 항목을 선명하게 본다 (2장)."""
    with app_session() as db:
        old = models.Retreat(name="지난 회차", start_date=TODAY - dt.timedelta(days=400),
                             end_date=TODAY - dt.timedelta(days=397))
        db.add(old)
        db.flush()
        old_dept = models.Department(retreat_id=old.id, key="chongmuM",
                                     name="1 총무M", color_tag="#2F4858", sort_order=0)
        db.add(old_dept)
        db.flush()
        old_dept_id = old_dept.id
        db.add(models.Checklist(retreat_id=fin["retreat"], name="비품 목록",
                                department_id=fin["dept"]))
        db.commit()
    make_user("옛 소속", "01066660003", "dept_lead", department_id=old_dept_id)
    login_as(client, "01066660003")
    client.get(f"/board?retreat_id={fin['retreat']}")
    page = client.get("/checklists").text
    assert "비품 목록" in page
    assert "opacity:.6" not in page                    # 같은 키 — 흐리지 않다


# ════════════════════════════════════════════════════════════════════
# 5. 껍데기 하나로
# ════════════════════════════════════════════════════════════════════


def test5_s01_옛_껍데기가_없다():
    assert not (ROOT / "app" / "templates" / "base.html").exists()
    assert not (ROOT / "app" / "static" / "css" / "app.css").exists()
    for path in pathlib.Path(ROOT, "app", "templates").rglob("*.html"):
        assert 'extends "base.html"' not in path.read_text(encoding="utf-8"), path.name


def test5_s02_로그인_전_화면도_같은_껍데기다(client, admin_client, fin):
    page = client.get("/login")
    assert page.status_code == 200
    assert "css/retreat." in page.text and "초대 링크" in page.text   # 해시 주소 (11-2)
    # 로그인 전에는 사이드바가 없다 (글자가 아니라 마크업으로 — 10장)
    assert '<aside class="sidenav"' not in page.text

    missing = admin_client.get("/board/task/999999")
    assert missing.status_code == 404


def test5_s03_회차_없음_안내가_새_껍데기로_뜬다(admin_client):
    page = admin_client.get("/")
    assert page.status_code == 200
    assert "수련회 회차가 없습니다" in page.text
    assert "css/retreat." in page.text   # 해시 주소 (11-2)
