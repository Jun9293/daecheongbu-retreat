"""못 이은 지출을 시트 결산 식 연결로 잇는다 (2026-09-15 · `scripts/재정짝잇기.py` · 봐둘것 BB-h).

실제 시트는 열지 않는다 — `tests/test_stage48.py` 의 합성 시트에 합계금액 식을 더해 쓴다. 값은 지어낸 것이다.
수는 합성 시트에서 정해지는 것만 쓴다(운영 수를 박지 않는다 · 11-3).
"""

from __future__ import annotations

import importlib.util
import pathlib
import sys

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import config, models
from tests.conftest import app_session
from tests.test_stage48 import _돌린다, 판  # noqa: F401 — 판 은 fixture

ROOT = pathlib.Path(__file__).resolve().parent.parent


def _짝잇기():
    이름 = "재정짝잇기_시험"
    spec = importlib.util.spec_from_file_location(이름, ROOT / "scripts" / "재정짝잇기.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[이름] = mod
    spec.loader.exec_module(mod)
    return mod


def _합계식(경로: pathlib.Path) -> None:
    """합성 시트 지출 탭 H2 에 「그 묶음 줄들의 예산 사용액 합」 식을 넣는다 — 시트 관행(7-1)."""
    wb = load_workbook(경로)
    wb["지출 상세내역"]["H2"] = "=SUM(G2:G4)"
    wb.save(경로)


def _잇기돌림(판, 실행, 사본=False):
    with app_session() as db:
        return 판["짝"].돌린다(db, 판["경로"], 판["retreat"], 실행, 사본=사본)


@pytest.fixture
def 짝판(판):
    _합계식(판["경로"])
    판["짝"] = _짝잇기()
    return 판


def _옛들여오기(판):
    """2026 여름 시트를 들인 때의 모양 — 그때는 글자로 이어 글자가 겹치거나 없는 줄(지출 6 · 7행)이 못 이은 채 들어갔다.
    지금 들여오기는 결산 식으로 이으므로(BB-i) 들인 뒤 그 둘을 비워 그 모양을 만든다. **글자도 시트 글자로 되돌린다** —
    못 이은 줄은 시트 글자를 두었으므로, 안 되돌리면 짝잇기가 글자를 덮는지를 아무 시험도 못 잰다(커밋 전 검토 M1)."""
    _돌린다(판, True)
    with app_session() as db:
        for x in db.scalars(select(models.ExpenseEntry).where(
                models.ExpenseEntry.retreat_id == 판["retreat"], models.ExpenseEntry.canceled_at.is_(None),
                models.ExpenseEntry.amount.in_((10000, 5000)))):
            x.budget_category_id = None
            if x.amount == 5000:
                x.level3a = "없는세부"
        db.commit()


def _미지정(판):
    with app_session() as db:
        return {e.id: e for e in db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == 판["retreat"], models.ExpenseEntry.canceled_at.is_(None),
            models.ExpenseEntry.budget_category_id.is_(None)))}


# ── 가) 결산 식을 따라가 예산 줄을 구한다 ──


def test51_a01_결산식을_따라가_지출_줄마다_예산_줄(짝판):
    식 = 짝판["mod"].결산식으로(짝판["경로"])
    # 예산 3행 → H2(=SUM(G2:G4)) → 지출 2·3·4 · 예산 4행 → H5 · 5행 → H6 · 6행 → H7 · 7행 → H8
    assert {r: 식[r] for r in (2, 3, 4, 5, 6, 7, 8)} == {2: {3}, 3: {3}, 4: {3}, 5: {4}, 6: {5}, 7: {6}, 8: {7}}
    # 들여오기가 이 식으로 잇는다 — 지출 줄마다 식이 가리키는 그 예산 줄
    계획 = _돌린다(짝판, False)
    이은 = {e["줄"]: e["예산줄"] for e in 계획.지출}
    assert 이은 and all(식[r] == {b} for r, b in 이은.items())


def test51_a02_들여오기는_결산_식으로_잇고_글자는_확인용():
    """사람이 정함(2026-09-15 · 봐둘것 BB-i) — 2027 시트부터 처음부터 결산 식으로 잇는다."""
    글 = (ROOT / "scripts" / "재정들여오기.py").read_text(encoding="utf-8")
    고른다 = 글[글.index("def 고른다("):글.index("def _이름(")]
    assert "판.결산식" in 고른다 and "열쇠.get(key) == 1" in 고른다


# ── 다) 미리보기 · 라) 실행 · 마) 두 번째 ──


def test51_b01_미지정만_채우고_이미_이은_것은_안_건드린다(짝판, capsys):
    _옛들여오기(짝판)
    전미지정 = _미지정(짝판)
    with app_session() as db:
        이은것 = {e.id: e.budget_category_id for e in db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == 짝판["retreat"], models.ExpenseEntry.canceled_at.is_(None),
            models.ExpenseEntry.budget_category_id.is_not(None)))}
    미리 = _잇기돌림(짝판, False)
    assert 미리.수["예산 항목 미지정"] == len(전미지정) > 0
    assert 미리.수["채울 것"] == len(미리.채울) == len(전미지정), "합성 시트의 못 이은 줄은 전부 식이 하나를 가리킨다"
    assert _미지정(짝판).keys() == 전미지정.keys(), "미리보기가 바꿨다"

    _잇기돌림(짝판, True)
    assert _미지정(짝판) == {}
    with app_session() as db:
        for 지출id, 예산id in 미리.채울:
            x, c = db.get(models.ExpenseEntry, 지출id), db.get(models.BudgetCategory, 예산id)
            assert x.budget_category_id == c.id and (x.level1, x.level2, x.level3a) == (c.level1, c.level2, c.level3)
        for 지출id, 예산id in 이은것.items():
            assert db.get(models.ExpenseEntry, 지출id).budget_category_id == 예산id, "이미 이은 지출을 건드렸다"
        assert db.scalars(select(models.ActivityLog).where(models.ActivityLog.action == 짝판["짝"].짝잇기행위)).one()

    두번째 = _잇기돌림(짝판, True)
    out = capsys.readouterr().out
    assert 두번째.채울 == [] and "채울 것이 없습니다" in out
    # 값을 안 찍는다 — 미리보기 · 실행 · 두 번째 출력 전부
    assert "시트 결산 식으로 이을 것" in out and "없는세부" not in out and "가명" not in out


def test51_b02_취소된_예산_항목을_가리키면_안_채우고_짝_표에(짝판):
    _옛들여오기(짝판)
    미리 = _잇기돌림(짝판, False)
    with app_session() as db:
        하나 = db.get(models.BudgetCategory, 미리.채울[0][1])
        하나.canceled_at = models._now()
        db.commit()
        취소id = 하나.id
    다시 = _잇기돌림(짝판, False)
    취소된 = sum(1 for _, 예산id in 미리.채울 if 예산id == 취소id)
    assert 다시.수["못 채움(가리키는 예산 항목이 취소됨)"] == 취소된 == len(다시.짝) >= 1
    assert len(다시.채울) == len(미리.채울) - 취소된
    _잇기돌림(짝판, True)
    assert len(_미지정(짝판)) == 취소된, "취소된 항목을 가리킨 지출만 미지정으로 남아야 한다"
    md = (pathlib.Path(config.DATA_DIR) / "재정짝잇기.real.md").read_text(encoding="utf-8")
    assert "취소 표시" in md


def test51_b03_들여온_기록이_없으면_멈춘다(짝판):
    with pytest.raises(짝판["mod"].멈춤, match="들여온 기록이 없습니다"):
        _잇기돌림(짝판, False)


def test51_b04_사본을_못_뜨면_아무것도_안_바꾼다(짝판):
    _옛들여오기(짝판)
    전 = _미지정(짝판).keys()
    with pytest.raises(짝판["mod"].멈춤, match="다릅니다"):
        _잇기돌림(짝판, True, 사본=True)
    assert _미지정(짝판).keys() == 전


def test51_b05_들인_뒤_지출이_바뀌어_순서가_안_맞으면_멈춘다(짝판):
    _옛들여오기(짝판)
    with app_session() as db:
        x = next(iter(db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == 짝판["retreat"], models.ExpenseEntry.canceled_at.is_(None)))))
        x.amount += 1
        db.commit()
    with pytest.raises(짝판["mod"].멈춤, match="순서대로 맞지 않습니다"):
        _잇기돌림(짝판, False)


# ── 막는 갈래 (11-3 · 커밋 전 검토) ──


def test51_c01_아는_꼴의_식만_줄로_푼다(짝판):
    풀기 = 짝판["mod"]._식의줄들
    탭 = 짝판["mod"].지출탭
    assert 풀기(f"='{탭}'!$H$5", 탭) == {5}
    # Excel 기본 범위 꼴 — 탭 이름이 범위 앞에 한 번(사람이 정함 · 2026-09-15) · 두 끝에 다 붙인 꼴도
    assert 풀기(f"='{탭}'!G2:G4", 탭) == {2, 3, 4} and 풀기(f"=SUM('{탭}'!$G$4:$G$2)", 탭) == {2, 3, 4}
    assert 풀기(f"='{탭}'!G2:'{탭}'!G4", 탭) == {2, 3, 4}
    assert 풀기(f"=G2:'{탭}'!G4", 탭) == set(), "뒤 끝에만 붙은 꼴은 모른다"
    assert 풀기(f"='{탭}'!G2:G4") == set(), "합계 식(탭 없음)에 탭 이름이 섞이면 모른다"
    assert 풀기("=sum(G2:G4)") == {2, 3, 4} and 풀기("=SUM(G4:G2)") == {2, 3, 4} and 풀기("=G2+G3") == {2, 3}
    # 모르는 꼴은 빈 것 — 틀린 한 줄로 조용히 채우지 않는다
    for 모름 in ("=LOG10(G2)", "=SUMIF(C:C,C5,G:G)", '=H5&"A1"', "=G2-G3", "=G2*2", "G2", ""):
        assert 풀기(모름) == set(), 모름
    assert 풀기(f"='{탭}'!H5-D3", 탭) == set(), "다른 탭 주소가 섞인 결산 식"
    assert 풀기("=H5", 탭) == set(), "결산 식은 지출 탭을 가리켜야 한다"


def test51_c02_창이_두_곳에서_맞으면_빈_것(짝판):
    창 = 짝판["짝"]._창
    assert 창([1, 2, 1, 2], [1, 2], lambda v: v) == [], "두 곳에서 맞는데 하나를 골랐다"
    assert 창([9, 1, 2], [1, 2], lambda v: v) == [1, 2]


def test51_c03_들인_뒤_예산_항목_이름을_고치면_멈춘다(짝판):
    _옛들여오기(짝판)
    with app_session() as db:
        c = db.scalars(select(models.BudgetCategory).where(
            models.BudgetCategory.retreat_id == 짝판["retreat"], models.BudgetCategory.canceled_at.is_(None))).first()
        c.level2 = (c.level2 or "") + "고침"
        db.commit()
    with pytest.raises(짝판["mod"].멈춤, match="순서대로 맞지 않습니다"):
        _잇기돌림(짝판, False)


def _결산식(판, 줄, 식):
    wb = load_workbook(판["경로"])
    wb["예산(실시간)"].cell(줄, 11, 식)
    wb.save(판["경로"])


def test51_c04_식이_예산_줄_둘을_가리키거나_모르는_꼴이면_짝_표에(짝판):
    _옛들여오기(짝판)                          # 들인 뒤 시트의 식만 바꾼다 — 글자와 순서는 그대로
    미리 = _잇기돌림(짝판, False)
    assert len(미리.채울) == 2, "합성 시트의 못 이은 줄은 둘(6 · 7행)"
    _결산식(짝판, 5, "='지출 상세내역'!H6+'지출 상세내역'!H7")   # 7행이 예산 5·6행 둘에 잡힌다
    둘 = _잇기돌림(짝판, False)
    assert ("결산 식이 가리키는 예산 줄이 2개" in {까닭 for _, 까닭 in 둘.짝}) and len(둘.채울) == 1
    _결산식(짝판, 5, "=SUMIF('지출 상세내역'!C:C,1,'지출 상세내역'!G:G)")   # 모르는 꼴 — 6행이 아무 데도 안 잡힌다
    모름 = _잇기돌림(짝판, False)
    assert "결산 식이 가리키는 예산 줄이 0개" in {까닭 for _, 까닭 in 모름.짝}


def test51_c05_식이_들여온_예산_줄이_아닌_줄을_가리키면_짝_표에(짝판, monkeypatch):
    _옛들여오기(짝판)
    monkeypatch.setattr(짝판["짝"], "결산식으로", lambda 경로: {6: {99}, 7: {99}})
    판2 = _잇기돌림(짝판, False)
    assert 판2.채울 == [] and len(판2.짝) == 2
    assert {까닭 for _, 까닭 in 판2.짝} == {"결산 식이 가리키는 줄이 들여온 예산 줄이 아님"}


def test51_c06_이미_이은_것이_식과_다르면_센다(짝판):
    _옛들여오기(짝판)
    assert _잇기돌림(짝판, False).수["이미 이은 것 중 결산 식이 가리키는 항목과 다른 것"] == 0
    with app_session() as db:
        이은 = db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == 짝판["retreat"], models.ExpenseEntry.canceled_at.is_(None),
            models.ExpenseEntry.budget_category_id.is_not(None))).first()
        다른 = db.scalars(select(models.BudgetCategory).where(
            models.BudgetCategory.retreat_id == 짝판["retreat"], models.BudgetCategory.canceled_at.is_(None),
            models.BudgetCategory.id != 이은.budget_category_id)).first()
        이은.budget_category_id = 다른.id
        db.commit()
    assert _잇기돌림(짝판, False).수["이미 이은 것 중 결산 식이 가리키는 항목과 다른 것"] == 1
