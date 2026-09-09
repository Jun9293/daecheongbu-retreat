"""도막 1 — 관리자 복구 · 계좌 문.

**계좌를 담을 자리는 있는데 읽는 데 문이 없었다** (도막 0 의 2-(c)·2-(d)) —
로그인만 하면 `viewer` 까지 지출 목록에서 계좌를 전부 보고, 엑셀로 통째로
받아 갔다. 화면 하나만 막으면 파일로 새고, 파일만 막으면 화면으로 샌다.

**막는 코드에는 막히는 쪽 시험이 함께 있다** (11-3) — 셋이 한 벌이다:
① 막혀야 할 것이 막히는가 ② **통과해야 할 것(admin)이 통과하는가**
③ 검사가 볼 것을 보고 있는가(계좌 값이 실제로 그 회차에 있는가).
"""

from __future__ import annotations

import datetime as dt
import io
import pathlib
import re

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import models
from tests.conftest import app_session, login_as, make_user

ROOT = pathlib.Path(__file__).resolve().parent.parent
TODAY = dt.date.today()

# **시험이 쓰는 계좌 문자열은 저장소가 이미 싣고 있는 것이다** —
# `app/templates/expenses.html` 의 placeholder. 새 문자열을 지어내면
# 그것이 또 검사에 걸리고 넘김 목록만 는다 (도막 0 의 8장 C-4).
계좌 = "국민 000000-00-000000"


@pytest.fixture
def 돈회차(admin_client):
    """지출 하나에 계좌가 든 회차 + 부서 하나."""
    with app_session() as db:
        r = models.Retreat(name="계좌 회차", meal_subsidy_per_person=8000,
                           start_date=TODAY + dt.timedelta(days=30),
                           end_date=TODAY + dt.timedelta(days=33))
        db.add(r)
        db.flush()
        dept = models.Department(retreat_id=r.id, key="hebron", name="5 헤브론",
                                 color_tag="#4A8A5C", sort_order=0)
        db.add(dept)
        db.flush()
        # **식대로 넣습니다** — 등록 폼의 「직전 입력값 제안」(7-2)이
        # 식대에서만 돌기 때문입니다. 그 제안이 계좌가 폼으로 새던 자리라,
        # 식대가 아니면 g05 가 목록만 보고 폼은 한 번도 안 봅니다.
        db.add(models.ExpenseEntry(
            retreat_id=r.id, department_id=dept.id, expense_date=TODAY,
            amount=10_000, subsidy_amount=10_000, is_meal_expense=True,
            meal_headcount=2, level3b="모임 식사비-1",
            payer_name="지출한 사람", payer_account=계좌))
        db.commit()
        ids = {"retreat": r.id, "dept": dept.id}
    admin_client.get(f"/expenses?retreat_id={ids['retreat']}")
    return ids


def _본다(client, ids):
    return client.get(f"/expenses?retreat_id={ids['retreat']}").text


def _폼값(page: str, 이름: str) -> str:
    """등록 폼의 `name="…"` 입력칸에 미리 채워진 값."""
    m = re.search(rf'name="{이름}"[^>]*value="([^"]*)"', page)
    return m.group(1) if m else ""


# ════════════════════════════════════════════════════════════════════
# 2-c. 막히는 쪽 — viewer · dept_lead
# ════════════════════════════════════════════════════════════════════


def test23_g01_viewer_는_지출_목록에서_계좌를_못_본다(client, 돈회차):
    """**빈 칸이 아니라 없는 칸이다** — 자리만 남기면 「무엇이 가려져
    있다」 가 보이고, 그것도 알 필요 없는 사실이다."""
    make_user("보는 사람", "01088880001", "viewer")
    login_as(client, "01088880001")
    page = _본다(client, 돈회차)
    assert "지출한 사람" in page, "지출 목록 자체는 보여야 한다 (막은 것은 계좌뿐)"
    assert 계좌 not in page
    assert "enote" not in page, "빈 칸을 남기지 않는다"


def test23_g02_부서_리더도_못_본다(client, 돈회차):
    """편집자라고 남의 계좌를 볼 이유가 없다 — 지출을 등록하는 것과
    남이 적은 계좌를 읽는 것은 다른 일이다."""
    make_user("헤브론 리더", "01088880002", "dept_lead",
              department_id=돈회차["dept"])
    login_as(client, "01088880002")
    page = _본다(client, 돈회차)
    assert "지출한 사람" in page
    assert 계좌 not in page


def test23_g03_admin_은_본다(admin_client, 돈회차):
    """② **통과해야 할 것이 통과하는가.** 이것이 없으면 「아무에게도 안
    보인다」 와 「제대로 막았다」 가 구별되지 않는다 — 그리고 ③ 그 회차에
    계좌 값이 실제로 있다는 것도 이 줄이 함께 잰다."""
    assert 계좌 in _본다(admin_client, 돈회차)


def _엑셀칸(res) -> tuple[list, list]:
    """받은 xlsx 의 두 시트 머리줄 (지출 상세내역 · 환급 대상자)."""
    wb = load_workbook(io.BytesIO(res.content))
    머리 = lambda ws: [c.value for c in next(ws.iter_rows(max_row=1))]
    return 머리(wb["지출 상세내역"]), 머리(wb["환급 대상자"])


def test23_g04_엑셀은_편집자가_받되_계좌_칸은_admin_만(client, admin_client, 돈회차):
    """**막을 것은 표가 아니라 계좌다.** 전에는 파일을 통째로 막았는데,
    그러면 화면은 계좌만 빼고 누구나 보는데 파일은 아무도 못 받아
    **같은 표인데 보는 사람이 갈린다** (5-8 · 7-3).

    셋이 한 벌 — ① viewer 는 403 ② dept_lead 는 받되 계좌 칸이 **아예 없다**
    ③ admin 파일에는 그 칸이 있고 **값도 들어 있다**(없으면 「칸이 없다」 가
    「자료가 없다」 로도 통과한다)."""
    주소 = f"/export/expenses.xlsx?retreat_id={돈회차['retreat']}"

    ok = admin_client.get(주소)
    assert ok.status_code == 200
    assert ok.headers["content-type"].startswith(
        "application/vnd.openxmlformats"), "받아지는 것이 엑셀인지도 본다"
    지출머리, 환급머리 = _엑셀칸(ok)
    assert "지출자 계좌" in 지출머리 and "계좌" in 환급머리
    assert _값에계좌있나(ok), "총무팀 파일에 계좌 값이 실제로 들어 있어야 한다 (③)"

    # ② 편집자는 표를 받는다 — 계좌 칸만 없다
    make_user("엑셀 리더", "01088880004", "dept_lead", department_id=돈회차["dept"])
    login_as(client, "01088880004")
    리더 = client.get(주소)
    assert 리더.status_code == 200, "편집자가 표를 못 받는다"
    지출머리2, 환급머리2 = _엑셀칸(리더)
    assert "지출자 계좌" not in 지출머리2, "빈 칸이 아니라 없는 칸이어야 한다"
    assert "계좌" not in 환급머리2
    assert "지출자" in 지출머리2, "계좌만 빼야 하는데 표까지 줄었다"
    # xlsx 는 zip 이라 바이트를 뒤져서는 못 본다 — 칸을 읽어 본다
    assert not _값에계좌있나(리더)

    # ① 열람 전용은 못 받는다
    make_user("엑셀 보는 사람", "01088880003", "viewer", department_id=돈회차["dept"])
    login_as(client, "01088880003")
    assert client.get(주소).status_code == 403, "열람 전용이 엑셀을 받았다"


def _값에계좌있나(res) -> bool:
    """두 시트 어느 칸에든 계좌 문자열이 들어 있는가."""
    wb = load_workbook(io.BytesIO(res.content))
    return any(
        c.value is not None and 계좌 in str(c.value)
        for 이름 in ("지출 상세내역", "환급 대상자")
        for row in wb[이름].iter_rows(min_row=2) for c in row
    )


def test23_g05_등록_폼이_남의_계좌를_미리_채우지_않는다(client, 돈회차):
    """목록에서 계좌를 가려 놓고 **폼으로 새면** 가린 뜻이 없다 —
    직전 지출의 계좌가 다음 사람의 입력칸에 뜨던 자리다."""
    make_user("등록하는 리더", "01088880005", "dept_lead",
              department_id=돈회차["dept"])
    login_as(client, "01088880005")
    page = _본다(client, 돈회차)
    assert 계좌 not in page
    # **이름과 계좌는 함께 떨어진다** — 계좌만 자기 것으로 바꾸면
    # (남의 이름, 내 계좌) 가 되어 환급이 엉뚱한 곳으로 간다
    assert "지출한 사람" not in _폼값(page, "payer_name"), "남의 이름이 폼에 남았다"
    # ② 총무팀에게는 그대로 채워 준다 — 원래 다 보는 자리고 직전 값이 편하다
    with app_session() as db:
        admin = db.scalars(select(models.User).where(
            models.User.role == "admin")).first()
        admin_phone = admin.phone_number
    login_as(client, admin_phone)
    관리자쪽 = _본다(client, 돈회차)
    assert 계좌 in 관리자쪽
    # ③ **검사가 볼 것을 실제로 보고 있는가** — 관리자 폼에는 직전 지출자
    # 이름이 그대로 들어 있어야 한다. 이 줄이 없으면 위의 `not in` 은
    # 「폼 자체를 못 찾았다」 로도 통과한다
    assert _폼값(관리자쪽, "payer_name") == "지출한 사람"


def test23_g06_누를_수_없는_엑셀_단추를_그리지_않는다(client, admin_client, 돈회차):
    """10장이 「pytest 는 화면을 못 봅니다」 에 **누를 수 없는 단추**를 이미
    올려 두었다. 파일을 막았으면 그리로 가는 단추도 함께 감춰야 한다 —
    안 그러면 누른 사람이 403 화면을 본다.

    ② 총무팀에게는 남아 있어야 한다(안 그러면 기능이 사라진 것이다)."""
    주소 = "/export/expenses.xlsx"
    # ① 못 받는 사람에게는 안 그린다
    make_user("칩 보는 사람", "01088880006", "viewer", department_id=돈회차["dept"])
    login_as(client, "01088880006")
    for 화면 in ("/expenses", "/budget"):
        글 = client.get(f"{화면}?retreat_id={돈회차['retreat']}").text
        assert 주소 not in 글, f"{화면} 에 누를 수 없는 엑셀 단추가 남았다"
    # ② **받을 수 있는 사람에게는 그린다** — 계좌가 보이는지와 다른 물음이다.
    # 전에는 칩을 계좌 판정으로 그려서, 표를 받을 수 있는 부서 리더에게
    # 단추가 없었다
    make_user("칩 받는 리더", "01088880007", "dept_lead", department_id=돈회차["dept"])
    login_as(client, "01088880007")
    for 화면 in ("/expenses", "/budget"):
        글 = client.get(f"{화면}?retreat_id={돈회차['retreat']}").text
        assert 주소 in 글, f"{화면} 에서 편집자의 엑셀 단추까지 없앴다"
    for 화면 in ("/expenses", "/budget"):
        글 = admin_client.get(f"{화면}?retreat_id={돈회차['retreat']}").text
        assert 주소 in 글, f"{화면} 에서 총무팀의 엑셀 단추까지 없앴다"


# ════════════════════════════════════════════════════════════════════
# 1. 관리자 링크 — 만드는 길이 하나인가
# ════════════════════════════════════════════════════════════════════


def test23_h01_관리자링크가_초대_링크와_같은_함수를_부른다():
    """**링크 만드는 코드를 새로 쓰지 않는다** (4-12) — 두 곳에서 만들면
    한쪽만 고쳐진다.

    **낱말만 잰다** — 실제로 도는지는 서버 컴퓨터에서 세 번 돌려 본
    결과가 `docs/review/최근.md` 에 있다."""
    글 = (ROOT / "scripts/관리자링크.py").read_text(encoding="utf-8")
    코드 = re.sub(r'"""[\s\S]*?"""', "", 글)
    assert "invites.issue(" in 코드 and "invites.invite_url(" in 코드
    # 스스로 토큰을 만들지 않는다 — 그 자리가 둘이 되는 첫걸음이다
    for 만드는말 in ("token_urlsafe", "hash_token", "InviteToken("):
        assert 만드는말 not in 코드, f"링크를 여기서 만들고 있다: {만드는말}"


def test23_h02_admin_이_아니면_거절한다():
    """**낱말만 잰다** — 실제 거절은 서버에서 돌린 결과가 보고에 있다.
    여기서는 그 갈래가 코드에 있는지만 본다(지우면 빨개진다)."""
    코드 = (ROOT / "scripts/관리자링크.py").read_text(encoding="utf-8")
    assert "person.role != perm.ADMIN" in 코드
    assert "발급하지 않았습니다" in 코드


def test23_h04_거절_갈래를_실제로_밟는다(client):
    """**낱말만 재면 그 경로를 처음 밟는 것이 실제로 막아야 할 순간입니다**
    (11-3). `h02` 는 문자열만 보므로 `perm.ADMIN` 이 없어져도 초록이고
    부를 때 터집니다 — 그리고 그 순간은 「관리자가 다 막힌 날」 입니다.

    셋을 함께 잽니다 — ① 관리자가 아니면 1 로 거절하는가
    ② **관리자에게는 실제로 발급되는가** ③ 거절 쪽은 토큰을 안 만드는가
    (안 늘어야 막은 것이다)."""
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "관리자링크", ROOT / "scripts/관리자링크.py")
    모듈 = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(모듈)

    리더 = make_user("링크 못 받는 리더", "01088880011", "dept_lead")
    관리자 = make_user("링크 받는 총무", "01088880012", "admin")

    def 토큰수(db):
        return db.query(models.InviteToken).count()

    with app_session() as db:
        전 = 토큰수(db)
        assert 모듈.링크(db, 리더) == 1, "관리자가 아닌데 발급했다"
        assert 토큰수(db) == 전, "거절해 놓고 토큰을 만들었다"
        # ② 통과해야 할 쪽이 통과하는가 — 이것이 없으면 「아무에게도 안
        # 나간다」 와 「제대로 막았다」 가 구별되지 않는다
        assert 모듈.링크(db, 관리자) == 0
        assert 토큰수(db) == 전 + 1
        # ③ 없는 id 도 같은 자리에서 받는다
        assert 모듈.링크(db, 999_999) == 1


def test23_h03_돌리는_법이_배포_안내에_있다():
    글 = (ROOT / "docs/배포-안내.md").read_text(encoding="utf-8")
    assert "관리자링크.py" in 글, "돌리는 방법이 안내서에 없다"


# ════════════════════════════════════════════════════════════════════
# 3 · 5. 남기는 자리
# ════════════════════════════════════════════════════════════════════


def test23_i01_도막0_보고가_따로_남아_있다():
    """`최근.md` 는 **매번 덮인다** (11-3) — 덮이는 파일 하나에만 두면
    다음 판이 덮는 순간 그 셈이 통째로 사라진다."""
    p = ROOT / "docs/review/2026-09-09-도막0.md"
    assert p.exists(), "도막 0 보고가 안 남았다"
    글 = p.read_text(encoding="utf-8")
    assert "도막 0" in 글 and "짝짓기 열쇠" in 글
    봐둘것 = (ROOT / "docs/봐둘것.md").read_text(encoding="utf-8")
    assert "2026-09-09-도막0.md" in 봐둘것, "봐둘것이 그 파일을 안 가리킨다"


def test23_i02_인계_문서가_있고_열세_항목이_다_있다():
    """새 채팅이 **처음 읽는 파일**이다. 항목이 빠지면 그 일이 없어진다."""
    p = ROOT / "docs/인계.md"
    assert p.exists()
    글 = p.read_text(encoding="utf-8")
    # **부분문자열로 세지 않는다** — `11)` 이 `1)` 를 만족시켜서
    # 1~3 번이 통째로 빠져도 초록이었다.
    번호 = {int(m) for m in re.findall(r"^\s*(\d+)\)", 글, re.M)}
    assert 번호 >= set(range(1, 14)), f"빠진 항목: {sorted(set(range(1,14)) - 번호)}"
    for 말 in ("무엇", "재료", "미정"):
        assert 말 in 글, f"항목의 세 칸 중 「{말}」 이 없다"
    assert "docs/review/2026-09-09-도막0.md" in 글


def test23_i03_최근보고를_덮기_전에_복사하라는_단계가_있다():
    """이 판이 겪은 자리다 — 도막 0 의 350줄이 덮일 뻔했다."""
    글 = (ROOT / "CLAUDE.md").read_text(encoding="utf-8")
    자리 = 글[글.index("### 4. 채팅에 붙여넣을 것을 만든다"):]
    assert "날짜 이름" in 자리 and "복사" in 자리


def test23_i04_얼린_판만_빼고_살아있는_보고는_그대로_본다():
    """**빼는 코드에는 빼면 안 되는 쪽 시험이 함께 있다** (11-3).

    ① 얼린 판이 빠지는가 ② **`최근.md` 는 그대로 보는가** ③ 검사가
    볼 것을 실제로 보고 있는가(목록이 비면 통과가 아니다)."""
    from tests.test_docs_shell import 볼파일, 얼린판

    assert 얼린판(ROOT / "docs/review/2026-09-09-도막0.md")
    assert not 얼린판(ROOT / "docs/review/최근.md")

    본것 = {p for p, _ in 볼파일()}
    assert ROOT / "docs/review/최근.md" in 본것, "살아 있는 보고를 안 본다"
    assert ROOT / "CLAUDE.md" in 본것
    assert not any(얼린판(p) for p in 본것), "얼린 판이 아직 목록에 있다"
    assert len(본것) > 30, f"볼 것이 너무 적다: {len(본것)}"
