"""UI 정리 판 1 — 예산 표 · 봉사팀 보기 · 작은 것 셋.

셋이 한 벌이다 (11-3) — ① 막혀야 할 것이 막히는가 ② 막히면 안 되는 것이
통과하는가 ③ 검사가 볼 것을 실제로 보고 있는가. 엑셀 둘은 **골든**(`tests/golden/`)
과 견준다 — 판 전후로 셀 값·병합·행 수가 같아야 한다. 골든을 다시 뜨려면
`GOLDEN_WRITE=1` 로 돌린다 (보고에 왜 다시 떴는지 적는다).
"""

from __future__ import annotations

import datetime as dt
import io
import json
import os
import pathlib
import re

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import models
from app.domain import budget as budget_domain
from app.domain import permissions as perm
from tests.conftest import app_session, login_as, make_user

ROOT = pathlib.Path(__file__).resolve().parent.parent
GOLDEN = ROOT / "tests" / "golden"
TODAY = dt.date.today()
OPEN = dt.date(2026, 8, 21)
CLOSE = dt.date(2026, 8, 23)


# ── 재료 ─────────────────────────────────────────────────────────────

@pytest.fixture
def 예산(admin_client):
    """구분 둘(식비·홍보) · 항목 셋 · 세부항목 넷 — 식비 › 본행사 식사에 세부 둘.
    단가·명수·횟수는 첫 줄에만 있다. **수입은 없다** (2-a 를 재려고)."""
    with app_session() as db:
        r = models.Retreat(name="예산 정리 회차", meal_subsidy_per_person=8000,
                           start_date=TODAY + dt.timedelta(days=40),
                           end_date=TODAY + dt.timedelta(days=43))
        db.add(r); db.flush()
        d = models.Department(retreat_id=r.id, key="chongmuM", name="1 총무M",
                              color_tag="#2F4858", sort_order=0)
        db.add(d); db.flush()

        def cat(l1, l2, l3, planned, order, **kw):
            row = models.BudgetCategory(retreat_id=r.id, level1=l1, level2=l2, level3=l3,
                                        planned_amount=planned, sort_order=order, **kw)
            db.add(row); db.flush()
            return row

        c1 = cat("식비", "본행사 식사", "자율배식", 0, 1, unit_price=8000, headcount=150, times=5)
        c1.planned_amount = budget_domain.planned_amount_of(8000, 150, 5, 0)
        c2 = cat("식비", "본행사 식사", "간식", 200_000, 2)
        c3 = cat("식비", "야식", "야식", 300_000, 3)
        c4 = cat("홍보", "포스터", "인쇄비", 100_000, 4)
        e = models.ExpenseEntry(retreat_id=r.id, budget_category_id=c4.id, level1="홍보",
                                level2="포스터", expense_date=TODAY, amount=180_000,
                                payer_name="다른사람", paid=True, subsidy_amount=180_000,
                                department_id=d.id)
        db.add(e); db.commit()
        ids = {"retreat": r.id, "cats": [c1.id, c2.id, c3.id, c4.id]}
    admin_client.get(f"/budget?retreat_id={ids['retreat']}")
    return ids


def _수입(retreat_id, amount):
    with app_session() as db:
        db.add(models.IncomeItem(retreat_id=retreat_id, name="수련회비", amount=amount, sort_order=1))
        db.commit()


@pytest.fixture
def 시간표(admin_client):
    """`test_staff_sheet.sheet_data` 를 줄인 것 + **선발대 09:30 시작 프로그램** —
    0-d 의 「09:00 칸이 둘로 갈라지는」 모양이 이 재료에 있다."""
    with app_session() as db:
        r = models.Retreat(name="2026 여름수련회 Belong", start_date=OPEN, end_date=CLOSE)
        db.add(r); db.flush()
        for order, (key, name, color) in enumerate(
            [("chongmuM", "1 총무M", "#2F4858"), ("hebron", "5 헤브론", "#4A8A5C"),
             ("koram", "6 코람데오", "#B44B42")]
        ):
            db.add(models.Department(retreat_id=r.id, key=key, name=name,
                                     color_tag=color, sort_order=order))

        def program(day, time, name, *, audience="staff", track="main", parallel=False,
                    end=None, items=(), host=None, place=None):
            p = models.Program(retreat_id=r.id, day=day, start_time=time, name=name,
                               host=host, place=place, audience=audience, track=track,
                               parallel=parallel, end_time=end, sort_order=0)
            db.add(p); db.flush()
            for order, (part, who, text) in enumerate(items):
                db.add(models.ProgramItem(program_id=p.id, phase="pre", part_key=part,
                                          assignee_name=who, text=text, sort_order=order,
                                          scope="team" if part in ("헤브론", "코람데오") else "person"))
            return p

        program("선발대", "09:30", "본당집합")                       # 반 칸에서 시작 (0-d)
        program("선발대", "10:00", "짐정리", place="본당",
                items=[("코람데오", "재하", "짐 나르기")])
        program("선발대", "14:00", "음향 및 무대설치", host="헤브론",
                items=[("헤브론", "헤브론", "음향 및 무대설치 (3.5h)")])
        program("1일차", "10:00", "봉사자 예배", audience="all", host="목사",
                items=[("헤브론", "건우", "음향 장비 확인")])
        program("1일차", "19:00", "집회1", audience="all")
        program("1일차", "19:00", "새친구", audience="all", parallel=True, end="23:00")
        program("1일차", "14:00", "리허설")
        program("2일차", "00:00", "광고", audience="all")
        program("2일차", "00:00", "야식", audience="all")
        program("2일차", "08:00", "아침식사", audience="all", host="총무팀")
        program("폐회", "10:30", "폐회 예배 · 성찬", audience="all", host="목사", place="강당")
        db.commit()
        ids = {"retreat": r.id}
    admin_client.get(f"/live/staff?retreat_id={ids['retreat']}")
    return ids


# ── 1. 예산 표 마크업 (5-a) ──────────────────────────────────────────

def _행들(html):
    body = html.split("<tbody>", 1)[1].split("</tbody>", 1)[0]
    return re.findall(r"<tr[^>]*>.*?</tr>", body, re.S)


def test26_a01_구분_제목_줄에_소계가_있고_별도_소계_행과_항목_행이_없다(admin_client, 예산):
    page = admin_client.get("/budget").text
    rows = _행들(page)
    assert not [r for r in rows if 'class="subtotal"' in r], "별도 소계 행이 남아 있다"
    assert not [r for r in rows if 'class="l2row"' in r], "항목 별도 줄이 남아 있다"
    heads = [r for r in rows if "l1row" in r]
    assert len(heads) == 2
    식비 = next(r for r in heads if "식비" in r)
    # 제목 줄 = 이름 · 세부항목 수 · 소계(예산·결산·차액) — 값은 summary 의 것
    assert "세부항목 3" in 식비 and "6,500,000" in 식비 and '<td class="r mono">0</td>' in 식비
    assert 'data-fold=' in 식비 and 'aria-expanded="true"' in 식비, "접기 단추가 없다"
    홍보 = next(r for r in heads if "홍보" in r)
    assert "세부항목 1" in 홍보 and "100,000" in 홍보 and "180,000" in 홍보 and "-80,000" in 홍보


def test26_a02_항목_라벨은_첫_세부항목에만_붙고_단가_명수_횟수_칸이_전_행에_있다(admin_client, 예산):
    page = admin_client.get("/budget").text
    data = [r for r in _행들(page) if 'class="l1row"' not in r and 'class="grand"' not in r]
    assert len(data) == 4
    labels = [re.search(r'class="itemlbl"[^>]*>([^<]*)<', r) for r in data]
    assert [m.group(1).strip() if m else None for m in labels] == ["본행사 식사", None, "야식", "포스터"]
    for r in data:
        assert r.count("<td") >= 11, "단가·명수·횟수 칸이 빠진 행이 있다"
        assert 'data-group="' in r, "접기가 가리킬 묶음이 없다"


def test26_a03_빈_단가는_대시이고_있으면_값이다(admin_client, 예산):
    page = admin_client.get("/budget").text
    data = [r for r in _행들(page) if 'class="l1row"' not in r and 'class="grand"' not in r]
    first, second = data[0], data[1]
    assert "8,000" in first and "150" in first and 'class="r mono dim">–<' not in first
    assert second.count('class="r mono dim">–<') == 3, "값이 없는 셋이 「–」 로 안 찍혔다"


# ── 2. 잔액 카드 (5-b) ───────────────────────────────────────────────

def _잔액카드(html):
    return re.search(r'<div class="kpi"><div class="l">잔액</div>.*?</div>\s*</div>', html, re.S).group(0)


def test26_b01_수입이_없으면_음수_대신_수입_미입력이고_빨강이_없다(admin_client, 예산):
    page = admin_client.get("/budget").text
    card = _잔액카드(page)
    assert "수입 미입력" in card and "수입을 넣으면 계산됩니다" in card
    assert "-" not in card.split("잔액")[1].split("수입")[0] and "var(--now)" not in card
    assert "overrun" not in page.split("잔액 (총 수입")[1][:300], "수입 표의 잔액 줄이 빨갛다"


def test26_b02_수입이_있으면_계산하고_음수면_빨갛다(admin_client, 예산):
    _수입(예산["retreat"], 1_000_000)          # 지출예산 6,600,000 보다 작다 → 음수
    page = admin_client.get("/budget").text
    card = _잔액카드(page)
    assert "수입 미입력" not in card and "-5,600,000" in card and "var(--now)" in card
    assert "overrun" in page.split("잔액 (총 수입")[1][:300]


def test26_b03_수입이_넉넉하면_빨강이_없다(admin_client, 예산):
    _수입(예산["retreat"], 9_000_000)
    page = admin_client.get("/budget").text
    card = _잔액카드(page)
    assert "2,400,000" in card and "var(--now)" not in card and "수입 미입력" not in card


# ── 3. 엑셀 둘은 판 전후로 같다 (5-c) ──────────────────────────────────

def _시트요약(ws):
    return {
        "max_row": ws.max_row, "max_col": ws.max_column,
        "merged": sorted(str(m) for m in ws.merged_cells.ranges),
        "cells": [[c.coordinate, c.value] for row in ws.iter_rows() for c in row if c.value not in (None, "")],
    }


def _골든(name, got):
    path = GOLDEN / f"{name}.json"
    if os.environ.get("GOLDEN_WRITE"):
        GOLDEN.mkdir(exist_ok=True)
        path.write_text(json.dumps(got, ensure_ascii=False, indent=1), encoding="utf-8")
    assert path.exists(), f"골든이 없다: {path} — GOLDEN_WRITE=1 로 한 번 떠 둔다"
    want = json.loads(path.read_text(encoding="utf-8"))
    assert got == want, f"{name} 엑셀이 골든과 다르다"


def test26_c01_예산_대비_집행_시트가_골든과_같다(admin_client, 예산):
    raw = admin_client.get(f"/export/expenses.xlsx?retreat_id={예산['retreat']}").content
    ws = load_workbook(io.BytesIO(raw))["예산 대비 집행"]
    got = _시트요약(ws)
    assert got["cells"], "시트가 비었다 — 검사가 아무것도 안 보고 있다"
    _골든("예산대비집행", got)


def test26_c02_봉사팀_시트가_골든과_같다(admin_client, 시간표):
    raw = admin_client.get("/live/staff.xlsx").content
    ws = load_workbook(io.BytesIO(raw))["봉사자 시간표"]
    got = _시트요약(ws)
    assert got["merged"] and got["cells"], "시트가 비었다"
    _골든("봉사팀", got)


def test26_c03_골든_검사가_다른_것을_실제로_잡는다(tmp_path, monkeypatch):
    """③ — 같은 요약을 한 글자만 바꿔 넣으면 빨개진다."""
    monkeypatch.setattr("tests.test_stage26.GOLDEN", tmp_path)
    got = {"max_row": 1, "max_col": 1, "merged": [], "cells": [["A1", "x"]]}
    monkeypatch.setenv("GOLDEN_WRITE", "1")
    _골든("가짜", got)
    monkeypatch.delenv("GOLDEN_WRITE")
    with pytest.raises(AssertionError):
        _골든("가짜", {**got, "cells": [["A1", "y"]]})


# ── 4-b. /draft 기본 칸 (5-d) ───────────────────────────────────────

@pytest.fixture
def 초안(admin_client):
    with app_session() as db:
        r = models.Retreat(name="초안 회차", start_date=OPEN, end_date=CLOSE)
        db.add(r); db.flush()
        ids = {"retreat": r.id}
        for i, (key, name) in enumerate((("hebron", "5 헤브론"), ("sketch", "4 스케치"))):
            d = models.Department(retreat_id=r.id, key=key, name=name, color_tag="#888", sort_order=i)
            db.add(d); db.flush(); ids[key] = d.id
        from app.domain import drafts as draft_domain
        draft_domain.open_draft(db, name="다음 회차", open_date=OPEN, close_date=CLOSE,
                                meal_subsidy=8000, department_keys=["hebron", "sketch"])
    return ids


def test26_d01_팀원으로만_둘이면_고르는_화면이_뜬다(client, 초안):
    make_user("팀원 둘", "01066660081", "general",
              departments=[(초안["hebron"], perm.MEMBER), (초안["sketch"], perm.MEMBER)])
    login_as(client, "01066660081")
    page = client.get("/draft").text
    assert "어느 칸을 채울지" in page
    assert 'href="/draft?department=hebron"' in page and 'href="/draft?department=sketch"' in page


def test26_d02_팀원으로_하나면_바로_그_부서다(client, 초안):
    make_user("팀원 하나", "01066660082", "general", departments=[(초안["sketch"], perm.MEMBER)])
    login_as(client, "01066660082")
    page = client.get("/draft").text
    assert 'href="/draft?department=hebron"' not in page
    assert "4 스케치" in page and "고르" not in page.split("<main", 1)[-1][:400]


def test26_d03_리더인_부서가_하나면_팀원이_더_있어도_그_부서다(client, 초안):
    make_user("리더 하나", "01066660083", "general",
              departments=[(초안["hebron"], perm.LEAD), (초안["sketch"], perm.MEMBER)])
    login_as(client, "01066660083")
    page = client.get("/draft").text
    assert 'href="/draft?department=sketch"' not in page and "5 헤브론" in page


# ── 3-b·4-a. CSS 와 마크업이 정한 대로인가 (계산값은 브라우저에서 — 5-e) ──

CSS = (ROOT / "app/static/css/retreat.css").read_text(encoding="utf-8")
PASTEL = ("FFF2CC", "D9EAD3", "CFE2F3", "FCE5CD", "E6D7F2", "D9D9D9")


def _ssheet_css():
    i = CSS.index(".sheetwrap{")
    j = CSS.index("/* ── 프로그램 만들기·고치기 창")
    return CSS[i:j]


def test26_e01_봉사팀_보기_CSS_에_엑셀_흉내가_없다():
    block = _ssheet_css()
    for hexcode in PASTEL:
        assert hexcode.lower() not in block.lower(), f"파스텔 {hexcode} 가 남았다"
    assert "--mono" not in block, "고정폭 글꼴이 남았다"
    assert "rgba(0,0,0,.55)" not in block and "double" not in block, "굵은 회색 격자가 남았다"
    assert "--line" in block or "--rule" in block


def test26_e02_화면은_구조의_색을_칠하지_않고_띠와_부서로_가른다():
    tpl = (ROOT / "app/templates/staff_sheet.html").read_text(encoding="utf-8")
    assert 'style="background:#{{ cell.fill }}"' not in tpl, "파스텔 채움을 아직 칠한다"
    assert "band-" in tpl and "--dc:" in tpl, "띠·부서 표시가 없다"


def test26_e03_사이드바_소속_줄은_두_줄까지_꺾이고_전체가_title_에_있다():
    tpl = (ROOT / "app/templates/retreat_base.html").read_text(encoding="utf-8")
    assert 'class="depts" title="{{ side_dept_name' in tpl
    assert "-webkit-line-clamp:2" in CSS.replace(" ", "") and ".sidefoot .who .depts" in CSS


def test26_e04_봉사팀_보기의_글자_하한과_그_예외():
    """칸 안 글자는 14px 하한(4-0)이다. **예외는 넘치는 칸의 `dense` 하나** —
    칸을 늘리면 시각 눈금이 어긋나므로 글자를 줄이는 자리(5-8). 크기는 CSS 가
    아니라 `staff_sheet` 가 정하므로 그 값을 잰다."""
    from app.domain import staff_sheet as ss
    assert ss.BODY_FONT_PX >= 14 and ss.HEAD_FONT_PX >= 14
    assert ss.DENSE_FONT_PX == 13, "예외의 값이 바뀌었다 — 이 시험의 이유도 다시 본다"
    assert ss.DENSE_FONT_PX < 14 and ss.DENSE_FONT_PX >= 12
    # 여백은 예전(3px 5px)의 두 배 이상, 줄 간격 1.35 이상 (3-c) — 줄 간격은 구조의
    # LINE_SCALE 하나에서 나온다 (e06 이 CSS 와 견준다)
    block = _ssheet_css()
    assert "padding:6px 10px" in block and ss.LINE_SCALE >= 1.35


def test26_e05_파트_색_짝은_전체일정_봉사자_공통색과_겹치지_않는다():
    """화면이 `fill` 에서 파트를 되돌리므로(라우터 `part_of_fill`), 파트 색 상수가
    전체일정·봉사자 공통색과 겹치면 그 칸에 엉뚱한 부서 선이 붙는다 (검토 9번)."""
    from app.domain import staff_sheet as ss
    keys = {ss.fill_for_part(p) for p in ss.team_parts()}
    assert ss.FILL_ALL not in keys and ss.FILL_STAFF not in keys
    assert len(keys) == len(ss.team_parts()), "파트끼리 색이 겹친다"


def test26_e06_화면의_여백_줄간격이_넘침_판정의_전제와_같다():
    """`dense_of` 가 전제하는 여백·줄 간격과 CSS `.cx` 가 같은 값이다 — 갈리면 안
    넘친다고 판정한 칸이 화면에서 조용히 잘린다 (검토 1번)."""
    from app.domain import staff_sheet as ss
    block = _ssheet_css()
    m = re.search(r"\.cx\{[^}]*padding:(\d+)px (\d+)px[^}]*line-height:var\(--lh\)", block)
    assert m, ".cx 의 여백·줄 간격이 정한 모양이 아니다"
    top, side = int(m.group(1)), int(m.group(2))
    assert side * 2 + 1 == ss.CELL_PAD_PX, "좌우 여백+테두리가 CELL_PAD_PX 와 다르다"
    assert top * 2 == ss.CELL_VPAD_PX, "위아래 여백이 CELL_VPAD_PX 와 다르다"
    assert ss.LINE_SCALE >= 1.35 and "line-height:1." not in block.split(".cx{")[1].split("}")[0]
    assert "width:max-content;min-width:100%" in block, "표를 창에 맞춰 줄이면 넘침 판정이 어긋난다"
