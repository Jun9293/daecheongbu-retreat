"""재정 차례 4 — 계좌 셋 · 영수증 잇기 표 · 원본 영수증 번호 (2026-09-14).

- 계좌는 은행 · 번호 · 예금주 셋으로 받고 보인다. 옛 한 칸은 남고 값도 안 바뀐다
- 영수증 한 장에 지출 여럿이 걸린다. 금액은 지출마다, 자동 번호는 영수증마다 하나
- 떼는 것은 잇기를 끊는 것이다 — 영수증 행은 남는다
- 지금 있는 영수증은 `scripts/영수증잇기옮기기.py` 가 처음 붙은 지출에 잇는다
- 원본 번호는 있을 때만 화면에 붙고 엑셀 칸에 들어간다
- 예산금액이 0 인 항목에서 집행률이 터지지 않는다

계좌·이름은 지어낸 값이다. 수는 박지 않는다(11-3) — 전후를 견준다.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import io
import pathlib

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select, text

from app import models
from app.domain import budget as B
from tests.conftest import app_session, login_as, make_user

TODAY = dt.date.today()
ROOT = pathlib.Path(__file__).resolve().parent.parent
은행, 번호, 예금주 = "지어낸은행", "999-0000-1111", "가명예금주"


@pytest.fixture
def 회차(admin_client):
    with app_session() as db:
        r = models.Retreat(name="차례4 회차", meal_subsidy_per_person=8000,
                           start_date=TODAY + dt.timedelta(days=30),
                           end_date=TODAY + dt.timedelta(days=32))
        db.add(r)
        db.flush()
        dept = models.Department(retreat_id=r.id, key="chongmuM", name="1 총무M", sort_order=0)
        db.add(dept)
        db.flush()
        db.commit()
        ids = {"retreat": r.id, "dept": dept.id}
    admin_client.get(f"/expenses?retreat_id={ids['retreat']}")
    return ids


def _등록(client, ids, amount, **extra):
    data = {"amount": str(amount), "department_id": str(ids["dept"]), "payer_name": "가명지출자"}
    data.update(extra)
    resp = client.post(f"/expenses/create?retreat_id={ids['retreat']}", data=data,
                       follow_redirects=True)
    assert resp.status_code == 200, resp.text
    with app_session() as db:
        return db.scalar(select(models.ExpenseEntry.id).order_by(models.ExpenseEntry.id.desc()))


# ── 가) 계좌 셋 — 셋으로 저장되고 셋으로 보이고, 못 보는 계정에는 안 보인다 ──


def test46_a01_계좌_셋이_셋으로_저장되고_보인다(client, admin_client, 회차):
    eid = _등록(admin_client, 회차, 10_000, payer_bank=은행,
               payer_account_number=번호, payer_account_holder=예금주)
    with app_session() as db:
        e = db.get(models.ExpenseEntry, eid)
        assert (e.payer_bank, e.payer_account_number, e.payer_account_holder) == (은행, 번호, 예금주)
    page = admin_client.get(f"/expenses?retreat_id={회차['retreat']}").text
    for 칸, 값 in (("acct-bank", 은행), ("acct-no", 번호), ("acct-holder", 예금주)):
        assert f'class="{칸}">' in page and 값 in page, 칸
    # 엑셀도 셋 — 볼 수 있는 사람의 파일
    wb = load_workbook(io.BytesIO(admin_client.get(
        f"/export/expenses.xlsx?retreat_id={회차['retreat']}").content))
    머리 = [c.value for c in wb["지출 상세내역"][1]]
    줄 = [c.value for c in wb["지출 상세내역"][2]]
    assert [줄[머리.index(h)] for h in ("은행", "계좌번호", "예금주")] == [은행, 번호, 예금주]

    # 같은 판에서 — 못 보는 계정(부서 리더)에는 셋 다 없다
    make_user("가명리더", "01077770001", "dept_lead", dept=회차["dept"])
    login_as(client, "01077770001")
    남 = client.get(f"/expenses?retreat_id={회차['retreat']}").text
    assert "가명지출자" in 남, "목록 자체는 보여야 한다"
    for 값 in (은행, 번호, 예금주):
        assert 값 not in 남
    assert "enote acct" not in 남, "빈 계좌 줄을 남기지 않는다"


def test46_a02_내_정보의_계좌_셋이_등록_폼에_채워진다(admin_client, 회차):
    admin_client.post("/me/update", data={"name": "총무 김간사", "bank_name": 은행,
                                           "account_number": 번호, "account_holder": 예금주})
    page = admin_client.get(f"/expenses?retreat_id={회차['retreat']}").text
    for 이름, 값 in (("payer_bank", 은행), ("payer_account_number", 번호),
                   ("payer_account_holder", 예금주)):
        assert f'name="{이름}" value="{값}"' in page, 이름


# ── 나) 옛 칸은 그대로 있고 값이 안 바뀌었다 ──


def test46_b01_옛_계좌_칸은_남고_값이_안_바뀐다(admin_client, 회차):
    eid = _등록(admin_client, 회차, 5_000)
    with app_session() as db:
        db.execute(text("UPDATE expense_entries SET payer_account='옛값그대로' WHERE id=:i"), {"i": eid})
        uid = db.scalar(select(models.User.id))
        db.execute(text("UPDATE users SET bank_account='옛값그대로' WHERE id=:i"), {"i": uid})
        db.commit()
        옛지출 = db.execute(text("SELECT payer_account FROM expense_entries WHERE id=:i"), {"i": eid}).scalar_one()
        옛사람 = db.execute(text("SELECT bank_account FROM users WHERE id=:i"), {"i": uid}).scalar_one()
    # 새 칸으로 고치고 등록하고 내려받아도
    admin_client.post("/me/update", data={"name": "총무 김간사", "bank_name": 은행})
    _등록(admin_client, 회차, 6_000, payer_bank=은행)
    admin_client.get(f"/export/expenses.xlsx?retreat_id={회차['retreat']}")
    with app_session() as db:
        assert db.execute(text("SELECT payer_account FROM expense_entries WHERE id=:i"), {"i": eid}).scalar_one() == 옛지출
        assert db.execute(text("SELECT bank_account FROM users WHERE id=:i"), {"i": uid}).scalar_one() == 옛사람
        # 옛 칸은 구조로 봉인됐다 — 옛 이름으로 읽으면 AttributeError
        with pytest.raises(AttributeError):
            db.get(models.ExpenseEntry, eid).payer_account
        with pytest.raises(AttributeError):
            db.get(models.User, uid).bank_account
    # 옛 이름을 읽는 곳이 없다 — **템플릿까지** 본다. Jinja 는 없는 속성을 빈 값으로 그려
    # 조용히 새므로 가장 먼저 봐야 할 자리다. 새 이름(payer_account_number 등)은 뒤에 `_` 가 붙어 걸리지 않는다
    import re
    옛 = re.compile("(?<![A-Za-z0-9_])(payer_account|bank_account)(?![A-Za-z0-9_])")
    볼것 = [p for p in (ROOT / "app").rglob("*") if p.suffix in (".py", ".html", ".js")]
    볼것 += list((ROOT / "scripts").rglob("*.py")) + [ROOT / "seed.py"]
    assert any(p.suffix == ".html" for p in 볼것), "템플릿을 안 보고 있다"
    for path in 볼것:
        src = path.read_text(encoding="utf-8")
        if path.name in ("models.py", "db.py"):          # 봉인한 자리 · 칸을 붙이는 자리
            continue
        assert not 옛.search(src), path
    assert 옛.search("x.payer_account,") and not 옛.search("x.payer_account_number"), "무늬가 새 이름을 거르거나 옛 이름을 놓친다"


# ── 다) 영수증 한 장에 지출 둘 — 금액은 따로, 자동 번호는 하나 ──


def _영수증번호(db, eid):
    return [r.number for r in db.get(models.ExpenseEntry, eid).receipts]


def test46_c01_영수증_하나에_지출_둘(admin_client, 회차):
    rid = 회차["retreat"]
    e1 = _등록(admin_client, 회차, 30_000, receipt_memo="결산 파일에 별첨")
    e2 = _등록(admin_client, 회차, 12_000)
    with app_session() as db:
        [번] = _영수증번호(db, e1)
        다음 = B.next_receipt_number(db, db.get(models.Retreat, rid))
    resp = admin_client.post(f"/expenses/{e2}/receipts/link?retreat_id={rid}",
                             data={"number": str(번)}, follow_redirects=True)
    assert resp.status_code == 200
    with app_session() as db:
        assert _영수증번호(db, e1) == [번] and _영수증번호(db, e2) == [번]
        assert db.scalar(select(func.count()).select_from(models.ExpenseReceipt)) == 1, "번호가 둘로 늘었다"
        assert B.next_receipt_number(db, db.get(models.Retreat, rid)) == 다음, "잇기가 새 번호를 먹었다"
        # 금액은 지출마다 그대로 — 영수증에 총액이 없다
        assert db.get(models.ExpenseEntry, e1).amount == 30_000
        assert db.get(models.ExpenseEntry, e2).amount == 12_000
        assert not hasattr(models.ExpenseReceipt, "amount")
        assert len(db.get(models.ExpenseEntry, e1).receipts[0].expenses) == 2
    page = admin_client.get(f"/expenses?retreat_id={rid}").text
    assert "지출 2건" in page
    # 없는 번호는 거절한다
    assert admin_client.post(f"/expenses/{e2}/receipts/link?retreat_id={rid}",
                             data={"number": "9999"}).status_code == 404


def test46_c02_남의_부서_지출에는_못_잇는다(client, admin_client, 회차):
    rid = 회차["retreat"]
    e1 = _등록(admin_client, 회차, 1_000, receipt_memo="별첨")
    with app_session() as db:
        other = models.Department(retreat_id=rid, key="hebron", name="5 헤브론", sort_order=1)
        db.add(other)
        db.commit()
        oid = other.id
    make_user("가명헤브론", "01077770002", "dept_lead", dept=oid)
    login_as(client, "01077770002")
    with app_session() as db:
        [receipt] = db.get(models.ExpenseEntry, e1).receipts
        번, rcpt_id = receipt.number, receipt.id
    assert client.post(f"/expenses/{e1}/receipts/link?retreat_id={rid}",
                       data={"number": str(번)}).status_code == 403
    # 떼기도 같은 문 — 막히는 쪽을 따로 잰다
    assert client.post(f"/expenses/{e1}/receipts/{rcpt_id}/detach?retreat_id={rid}").status_code == 403
    with app_session() as db:
        assert _영수증번호(db, e1) == [번], "남의 부서가 뗐다"


def test46_c03_다른_회차의_영수증에는_못_잇는다(admin_client, 회차):
    """번호는 회차 안에서만 매겨진다 — 같은 번호가 다른 회차에도 있다(2장 · 인계 26)."""
    rid = 회차["retreat"]
    with app_session() as db:
        b = models.Retreat(name="다른 회차", meal_subsidy_per_person=8000,
                           start_date=TODAY + dt.timedelta(days=90), end_date=TODAY + dt.timedelta(days=92))
        db.add(b)
        db.flush()
        other = models.ExpenseEntry(retreat_id=b.id, amount=1, subsidy_amount=1, expense_date=TODAY)
        db.add(other)
        other.attach_receipt(models.ExpenseReceipt(number=1, memo="다른 회차"))
        db.commit()
        bid = b.id
    e1 = _등록(admin_client, 회차, 1_000)
    assert admin_client.post(f"/expenses/{e1}/receipts/link?retreat_id={rid}",
                             data={"number": "1"}).status_code == 404
    # 다른 회차 영수증의 id 로 떼기 — 이 지출에 이어진 줄이 없으니 404
    with app_session() as db:
        other_rcpt = db.scalar(select(models.ExpenseReceipt.id).where(models.ExpenseReceipt.memo == "다른 회차"))
    assert admin_client.post(f"/expenses/{e1}/receipts/{other_rcpt}/detach?retreat_id={rid}").status_code == 404
    with app_session() as db:
        assert db.get(models.ExpenseEntry, e1).receipts == []
        assert B.next_receipt_number(db, db.get(models.Retreat, rid)) == 1
        assert B.next_receipt_number(db, db.get(models.Retreat, bid)) == 2


# ── 라) 떼면 잇기가 끊기고 영수증 행은 남는다 ──


def test46_d01_떼면_영수증_행이_남는다(admin_client, 회차):
    rid = 회차["retreat"]
    e1 = _등록(admin_client, 회차, 1_000, receipt_memo="잘못 붙인 것")
    with app_session() as db:
        receipt = db.get(models.ExpenseEntry, e1).receipts[0]
        rcpt_id, 번 = receipt.id, receipt.number
    resp = admin_client.post(f"/expenses/{e1}/receipts/{rcpt_id}/detach?retreat_id={rid}",
                             follow_redirects=True)
    assert resp.status_code == 200
    with app_session() as db:
        assert db.get(models.ExpenseReceipt, rcpt_id) is not None, "영수증 행이 지워졌다"
        assert db.get(models.ExpenseEntry, e1).receipts == []
        links = db.scalars(select(models.ExpenseReceiptLink).where(
            models.ExpenseReceiptLink.receipt_id == rcpt_id)).all()
        assert len(links) == 1 and links[0].detached_at is not None, "끊긴 줄도 남아야 한다"
        assert B.unlinked_receipt_count(db, db.get(models.Retreat, rid)) == 0, "뗀 것은 옮길 것이 아니다"
        assert B.next_receipt_number(db, db.get(models.Retreat, rid)) == 번 + 1, "뗀 번호를 다시 쓰면 안 된다"
    with app_session() as db:
        assert e1 in [e.id for e in B.no_receipt_entries(db, db.get(models.Retreat, rid))]
    # 번호로 다시 이을 수 있다
    admin_client.post(f"/expenses/{e1}/receipts/link?retreat_id={rid}", data={"number": str(번)})
    with app_session() as db:
        assert _영수증번호(db, e1) == [번]
    # 이어져 있지 않은 것을 떼려 하면 404
    e2 = _등록(admin_client, 회차, 2_000)
    assert admin_client.post(f"/expenses/{e2}/receipts/{rcpt_id}/detach?retreat_id={rid}").status_code == 404


# ── 마) 지금 있는 영수증이 전환 뒤에도 같은 지출에 걸린다 ──


def _스크립트():
    path = ROOT / "scripts" / "영수증잇기옮기기.py"
    spec = importlib.util.spec_from_file_location("영수증잇기옮기기", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test46_e01_옛_영수증을_잇는_전환(admin_client, 회차):
    rid = 회차["retreat"]
    ids = [_등록(admin_client, 회차, 1_000 * (i + 1), receipt_memo=f"메모{i}") for i in range(3)]
    ids.append(_등록(admin_client, 회차, 500))                       # 영수증 없는 지출
    with app_session() as db:
        # 전환 전의 모양으로 되돌린다 — 잇기 표가 생기기 전에는 expense_id 만 있었다
        전 = dict(db.execute(text("SELECT id, expense_id FROM expense_receipts")).all())
        assert 전, "견줄 영수증이 없으면 이 시험은 아무것도 안 본다"
        db.execute(text("DELETE FROM expense_receipt_links"))
        db.commit()
        assert all(db.get(models.ExpenseEntry, i).receipts == [] for i in ids)
        assert B.unlinked_receipt_count(db, db.get(models.Retreat, rid)) == len(전)
    assert "아직 지출에 이어지지 않은 영수증" in admin_client.get(f"/expenses?retreat_id={rid}").text
    # 창이 열린 동안 옛 영수증을 다른 지출에 번호로 이으면 스크립트가 건너뛰어 원래 지출을 잃는다 — 막는다
    with app_session() as db:
        옛번호 = db.get(models.ExpenseReceipt, next(iter(전))).number
    assert admin_client.post(f"/expenses/{ids[-1]}/receipts/link?retreat_id={rid}",
                             data={"number": str(옛번호)}).status_code == 409
    with app_session() as db:
        assert db.scalar(select(func.count()).select_from(models.ExpenseReceiptLink)) == 0

    mod = _스크립트()
    with app_session() as db:
        mod.옮기기(db, False)                                        # 미리보기는 안 바꾼다
        assert mod.센다(db)["이을것"] == len(전)
        mod.옮기기(db, True, 사본=False)
    with app_session() as db:
        후 = {r.id: [e.id for e in r.expenses] for r in db.scalars(select(models.ExpenseReceipt))}
        assert 후 == {k: [v] for k, v in 전.items()}
        assert mod.센다(db)["이을것"] == 0
        before = mod.센다(db)
        mod.옮기기(db, True, 사본=False)                              # 두 번째는 이을 것이 없다
        assert mod.센다(db) == before
    assert "아직 지출에 이어지지 않은 영수증" not in admin_client.get(f"/expenses?retreat_id={rid}").text
    # 한 벌 — 창이 닫힌 뒤에는 같은 옛 영수증을 번호로 이을 수 있다
    assert admin_client.post(f"/expenses/{ids[-1]}/receipts/link?retreat_id={rid}",
                             data={"number": str(옛번호)}, follow_redirects=True).status_code == 200
    with app_session() as db:
        assert 옛번호 in _영수증번호(db, ids[-1])


def test46_e02_뗀_영수증은_전환이_다시_붙이지_않는다(admin_client, 회차):
    rid = 회차["retreat"]
    e1 = _등록(admin_client, 회차, 1_000, receipt_memo="뗄 것")
    with app_session() as db:
        rcpt_id = db.get(models.ExpenseEntry, e1).receipts[0].id
    admin_client.post(f"/expenses/{e1}/receipts/{rcpt_id}/detach?retreat_id={rid}")
    mod = _스크립트()
    with app_session() as db:
        mod.옮기기(db, True, 사본=False)
        assert db.get(models.ExpenseEntry, e1).receipts == []


# ── 바) 원본 번호 — 화면과 엑셀에 보이고, 비면 그 자리가 안 그려진다 ──


def test46_f01_원본_번호가_보이고_비면_안_그려진다(admin_client, 회차):
    rid = 회차["retreat"]
    e1 = _등록(admin_client, 회차, 1_000, receipt_memo="별첨", receipt_original_no="시트17")
    e2 = _등록(admin_client, 회차, 2_000, receipt_memo="별첨")      # 원본 번호 없음
    with app_session() as db:
        r1 = db.get(models.ExpenseEntry, e1).receipts[0]
        r2 = db.get(models.ExpenseEntry, e2).receipts[0]
        assert r1.original_no == "시트17" and r2.original_no is None
        n1, n2 = r1.number, r2.number
    page = admin_client.get(f"/expenses?retreat_id={rid}").text
    assert f"영수증 {n1} · 원본 시트17" in page
    assert f"영수증 {n2} · 원본" not in page, "빈 원본 번호 자리가 그려졌다"
    assert f"영수증 {n2} · 별첨" in page

    wb = load_workbook(io.BytesIO(admin_client.get(f"/export/expenses.xlsx?retreat_id={rid}").content))
    for 시트 in ("지출 상세내역",):
        ws = wb[시트]
        머리 = [c.value for c in ws[1]]
        칸 = 머리.index("원본 영수증번호")
        값 = {row[머리.index("영수증번호")].value: row[칸].value for row in ws.iter_rows(min_row=2)}
        assert 값[str(n1)] == "시트17" and 값[str(n2)] is None


def test46_f02_원본_번호는_회차_안에서_겹칠_수_있다(admin_client, 회차):
    _등록(admin_client, 회차, 1_000, receipt_memo="a", receipt_original_no="3")
    _등록(admin_client, 회차, 1_000, receipt_memo="b", receipt_original_no="3")
    with app_session() as db:
        rows = db.scalars(select(models.ExpenseReceipt)).all()
        assert [r.original_no for r in rows] == ["3", "3"]
        assert len({r.number for r in rows}) == 2, "자동 번호는 겹치면 안 된다"


# ── 사) 예산금액이 0 인 항목 — 0 으로 나누는 자리를 만들어 잰다 ──


def test46_g01_예산금액_0_인_항목에서_집행률이_안_터진다(admin_client, 회차):
    rid = 회차["retreat"]
    with app_session() as db:
        zero = models.BudgetCategory(retreat_id=rid, level1="교통", level2="버스대여",
                                     planned_amount=0, sort_order=1)
        db.add(zero)
        db.commit()
        zid = zero.id
        s = B.build_budget_summary(db, retreat=db.get(models.Retreat, rid))
        assert s.total_planned == 0                                   # 분모가 0 인 자리
        assert s.progress_pct == 0.0 and s.categories[0].progress_pct == 0.0
        assert s.categories[0].ratio_pct == 0.0
    # 걸린 지출이 있어도(분자 > 0, 분모 0)
    _등록(admin_client, 회차, 50_000, budget_category_id=str(zid))
    with app_session() as db:
        s = B.build_budget_summary(db, retreat=db.get(models.Retreat, rid))
        row = s.categories[0]
        assert row.spent > 0 and row.planned == 0
        assert row.progress_pct == 0.0 and s.progress_pct == 0.0
        assert row.is_over_budget, "0 예산에 쓴 돈은 초과로 칠해져야 한다"
    for 화면, 글 in (("/", "%"), ("/budget", "버스대여"), ("/expenses", "버스대여"),
                   (f"/settings/retreats/{rid}", "%")):
        resp = admin_client.get(f"{화면}?retreat_id={rid}")
        assert resp.status_code == 200, (화면, resp.status_code)
        assert 글 in resp.text, 화면
    assert admin_client.get(f"/export/expenses.xlsx?retreat_id={rid}").status_code == 200
