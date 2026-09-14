"""재정 차례 6 — 시트 들여오기 (2026-09-14).

**실제 시트는 열지 않는다.** 시트와 같은 꼴(탭 이름 · 머리글 · 병합 · 표지 줄 · 초록 채움)의
합성 xlsx 를 만들어 `scripts/재정들여오기.py` 를 돌린다. 이름·계좌·금액은 지어낸 값이다.
수는 박지 않는다(11-3) — 스크립트가 세운 계획의 수와 DB 에 들어간 수를 견준다.
"""

from __future__ import annotations

import ast
import datetime as dt
import importlib.util
import json
import pathlib
import sys

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import func, select

from app import config, models
from tests.conftest import app_session
from tests.test_stage47 import _가르는곳들

ROOT = pathlib.Path(__file__).resolve().parent.parent
상한 = 8000
지어낸계좌 = "지어낸은행 999-0000-1111 "
지어낸이름 = "가명지출자"


def _스크립트():
    이름 = "재정들여오기_시험"
    spec = importlib.util.spec_from_file_location(이름, ROOT / "scripts" / "재정들여오기.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[이름] = mod          # dataclass 가 모듈을 찾는다
    spec.loader.exec_module(mod)
    return mod


def _시트(경로: pathlib.Path, *, 둘째판: bool = False) -> None:
    """시트와 같은 꼴의 합성 파일. 예산 탭은 한 칸 밀려 있고(B 부터), 항목은 C:D 병합,
    지출 탭은 항목 B:C 병합 · 세부항목-1 병합 · 영수증 칸 초록 채움까지 흉내 낸다."""
    wb = Workbook()
    b = wb.active
    b.title = "예산(실시간)"
    머리 = [None, "구분", "항목", None, "세부항목", "단가", "명수", "횟수", "예산금액", "비율", "결산금액", "차액"]
    for c, v in enumerate(머리, start=1):
        b.cell(2, c, v)
    예산줄 = [  # 줄, 구분, 항목, 세부, 단가, 명수, 횟수, 예산금액, 가리키는 지출 줄
        (3, "그 외", "준비지원", "모임 식사비", None, None, None, 300000, 2),
        (4, None, None, "간식", 1000, 10, 2, 20000, 5),
        (5, "홍보", "포스터", "인쇄", None, None, None, 0, 6),          # 예산금액 0 도 들인다
        (6, None, None, "인쇄", None, None, None, 50000, 7),            # 같은 글자 둘 — 못 잇는다
        (7, "교통", "버스대여", None, 100, 2, 1, 999, 8),               # 셈이 예산금액과 다름
    ]
    for r, 구분, 항목, 세부, 단가, 명수, 횟수, 금액, 지출줄 in 예산줄:
        b.cell(r, 2, 구분); b.cell(r, 3, 항목); b.cell(r, 5, 세부)
        b.cell(r, 6, 단가); b.cell(r, 7, 명수); b.cell(r, 8, 횟수); b.cell(r, 9, 금액)
        b.cell(r, 11, f"='지출 상세내역'!H{지출줄}")
    b.merge_cells("B3:B4"); b.merge_cells("C3:D4"); b.merge_cells("B5:B6"); b.merge_cells("C5:D6")
    b.cell(9, 2, "총 예상 지출 (가)"); b.cell(9, 9, "=SUM(I3:I7)")
    수입 = [(10, "교개협 지원", 5000, 10, 1, 50000), (11, "=B10", None, None, None, 50000),
          (12, "후원금", None, None, None, 70000), (13, "기관부담금 사용", None, None, None, 10000)]
    for r, 이름, 단가, 명수, 횟수, 금액 in 수입:
        b.cell(r, 2, 이름); b.cell(r, 6, 단가); b.cell(r, 7, 명수); b.cell(r, 8, 횟수); b.cell(r, 9, 금액)
    b.cell(14, 2, "총 예상 수입 (나)")
    b.cell(16, 3, "납부내역")   # 안 들인다

    e = wb.create_sheet("지출 상세내역")
    머리 = ["구분", "항목", None, "세부항목-1", "세부항목-2", "세부항목-3", "예산 사용액", "합계금액", None,
          "영수증번호", "지출일자", "금액", "지원금액", "비고", "지급여부", "지급일", "지출자", "지출자 계좌"]
    for c, v in enumerate(머리, start=1):
        e.cell(1, c, v)
    날 = dt.date(2026, 8, 1)
    지출 = [  # 줄, 구분, 항목, 세부1, 세부3, 번호, 금액, 지원(값 또는 식), 비고, 지급, 지출자, 계좌, 초록
        (2, "그 외", "준비지원", "모임 식사비", "헤브론 모임", 1, 130600, f"={상한}*12", "12명 가명들", True, 지어낸이름, 지어낸계좌, False),
        (3, None, None, "모임 식사비", "코람데오 모임", 1, 50000, 45000, "5명 가명들", False, 지어낸이름, 지어낸계좌, False),  # 손 셈이 앱 셈과 다름
        (4, None, None, "모임 식사비", "총무팀 모임", "2, 3, 3", 30000, 30000, "3명과 4명", False, "수련회계좌", None, False),
        (5, None, None, "간식", None, None, 20000, 20000, None, True, "수련회계좌", None, True),
        (6, "홍보", "포스터", "인쇄", None, 4, 10000, 10000, "비고 글", False, 지어낸이름, "빈칸없는번호", False),
        (7, None, None, "없는세부", None, "결산 영수증 파일에 따로 첨부", 5000, 5000, None, False, 지어낸이름, None, False),
        (8, "교통", "버스대여", None, None, None, 7000, 7000, None, False, 지어낸이름, "지어낸은행 999-00가-1111", False),  # 꼴에 걸림
    ]
    for r, 구분, 항목, 세부1, 세부3, 번호, 금액, 지원, 비고, 지급, 지출자, 계좌, 초록 in 지출:
        for c, v in ((1, 구분), (2, 항목), (4, 세부1), (6, 세부3), (10, 번호), (11, 날), (12, 금액), (13, 지원),
                     (14, 비고), (15, 지급), (16, 날 if 지급 else None), (17, 지출자), (18, 계좌)):
            e.cell(r, c, v)
        e.cell(r, 7, f"=M{r}")
        if 초록:
            e.cell(r, 10).fill = PatternFill("solid", fgColor="FFB6D7A8")
    e.merge_cells("A2:A5"); e.merge_cells("B2:C5"); e.merge_cells("A6:A7"); e.merge_cells("B6:C7")
    e.cell(9, 1, "합계 줄")      # 금액 없는 줄은 안 들인다

    rc = wb.create_sheet("영수증")
    for i, n in enumerate((1, 2, 3, 4), start=1):
        rc.cell(i, 1, f"({n}) 가명 영수증")
    if 둘째판:
        rc.cell(10, 1, "(99) 없는 번호")
    wb.save(경로)


@pytest.fixture
def 판(admin_client, tmp_path):
    with app_session() as db:
        r = models.Retreat(name="들여오기 회차", meal_subsidy_per_person=상한,
                           start_date=dt.date(2026, 8, 21), end_date=dt.date(2026, 8, 23))
        db.add(r); db.flush()
        for i, (key, name) in enumerate([("hebron", "5 헤브론"), ("koram", "6 코람데오"),
                                         ("chongmu", "1 총무팀"), ("chongmuM", "1 총무M")]):
            db.add(models.Department(retreat_id=r.id, key=key, name=name, sort_order=i))
        # seed 흉내 — 들여오기가 취소 표시로 내린다
        c = models.BudgetCategory(retreat_id=r.id, level1="옛", level2="seed", planned_amount=1, sort_order=1)
        db.add(c); db.flush()
        e = models.ExpenseEntry(retreat_id=r.id, budget_category_id=c.id, amount=1, subsidy_amount=1)
        db.add(e); db.flush()
        e.attach_receipt(models.ExpenseReceipt(number=1, memo="seed"))
        db.commit()
        rid, seed_c, seed_e = r.id, c.id, e.id
    경로 = tmp_path / "합성.xlsx"
    _시트(경로)
    return {"retreat": rid, "seed_c": seed_c, "seed_e": seed_e, "경로": 경로, "mod": _스크립트()}


def _돌린다(판, 실행, capsys=None):
    with app_session() as db:
        계획 = 판["mod"].돌린다(db, 판["경로"], 판["retreat"], 실행, 사본=False)
    return 계획


# ── 가) seed 가 취소 표시로 내려가고 새 항목이 계획만큼 들어온다 ──


def test48_a01_seed_를_내리고_예산_수입이_계획만큼(판):
    계획 = _돌린다(판, True)
    with app_session() as db:
        assert db.get(models.BudgetCategory, 판["seed_c"]).canceled_at is not None
        assert db.get(models.ExpenseEntry, 판["seed_e"]).canceled_at is not None, "seed 지출이 안 내려갔다"
        산예산 = db.scalars(select(models.BudgetCategory).where(
            models.BudgetCategory.retreat_id == 판["retreat"], models.BudgetCategory.canceled_at.is_(None))).all()
        assert len(산예산) == 계획.수["예산 항목"] == len(계획.예산) > 0
        assert 계획.수["seed 예산 항목 취소"] == 1 and 계획.수["seed 지출 취소"] == 1
        assert any(c.planned_amount == 0 for c in 산예산), "예산금액 0 인 항목도 들어와야 한다"
        # 셈이 맞는 줄만 단가·명수·횟수를 넣는다(고치기가 같은 식으로 다시 세므로)
        for c in 산예산:
            if c.unit_price is not None:
                assert c.unit_price * c.headcount * c.times == c.planned_amount
        수입 = db.scalars(select(models.IncomeItem).where(models.IncomeItem.retreat_id == 판["retreat"])).all()
        assert len(수입) == 계획.수["수입"] and "납부내역" not in {i.name for i in 수입}
        assert "=B10" not in {i.name for i in 수입}, "식으로 옮겨 적은 사본 줄이 들어왔다"
        assert sum(1 for i in 수입 if i.unit_price is not None) == 계획.수["수입 · 단가·명수로"]


# ── 나) 지출 · 식대 인원 ──


def test48_b01_지출과_식대_인원(판):
    계획 = _돌린다(판, True)
    with app_session() as db:
        산 = {e.amount: e for e in db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == 판["retreat"], models.ExpenseEntry.canceled_at.is_(None)))}
        assert len(산) == 계획.수["지출"] == len(계획.지출)
        식 = 산[130600]
        assert (식.is_meal_expense, 식.meal_headcount, 식.subsidy_amount) == (True, 12, 96000)
        비고 = 산[50000]
        assert (비고.is_meal_expense, 비고.meal_headcount, 비고.subsidy_amount) == (True, 5, 40000)
        못뽑음 = 산[30000]                              # 비고에 「명」 이 둘 — 인원을 비운다
        assert (못뽑음.is_meal_expense, 못뽑음.meal_headcount) == (True, None)
        assert 못뽑음.subsidy_amount == 30000, "인원을 못 뽑은 줄은 시트의 손 셈 그대로"
        assert 계획.수["지출 · 식대 · 인원 식"] == 1 and 계획.수["지출 · 식대 · 인원 비고"] == 1
        assert 계획.수["지출 · 식대 · 인원 못뽑음"] == 1
        assert all(e.note is None for e in 산.values()), "비고는 안 들인다"
        # 지급여부·지급일 그대로
        assert 식.paid and 식.paid_date == dt.date(2026, 8, 1) and not 비고.paid and 비고.paid_date is None
    # 식만 저장된 파일에서도 곱셈 식의 값을 센다 — 시트와 같으면(2줄) 「다름」 에 안 들고,
    # 손 셈이 앱 셈과 다른 줄(3줄)만 세어 짝 표 「식대」 에 줄 번호가 남는다
    assert 계획.수["지출 · 식대 · 앱이 다시 센 지원금액이 시트와 다름"] == 1
    assert [r for r, _, _ in 계획.짝["식대"]] == [3]
    md = (pathlib.Path(config.DATA_DIR) / "재정짝표.real.md").read_text(encoding="utf-8")
    assert "## 식대 — 1건" in md


# ── 다) 계좌 셋 — split_account 를 실제로 부른다 ──


def test48_c01_계좌는_split_account_로(판, monkeypatch):
    불린 = []
    원래 = 판["mod"].split_account
    monkeypatch.setattr(판["mod"], "split_account", lambda raw, name=None: 불린.append(1) or 원래(raw, name))
    계획 = _돌린다(판, True)
    assert len(불린) == 계획.수["지출 · 계좌 셋으로 가름"] > 0
    with app_session() as db:
        e = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 130600)).one()
        assert (e.payer_bank, e.payer_account_number, e.payer_account_holder) == ("지어낸은행", "999-0000-1111", 지어낸이름)
        계좌없음 = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 20000)).one()
        assert (계좌없음.payer_bank, 계좌없음.payer_account_number, 계좌없음.payer_account_holder) == (None, None, None)
    assert 계획.수["지출 · 계좌 끝 빈칸 뗌"] >= 1


def test48_c03_꼴에_걸린_계좌는_빼고_들이고_짝_표에(판):
    """게이트(budget.account_problem · 봐둘것 BB-c)에 걸린 줄은 계좌 없이 들어가고 짝 표에 남는다 — 전체는 안 멈춘다."""
    계획 = _돌린다(판, True)
    걸린줄 = {r for r, _, _ in 계획.짝["계좌"]}
    assert 걸린줄 == {6, 8}, "은행 없는 번호(6)와 글자 섞인 번호(8)가 걸려야 한다"
    assert 계획.수["지출 · 계좌 꼴에 걸려 계좌 없이 들임"] == len(걸린줄)
    with app_session() as db:
        산 = {e.amount: e for e in db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == 판["retreat"], models.ExpenseEntry.canceled_at.is_(None)))}
        assert len(산) == 계획.수["지출"], "게이트에 걸렸다고 지출이 빠지면 안 된다"
        for 금액 in (10000, 7000):
            e = 산[금액]
            assert (e.payer_bank, e.payer_account_number, e.payer_account_holder) == (None, None, None)
        # 게이트를 지난 계좌는 그대로 들어간다
        assert 산[130600].payer_account_number == "999-0000-1111"
        계좌든 = sum(1 for e in 산.values() if e.payer_account_number)
        assert 계좌든 == 계획.수["지출 · 계좌 셋으로 가름"] - len(걸린줄)
    md = (pathlib.Path(config.DATA_DIR) / "재정짝표.real.md").read_text(encoding="utf-8")
    assert f"## 계좌 — {len(걸린줄)}건" in md and "999-00가" not in md, "짝 표에 까닭만 — 값은 안 적는다"


def test48_c02_스크립트가_계좌를_스스로_가르지_않는다():
    tree = ast.parse((ROOT / "scripts" / "재정들여오기.py").read_text(encoding="utf-8"))
    assert not _가르는곳들(tree), "들여오기 스크립트가 계좌 가르기를 다시 적었다"
    assert any(isinstance(n, ast.Call) and getattr(n.func, "id", None) == "split_account" for n in ast.walk(tree))


# ── 라) 영수증 · 잇기 · 번호 둘 ──


def test48_d01_영수증과_잇기(판):
    계획 = _돌린다(판, True)
    with app_session() as db:
        새 = db.scalars(select(models.ExpenseReceipt).join(
            models.ExpenseEntry, models.ExpenseEntry.id == models.ExpenseReceipt.expense_id).where(
            models.ExpenseEntry.retreat_id == 판["retreat"], func.coalesce(models.ExpenseReceipt.memo, "") != "seed")).all()
        assert len(새) == 계획.수["영수증"]
        원본 = {r.original_no: r for r in 새 if r.original_no}
        # 한 번호가 두 지출에 — 영수증 한 장에 잇기 둘, 금액은 지출마다
        assert sorted(e.amount for e in 원본["1"].expenses) == [50000, 130600]
        # 한 칸에 번호 둘 — 영수증 둘이 한 지출에
        e30 = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 30000)).one()
        assert sorted(r.original_no for r in e30.receipts) == ["2", "3"], "한 칸에 같은 번호가 두 번이면 한 번만 잇는다"
        assert 계획.수["지출 · 번호 둘 이상 적힌 줄"] == 1
        # 관계(receipts)는 같은 행을 한 번만 돌려주므로 잇기 줄을 직접 센다 — 겹친 번호가 줄 둘을 세우는지
        assert db.scalar(select(func.count()).select_from(models.ExpenseReceiptLink).where(
            models.ExpenseReceiptLink.expense_id == e30.id)) == 2
        # 초록 칸(번호 없음)과 번호 대신 글 — 원본 번호 없는 메모 영수증
        for 금액 in (20000, 5000):
            e = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 금액)).one()
            assert [r.original_no for r in e.receipts] == [None] and e.receipts[0].memo == 판["mod"].따로첨부
        # 번호도 초록도 없는 줄은 영수증 없음
        assert db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 7000)).one().receipts == []
        # 자동 번호는 seed 다음부터 안 겹친다
        번호 = [r.number for r in db.scalars(select(models.ExpenseReceipt))]
        assert len(번호) == len(set(번호)) and min(r.number for r in 새) > 1
        잇기 = db.scalar(select(func.count()).select_from(models.ExpenseReceiptLink).join(
            models.ExpenseReceipt).where(func.coalesce(models.ExpenseReceipt.memo, "") != "seed"))
        assert 잇기 == 계획.수["잇기"]
        # 번호로 들인 영수증은 메모를 안 단다 — 원본 번호가 이미 말한다 · 초록만 따로 첨부
        assert 원본["2"].memo is None


# ── 마) 짝 표 ──


def test48_e01_짝_표가_남고_짝_표로_고친다(판):
    계획 = _돌린다(판, False)
    md = (pathlib.Path(config.DATA_DIR) / "재정짝표.real.md").read_text(encoding="utf-8")
    for 이름, 목록 in 계획.짝.items():
        assert f"## {이름} — {len(목록)}건" in md
    못이은 = {r for r, _, _ in 계획.짝["예산"]}
    assert {6, 7} <= 못이은, "같은 글자 둘인 줄과 글자 없는 줄이 짝 표에 있어야 한다"
    assert any(사유 == "영수증 번호 없음" for _, 사유, _ in 계획.짝["영수증"])
    # 부서 — 이름 있는 줄은 붙고, 없는 줄은 없음으로 센다
    붙음 = {e["줄"]: e["부서키"] for e in 계획.지출}
    assert (붙음[2], 붙음[3], 붙음[4]) == ("hebron", "koram", "chongmu")
    # 짝 표 json 으로 한 줄을 잇는다 — 다시 돌리면 그 줄이 이어진다
    (pathlib.Path(config.DATA_DIR) / "재정짝표.real.json").write_text(
        json.dumps({"예산": {"7": 6}, "부서": {"8": "koram"}}), encoding="utf-8")
    try:
        계획2 = _돌린다(판, True)
        assert 계획2.수["지출 · 예산 항목 짝 표로 이음"] == 1 and 계획2.수["지출 · 부서 짝 표로"] == 1
        with app_session() as db:
            e = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 5000)).one()
            assert e.budget_category is not None and e.budget_category.level3 == "인쇄"
            assert db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 7000)).one().department.key == "koram"
    finally:
        (pathlib.Path(config.DATA_DIR) / "재정짝표.real.json").unlink()


def test48_e02_이은_줄은_예산_항목_글자_못_이은_줄은_시트_글자(판):
    _돌린다(판, True)
    with app_session() as db:
        이은 = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 130600)).one()
        c = 이은.budget_category
        assert (이은.level1, 이은.level2, 이은.level3a) == (c.level1, c.level2, c.level3)
        못 = db.scalars(select(models.ExpenseEntry).where(models.ExpenseEntry.amount == 5000)).one()
        assert 못.budget_category_id is None and 못.level3a == "없는세부"


# ── 바) 두 번 들이지 않는다 · 아) 미리보기 수 = 실행 뒤 수 ──


def test48_f01_두_번째는_이미_들여왔습니다(판):
    _돌린다(판, True)
    with app_session() as db:
        전 = 판["mod"].센다(db, db.get(models.Retreat, 판["retreat"]))
    with pytest.raises(판["mod"].멈춤, match="이미 들여왔습니다"):
        _돌린다(판, False)
    with pytest.raises(판["mod"].멈춤, match="이미 들여왔습니다"):
        _돌린다(판, True)
    with app_session() as db:
        assert 판["mod"].센다(db, db.get(models.Retreat, 판["retreat"])) == 전


def test48_g01_미리보기_수가_실행_뒤_수와_같다(판, capsys):
    with app_session() as db:
        r = db.get(models.Retreat, 판["retreat"])
        전 = 판["mod"].센다(db, r)
    미리 = _돌린다(판, False)
    with app_session() as db:
        assert 판["mod"].센다(db, db.get(models.Retreat, 판["retreat"])) == 전, "미리보기가 바꿨다"
    실행 = _돌린다(판, True)
    assert 미리.수 == 실행.수
    with app_session() as db:
        후 = 판["mod"].센다(db, db.get(models.Retreat, 판["retreat"]))
    assert 후["산 예산 항목"] == 미리.수["예산 항목"]
    assert 후["산 지출"] == 미리.수["지출"] and 후["산 수입"] == 미리.수["수입"]
    assert 후["영수증"] - 전["영수증"] == 미리.수["영수증"]
    assert 후["산 잇기"] - 전["산 잇기"] == 미리.수["잇기"]
    # 값을 찍지 않는다 — 지어낸 이름·계좌·금액이 출력에 없다
    out = capsys.readouterr().out
    for 값 in (지어낸이름, "999-0000-1111", "130600", "130,600", "가명들"):
        assert 값 not in out, 값


def test48_g02_사본을_못_뜨면_아무것도_안_바꾼다(판):
    """세션이 여는 DB 와 사본 뜰 파일(DATA_DIR/app.db)이 다르면 멈춘다 — 시험 DB 가 바로 그 경우다."""
    with app_session() as db:
        전 = 판["mod"].센다(db, db.get(models.Retreat, 판["retreat"]))
        with pytest.raises(판["mod"].멈춤, match="다릅니다"):
            판["mod"].돌린다(db, 판["경로"], 판["retreat"], True, 사본=True)
    with app_session() as db:
        assert 판["mod"].센다(db, db.get(models.Retreat, 판["retreat"])) == 전


def test48_g03_회차를_안_주고_열린_회차가_둘이면_멈춘다(판):
    with app_session() as db:
        db.add(models.Retreat(name="다른 열린 회차", meal_subsidy_per_person=상한, start_date=dt.date(2027, 8, 20)))
        db.commit()
    with app_session() as db:
        with pytest.raises(판["mod"].멈춤, match="어디에 넣을지"):
            판["mod"].돌린다(db, 판["경로"], None, False, 사본=False)


def test48_g04_넣다가_어긋나면_아무것도_안_바꾼다(판, monkeypatch):
    """막는 쪽 시험 (11-3) — 넣는 중에 터지거나, 넣은 뒤 DB 를 다시 센 수가 계획과 다르면 되돌린다."""
    def 센_수():
        with app_session() as db:
            r = db.get(models.Retreat, 판["retreat"])
            return 판["mod"].센다(db, r), db.scalar(select(func.count()).select_from(models.ActivityLog))

    전 = 센_수()
    원래 = 판["mod"].넣는다

    def 터진다(db, retreat, 계획):
        원래(db, retreat, 계획)
        raise RuntimeError("넣는 중에 터짐")

    monkeypatch.setattr(판["mod"], "넣는다", 터진다)
    with pytest.raises(RuntimeError):
        _돌린다(판, True)
    assert 센_수() == 전

    def 하나_더(db, retreat, 계획):
        원래(db, retreat, 계획)
        db.add(models.IncomeItem(retreat_id=retreat.id, name="계획에 없는 줄", amount=1, sort_order=99))
        db.flush()

    monkeypatch.setattr(판["mod"], "넣는다", 하나_더)
    with pytest.raises(판["mod"].멈춤, match="되돌렸습니다"):
        _돌린다(판, True)
    assert 센_수() == 전
