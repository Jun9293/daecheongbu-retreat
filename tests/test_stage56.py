"""총무님 확인용 표 — `.md` 정본과 엑셀이 같은 말을 하는가 (2026-09-16 · `scripts/업무대조확인표.py`).

값은 전부 지어낸 것이다. **고장을 심어 잰다** — 엑셀에서 읽는 칸을 한 칸 밀면 빨개져야 한다.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from scripts import 업무대조확인표 as 확인표


def 만든표() -> 확인표.표:
    return 확인표.표(
        짝=[[1, "B", "구분", 12, "명찰 제작 (Sub · 1 총무팀 · 2026-06-01)", 34, "명찰 제작 (main · chongmu · 2026-06-01)", ""],
           [2, "C", "-", 41, "버스 대절 (Main · 7 재정 · 2026-07-01)", 77, "버스 섭외 (main · jaejeong · 2026-07-01)", ""]],
        노션만=[[1, 5, "포스터 인쇄", "Main · 4 스케치 · 2026-05-01", "", ""],
              [2, 9, "간식 구매", "Sch · (담당 없음) · (날짜 없음)", "", ""]],
        자동확정=[[1, 3, "식대 정산 (Main · 7 재정)", 21, "식대 정산 (main · jaejeong)"]])


def test56_a01_md_에_장이_셋이고_A_는_표시_칸이_없다():
    글 = 확인표.md로(만든표())
    assert "## 1. 이 둘이 같은 업무인가요?" in 글 and "## 2. 노션에만 있는 업무" in 글
    assert "## 3. 자동 확정 — 참고용, 표시하지 않으셔도 됩니다" in 글
    자동장 = 글.split("## 3.")[1]
    assert "맞다/아니다/모르겠다" not in 자동장, "자동 확정 장에 표시 칸이 있으면 총무님이 또 본다"
    # 표시 칸은 짝 · 노션만 표에만
    assert 글.count("맞다/아니다/모르겠다") == 1      # 짝 표 머리 한 줄뿐


def test56_a02_md_읽기는_노션_줄을_열쇠로_표시만_꺼낸다():
    그표 = 만든표()
    그표.짝[0][7] = "맞다"
    그표.노션만[1][4] = "안 넣음"
    그표.노션만[1][5] = "간식은 준비물에 묶임"
    읽은것 = 확인표.md읽기(확인표.md로(그표))
    assert 읽은것["짝"] == {"12": ("맞다",), "41": ("",)}
    assert 읽은것["노션만"] == {"5": ("", ""), "9": ("안 넣음", "간식은 준비물에 묶임")}


def test56_b01_엑셀에_고르기가_붙고_열쇠_칸은_숨는다(tmp_path):
    경로 = tmp_path / "확인용.xlsx"
    확인표.엑셀로(경로, 만든표())
    wb = load_workbook(경로)
    assert wb.sheetnames == ["1. 짝 확인", "2. 노션에만", "3. 자동 확정(참고)"]
    짝장 = wb["1. 짝 확인"]
    고르기 = list(짝장.data_validations.dataValidation)
    assert 고르기 and all(말 in 고르기[0].formula1 for 말 in 확인표.짝표시)
    assert "H2:H3" in str(고르기[0].sqref)
    노션만장 = wb["2. 노션에만"]
    assert all(말 in list(노션만장.data_validations.dataValidation)[0].formula1 for 말 in 확인표.노션만표시)
    assert not list(wb["3. 자동 확정(참고)"].data_validations.dataValidation), "참고용 장에는 고를 것이 없다"
    # 열쇠 칸은 숨긴다 — 총무님이 건드릴 자리가 아니다
    머리 = [c.value for c in 짝장[1]]
    from openpyxl.utils import get_column_letter
    for 이름 in ("노션 줄", "run id"):
        글자 = get_column_letter(머리.index(이름) + 1)
        assert 짝장.column_dimensions[글자].hidden
    assert not 짝장.column_dimensions[get_column_letter(머리.index("앱 쪽") + 1)].hidden


def test56_b02_엑셀에_적은_표시를_읽으면_md_와_같은_답이다(tmp_path):
    """수용 기준 라) — 다음 판이 엑셀을 읽어도 정본과 같은 판정이어야 한다."""
    그표, 경로 = 만든표(), tmp_path / "확인용.xlsx"
    확인표.엑셀로(경로, 그표)
    wb = load_workbook(경로)
    wb["1. 짝 확인"]["H2"] = "맞다"
    wb["2. 노션에만"]["E3"] = "안 넣음"
    wb["2. 노션에만"]["F3"] = "간식은 준비물에 묶임"
    wb.save(경로)

    그표.짝[0][7] = "맞다"                  # .md 쪽에 같은 표시를 적으면
    그표.노션만[1][4] = "안 넣음"
    그표.노션만[1][5] = "간식은 준비물에 묶임"
    md글 = 확인표.md로(그표)
    assert 확인표.엑셀읽기(경로) == 확인표.md읽기(md글)
    assert 확인표.견준다(md글, 경로) == []


def test56_b03_한쪽만_적혀_있으면_견준다가_말한다(tmp_path):
    그표, 경로 = 만든표(), tmp_path / "확인용.xlsx"
    확인표.엑셀로(경로, 그표)
    wb = load_workbook(경로)
    wb["1. 짝 확인"]["H2"] = "아니다"
    wb.save(경로)
    다름 = 확인표.견준다(확인표.md로(그표), 경로)          # .md 는 빈 채
    assert 다름 and "짝" in 다름[0]


def test56_b04_읽는_칸이_한_칸_밀리면_빨개진다(tmp_path, monkeypatch):
    """심는 고장 = 엑셀에서 표시 칸을 이름이 아니라 자리로 집기."""
    그표, 경로 = 만든표(), tmp_path / "확인용.xlsx"
    확인표.엑셀로(경로, 그표)
    wb = load_workbook(경로)
    wb["1. 짝 확인"]["H2"] = "맞다"
    wb.save(경로)
    그표.짝[0][7] = "맞다"
    md글 = 확인표.md로(그표)
    assert 확인표.견준다(md글, 경로) == []

    진짜 = 확인표.엑셀읽기

    def 한칸밀림(p):
        답 = 진짜(p)
        답["짝"] = {열쇠: ("",) for 열쇠 in 답["짝"]}      # 표시 칸을 엉뚱한 데서 읽은 꼴
        return 답
    monkeypatch.setattr(확인표, "엑셀읽기", 한칸밀림)
    assert 확인표.견준다(md글, 경로), "고장을 심었는데 견준다가 아무 말도 안 한다"


def test56_c01_표에_세로줄이_든_제목이_들어와도_칸이_안_쪼개진다():
    그표 = 만든표()
    그표.노션만[0][2] = "포스터 | 인쇄\t(2차)"
    줄 = [l for l in 확인표.md로(그표).splitlines() if l.startswith("| 1 | 5 |")][0]
    assert 줄.count("|") == len(확인표.노션만머리) + 1
    assert 확인표.md읽기(확인표.md로(그표))["노션만"]["5"] == ("", "")


def test56_c02_자동_확정_장은_표시로_안_읽힌다():
    그표 = 만든표()
    읽은것 = 확인표.md읽기(확인표.md로(그표))
    assert set(읽은것["짝"]) == {"12", "41"} and set(읽은것["노션만"]) == {"5", "9"}


@pytest.mark.parametrize("빠진장", ["1. 짝 확인", "2. 노션에만"])
def test56_c03_장이_없으면_그_갈래는_비어_돌아온다(tmp_path, 빠진장):
    경로 = tmp_path / "확인용.xlsx"
    확인표.엑셀로(경로, 만든표())
    wb = load_workbook(경로)
    del wb[빠진장]
    wb.save(경로)
    답 = 확인표.엑셀읽기(경로)
    갈래 = "짝" if 빠진장.startswith("1") else "노션만"
    assert 답[갈래] == {} and 답["짝" if 갈래 == "노션만" else "노션만"]
