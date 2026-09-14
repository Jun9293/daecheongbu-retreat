"""재정 — 지우는 길을 막고 세는 자리를 모은다 (2026-09-14).

- 예산 항목 · 수입은 지우지 않고 취소 표시를 토글한다 (0장 · 7-3)
- 걸린 지출이 있는 예산 항목도 취소된다(사람이 정한 나-ㄴ). 그 지출은 합계에서
  사라지지 않고 `canceled_category_spent` 로 든다
- 재정 숫자를 세는 곳은 `domain/budget.py` 하나다 — 밖에서 다시 세는지는 코드에서
  끌어낸 이름으로 잰다(손으로 적은 목록과 견주지 않는다)

숫자는 박지 않는다(11-3) — 값끼리, 그리고 **옛 식**(69dcaa5 의 라우터·엑셀에 있던
식을 그대로 옮긴 것)과 새 함수가 같은 자료에서 같은지 본다.
"""

from __future__ import annotations

import ast
import dataclasses
import datetime as dt
import inspect as pyinspect
import io
import pathlib
import re

import pytest
from openpyxl import load_workbook
from sqlalchemy import Integer, func, inspect, select

from app import models
from app.domain import budget as B
from tests.conftest import app_session, login_as, make_user

TODAY = dt.date.today()
ROOT = pathlib.Path(__file__).resolve().parent.parent


# ════════════════════════════════════════════════════════════════════
# 재정 세계 — 필터마다 걸리는 행이 적어도 하나씩 있게
# ════════════════════════════════════════════════════════════════════


@pytest.fixture
def 돈(admin_client):
    with app_session() as db:
        retreat = models.Retreat(
            name="재정 시험 회차", meal_subsidy_per_person=8000,
            start_date=TODAY + dt.timedelta(days=30), end_date=TODAY + dt.timedelta(days=32))
        db.add(retreat)
        db.flush()
        dept = models.Department(retreat_id=retreat.id, key="chongmuM", name="1 총무M",
                                 color_tag="#2F4858", sort_order=0)
        db.add(dept)
        db.flush()

        def cat(l1, l2, planned, order, **kw):
            row = models.BudgetCategory(retreat_id=retreat.id, level1=l1, level2=l2,
                                        planned_amount=planned, sort_order=order, **kw)
            db.add(row)
            db.flush()
            return row

        c_meal = cat("그 외", "준비지원", 1_000_000, 1)
        c_snack = cat("식비", "야식", 300_000, 2, unit_price=3000, headcount=50, times=2)
        c_poster = cat("홍보", "포스터", 100_000, 3)
        c_only_canceled = cat("홍보", "현수막", 50_000, 4)

        def spend(category, amount, *, payer, paid=False, receipt=False, meal=None,
                  canceled=False):
            e = models.ExpenseEntry(
                retreat_id=retreat.id,
                budget_category_id=category.id if category else None,
                expense_date=TODAY, amount=amount, payer_name=payer, paid=paid,
                department_id=dept.id, subsidy_amount=amount)
            if meal:
                e.is_meal_expense = True
                e.meal_headcount = meal
                e.subsidy_amount = min(amount, meal * 8000)
                e.personal_burden_amount = amount - e.subsidy_amount
            if receipt:
                e.attach_receipt(models.ExpenseReceipt(number=spend.n, memo="별첨"))
                spend.n += 1
            if canceled:
                e.canceled_at = dt.datetime.now()
            db.add(e)
            db.flush()
            return e

        spend.n = 1
        spend(c_meal, 130_600, payer="가명하나", meal=12, receipt=True)       # 식대 · 환급 · 미지급
        spend(c_meal, 68_900, payer=B.RETREAT_ACCOUNT, meal=9, paid=True)    # 식대 · 영수증 없음
        spend(c_snack, 120_000, payer="가명둘", paid=True, receipt=True)     # 환급 대상이지만 지급됨
        spend(c_poster, 180_000, payer="가명셋")                             # 환급 · 미지급 · 영수증 없음
        spend(None, 70_000, payer=B.RETREAT_ACCOUNT, receipt=True)            # 예산 항목 미지정
        spend(c_snack, 40_000, payer="가명넷", canceled=True)                 # 취소된 지출
        spend(c_only_canceled, 30_000, payer="가명다섯", canceled=True)       # 취소된 지출만 걸린 항목
        db.add(models.IncomeItem(retreat_id=retreat.id, name="수련회비", unit_price=50_000,
                                 headcount=100, amount=5_000_000, sort_order=1))
        db.add(models.IncomeItem(retreat_id=retreat.id, name="후원금", amount=700_000,
                                 sort_order=2))
        db.commit()
        ids = {"retreat": retreat.id, "dept": dept.id,
               "cats": {"식대": c_meal.id, "야식": c_snack.id, "포스터": c_poster.id,
                        "현수막": c_only_canceled.id}}
    admin_client.get(f"/board?retreat_id={ids['retreat']}")
    return ids


def _summary(rid):
    with app_session() as db:
        return B.build_budget_summary(db, retreat=db.get(models.Retreat, rid))


def _live_settlement_sum(rid) -> int:
    """취소 안 된 지출 전부의 집행 합 — summary 를 거치지 않고 행에서 바로 셈."""
    with app_session() as db:
        rows = db.scalars(select(models.ExpenseEntry).where(
            models.ExpenseEntry.retreat_id == rid,
            models.ExpenseEntry.canceled_at.is_(None)))
        return sum(e.settlement_amount for e in rows)


# ── 가) 예산 항목 · 수입 — 취소하면 행이 남고 합계에서 빠지고, 되살리면 돌아온다 ──


def test45_a01_예산_항목_취소_되살림(admin_client, 돈):
    rid, cid = 돈["retreat"], 돈["cats"]["포스터"]
    before = _summary(rid)
    planned = next(r.planned for r in before.categories if r.category.id == cid)

    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    with app_session() as db:
        assert db.get(models.BudgetCategory, cid) is not None, "행이 지워졌다"
    mid = _summary(rid)
    assert cid not in {r.category.id for r in mid.categories}
    assert cid in {r.category.id for r in mid.canceled_categories}
    assert mid.total_planned == before.total_planned - planned
    assert all(cid not in {r.category.id for r in g.rows} for g in mid.groups)
    assert cid not in {r.category.id for r in mid.unspent_rows}
    page = admin_client.get(f"/budget?retreat_id={rid}").text
    assert "취소됨" in page and "되살리기" in page

    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    after = _summary(rid)
    assert after.total_planned == before.total_planned
    assert after.total_spent == before.total_spent
    assert [r.category.id for r in after.categories] == [r.category.id for r in before.categories]


def test45_a02_수입_취소_되살림(admin_client, 돈):
    rid = 돈["retreat"]
    before = _summary(rid)
    first = before.active_incomes[0]

    admin_client.post(f"/budget/incomes/{first.id}/cancel?retreat_id={rid}", follow_redirects=True)
    mid = _summary(rid)
    assert first.id in {i.id for i in mid.incomes}, "행이 지워졌다"
    assert first.id not in {i.id for i in mid.active_incomes}
    assert mid.total_income == before.total_income - first.amount
    assert mid.balance == before.balance - first.amount

    admin_client.post(f"/budget/incomes/{first.id}/cancel?retreat_id={rid}", follow_redirects=True)
    assert _summary(rid).total_income == before.total_income

    # 전부 취소하면 「수입 미입력」 으로 돌아간다 — 취소된 행이 있어도 계산하지 않는다
    for i in before.active_incomes:
        admin_client.post(f"/budget/incomes/{i.id}/cancel?retreat_id={rid}", follow_redirects=True)
    page = admin_client.get(f"/budget?retreat_id={rid}").text
    assert "수입 미입력" in page and "취소됨" in page


def test45_a03_막히는_쪽(client, admin_client, 돈):
    """총무팀이 아니면 못 바꾸고, 다른 회차의 id 는 404 이며, 지우던 주소는 없다."""
    rid, cid = 돈["retreat"], 돈["cats"]["야식"]
    make_user("재정 팀원", "01045450001", "dept_lead", dept=돈["dept"])
    login_as(client, "01045450001")
    client.get(f"/board?retreat_id={rid}")
    resp = client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=False)
    assert resp.status_code == 403
    with app_session() as db:
        assert db.get(models.BudgetCategory, cid).canceled_at is None

    with app_session() as db:
        other = models.Retreat(name="다른 회차", meal_subsidy_per_person=8000)
        db.add(other)
        db.commit()
        other_id = other.id
    resp = admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={other_id}")
    assert resp.status_code == 404

    for path in (f"/budget/categories/{cid}/delete", f"/budget/incomes/1/delete"):
        assert admin_client.post(f"{path}?retreat_id={rid}").status_code in (404, 405)
    with app_session() as db:
        assert db.get(models.BudgetCategory, cid) is not None


# ── 나) 걸린 지출이 있어도 취소되고, 그 지출은 합계에서 사라지지 않는다 ──────────


def test45_b01_걸린_지출은_취소된_항목_줄로_세어진다(admin_client, 돈):
    rid = 돈["retreat"]
    before = _summary(rid)
    assert before.total_spent == _live_settlement_sum(rid)        # 불변식이 원래 참인가

    for key in ("야식", "현수막"):     # 산 지출이 걸린 것 · 취소된 지출만 걸린 것
        cid = 돈["cats"][key]
        spent = next(r.spent for r in before.categories if r.category.id == cid)
        resp = admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}",
                                 follow_redirects=True)
        assert resp.status_code == 200 and "옮겨주세요" not in resp.text
        after = _summary(rid)
        assert cid in {r.category.id for r in after.canceled_categories}
        assert after.total_spent == _live_settlement_sum(rid), "걸린 지출이 합계에서 사라졌다"
        assert after.total_spent == before.total_spent
        before = after
        with app_session() as db:
            linked = db.scalars(select(models.ExpenseEntry).where(
                models.ExpenseEntry.budget_category_id == cid)).all()
            assert linked, "걸린 지출이 떨어져 나갔다"

    s = _summary(rid)
    assert s.canceled_category_spent == sum(r.spent for r in s.canceled_categories)
    assert s.canceled_category_spent > 0          # 산 지출이 걸린 야식이 들어 있다

    # 화면 — 예산에는 따로 줄이, 지출 그룹 머리에는 「취소된 예산 항목」 이
    budget_page = admin_client.get(f"/budget?retreat_id={rid}").text
    assert "(취소된 예산 항목에 걸린 지출)" in budget_page
    assert f"{s.canceled_category_spent:,}" in budget_page
    assert "취소된 예산 항목" in admin_client.get(f"/expenses?retreat_id={rid}").text


def test45_b01b_걸린_지출이_있는_항목을_취소하면_집행률이_오른다(admin_client, 돈):
    """지금의 모양을 잰다 — 예산(분모)은 빠지고 집행(분자)은 남는다(7-3 에 적음).
    이 모양을 둔다(2026-09-14 에 사람이 정함 · 봐둘것 BA-f). 바뀌면 이 시험을 고친다."""
    rid, cid = 돈["retreat"], 돈["cats"]["야식"]
    before = _summary(rid)
    row = next(r for r in before.categories if r.category.id == cid)
    assert row.planned > 0 and row.spent > 0
    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    after = _summary(rid)
    assert after.total_planned == before.total_planned - row.planned
    assert after.total_spent == before.total_spent
    assert after.progress_pct > before.progress_pct
    assert after.total_remaining == before.total_remaining - row.planned


def test45_b02_취소된_항목에는_새_지출을_못_붙인다(admin_client, 돈):
    rid, cid = 돈["retreat"], 돈["cats"]["포스터"]
    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    with app_session() as db:
        n = db.scalar(select(func.count()).select_from(models.ExpenseEntry))
    resp = admin_client.post(f"/expenses/create?retreat_id={rid}", data={
        "budget_category_id": str(cid), "amount": "1000", "department_id": str(돈["dept"])})
    assert resp.status_code == 400
    with app_session() as db:
        assert db.scalar(select(func.count()).select_from(models.ExpenseEntry)) == n
    # 지출 등록의 선택지에도 없다
    page = admin_client.get(f"/expenses?retreat_id={rid}").text
    assert f'<option value="{cid}">' not in page


# ── 다) 엑셀과 복제에서 빠진다 ─────────────────────────────────────────────


def test45_c01_엑셀에서_빠진다(admin_client, 돈):
    rid = 돈["retreat"]
    cid = 돈["cats"]["야식"]
    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    s = _summary(rid)
    income = s.active_incomes[0]
    admin_client.post(f"/budget/incomes/{income.id}/cancel?retreat_id={rid}", follow_redirects=True)
    s = _summary(rid)

    resp = admin_client.get(f"/export/expenses.xlsx?retreat_id={rid}")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content))["예산 대비 집행"]
    labels = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    firsts = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    canceled_name = next(r.category.level2 for r in s.canceled_categories)
    assert canceled_name not in labels, "취소된 예산 항목이 파일에 실렸다"
    r = firsts["(취소된 예산 항목에 걸린 지출)"]
    assert ws.cell(row=r, column=5).value == s.canceled_category_spent
    assert income.name not in firsts, "취소된 수입이 파일에 실렸다"
    assert ws.cell(row=firsts["총 수입"], column=4).value == s.total_income
    total = firsts["총계"]
    assert ws.cell(row=total, column=4).value == s.total_planned
    assert ws.cell(row=total, column=5).value == s.total_spent


def test45_c02_복제에서_빠진다(돈):
    from app.domain.clone import clone_retreat

    rid, cid = 돈["retreat"], 돈["cats"]["포스터"]
    with app_session() as db:
        db.get(models.BudgetCategory, cid).canceled_at = dt.datetime.now()
        db.commit()
        source = db.get(models.Retreat, rid)
        names = {(c.level1, c.level2) for c in source.budget_categories if c.canceled_at is None}
        canceled = (db.get(models.BudgetCategory, cid).level1, db.get(models.BudgetCategory, cid).level2)
        new = clone_retreat(db, source=source, name="복제", start_date=None, end_date=None)
        db.commit()
        copied = {(c.level1, c.level2) for c in new.budget_categories}
    assert canceled not in copied
    assert copied == names


# ── 라) 기록에 그때의 금액이 남는다 ────────────────────────────────────────


def test45_d01_기록에_그때의_값이_남는다(admin_client, 돈):
    rid, cid = 돈["retreat"], 돈["cats"]["야식"]
    with app_session() as db:
        cat = db.get(models.BudgetCategory, cid)
        want = {"planned_amount": cat.planned_amount, "unit_price": cat.unit_price,
                "headcount": cat.headcount, "times": cat.times}
        income = db.scalars(select(models.IncomeItem).where(
            models.IncomeItem.retreat_id == rid)).first()
        income_want = {"amount": income.amount, "unit_price": income.unit_price,
                       "headcount": income.headcount}
        income_id = income.id

    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    admin_client.post(f"/budget/categories/{cid}/cancel?retreat_id={rid}", follow_redirects=True)
    admin_client.post(f"/budget/incomes/{income_id}/cancel?retreat_id={rid}", follow_redirects=True)

    with app_session() as db:
        logs = {(l.action, l.target_id): l for l in db.scalars(select(models.ActivityLog))}
    for action in ("예산항목_취소", "예산항목_되살림"):
        log = logs[(action, cid)]
        assert {k: log.before_value[k] for k in want} == want
        assert f"{want['planned_amount']:,}원" in log.summary
    log = logs[("수입_취소", income_id)]
    assert {k: log.before_value[k] for k in income_want} == income_want


# ── 마) 모으기 전후의 수가 같다 — 옛 식과 새 함수를 같은 자료로 ─────────────
#
# 아래 「옛_」 함수는 69dcaa5 의 routers/expenses.py · domain/budget_xlsx.py ·
# routers/budget.py · routers/settings.py 에 있던 식을 **그대로** 옮긴 것이다.


def 옛_필터(entries, name):
    if name == "meal":
        return [e for e in entries if e.is_meal_expense and e.canceled_at is None]
    if name == "unpaid":
        return [e for e in entries if not e.paid and e.canceled_at is None]
    if name == "refund":
        return [e for e in entries
                if B.is_refund_target(e) and not e.paid and e.canceled_at is None]
    if name == "noreceipt":
        return [e for e in entries if not e.receipts and e.canceled_at is None]
    return entries


def 옛_합(entries):
    live = [e for e in entries if e.canceled_at is None]
    return {
        "amount": sum(e.amount for e in live),
        "subsidy": sum(e.subsidy_amount for e in live if e.is_meal_expense),
        "burden": sum(e.personal_burden_amount for e in live if e.is_meal_expense),
        "settlement": sum(e.settlement_amount for e in live),
    }


def 옛_환급시트(entries):
    rows = [e for e in entries if B.is_refund_target(e)]
    unpaid = [e for e in rows if not e.paid]
    return [e.id for e in rows], len(unpaid), sum(e.settlement_amount for e in unpaid)


def 옛_수입식(unit, head, amount):
    return unit * head if (unit is not None and head is not None) else amount


def 옛_지출통계(db, retreat_id):
    live = (models.ExpenseEntry.retreat_id == retreat_id,
            models.ExpenseEntry.canceled_at.is_(None))
    count = db.scalar(select(func.count()).select_from(models.ExpenseEntry).where(*live)) or 0
    total = db.scalar(
        select(func.coalesce(func.sum(models.ExpenseEntry.amount), 0)).where(*live)) or 0
    return int(count), int(total)


def test45_e01_모으기_전후가_같다(admin_client, 돈):
    rid = 돈["retreat"]
    with app_session() as db:
        retreat = db.get(models.Retreat, rid)
        entries = B.entries_of(db, retreat)
        for name in B.FILTERS:
            new = B.filter_entries(entries, name)
            assert [e.id for e in new] == [e.id for e in 옛_필터(entries, name)], name
            assert new, f"{name} 에 걸린 행이 없어 아무것도 안 견줬다"
            t = B.list_totals(new)
            assert vars(t) == 옛_합(new), name
        assert [e.id for e in B.refund_entries(db, retreat)] == [e.id for e in 옛_필터(entries, "refund")]
        assert [e.id for e in B.no_receipt_entries(db, retreat)] == [e.id for e in 옛_필터(entries, "noreceipt")]

        live = [e for e in entries if e.canceled_at is None]
        # 옛 파일은 라우터가 먼저 거른 목록을 받았다 — 새 함수는 **안 거른 목록**을 받아도 같아야 한다
        assert any(e.canceled_at is not None and B.is_refund_target(e) for e in entries), \
            "취소된 환급 대상이 없어 거름을 안 쟀다"
        sheet = B.refund_sheet(entries)
        assert ([e.id for e in sheet.rows], sheet.unpaid_count, sheet.unpaid_total) == 옛_환급시트(live)
        assert sheet.unpaid_count == len(B.filter_entries(entries, "refund"))
        assert B.expense_stats(db, rid) == 옛_지출통계(db, rid)
    for unit, head, amount in ((50_000, 100, 1), (None, 100, 777), (3000, None, 5), (None, None, 0)):
        assert B.income_amount_of(unit, head, amount) == 옛_수입식(unit, head, amount)

    # 화면이 내는 수도 옛 식과 같다 — 필터마다 지표 넷
    for name in B.FILTERS:
        with app_session() as db:
            old = 옛_합(옛_필터(B.entries_of(db, db.get(models.Retreat, rid)), name))
        page = admin_client.get(f"/expenses?retreat_id={rid}&filter={name}").text
        for key in ("settlement", "amount", "subsidy", "burden"):
            assert f"{old[key]:,}원" in page, (name, key)


# ── 바) 라우터 · 템플릿 · 엑셀 쪽에 세는 코드가 없다 — 이름을 코드에서 끌어낸다 ──


def _재정이름() -> set[str]:
    """재정 모델의 정수 칸 · 금액 속성 이름 — 모델에서 끌어낸다(손으로 안 적는다)."""
    names: set[str] = set()
    for model in (models.ExpenseEntry, models.BudgetCategory, models.IncomeItem):
        for col in inspect(model).columns:
            if isinstance(col.type, Integer) and not col.foreign_keys and not col.primary_key \
                    and col.key not in ("sort_order",) and not col.key.startswith("_"):
                names.add(col.key)
        names |= {n for n, _ in pyinspect.getmembers(model, lambda o: isinstance(o, property))
                  if n.endswith("_amount")}
    # domain 이 내는 묶음의 칸과 수 속성 — 라우터·템플릿이 그 값끼리 다시 셈하는 것도 잡는다
    # (커밋 전 검토 [C]). 목록 칸(categories 등)은 수가 아니라 뺀다
    for cls in (B.CategorySummary, B.GroupSummary, B.BudgetSummary, B.ListTotals, B.RefundSheet):
        for f in dataclasses.fields(cls):
            if f.type in ("int", "float", int, float):
                names.add(f.name)
        names |= {n for n, o in pyinspect.getmembers(cls, lambda o: isinstance(o, property))
                  if "int" in str(pyinspect.signature(o.fget).return_annotation)
                  or "float" in str(pyinspect.signature(o.fget).return_annotation)}
    return names


def _파이썬에서_센다(src: str, names: set[str]) -> list[str]:
    tree = ast.parse(src)
    hits = []

    def 재정속성(node):
        return any(isinstance(n, ast.Attribute) and n.attr in names for n in ast.walk(node))

    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            f = node.func
            fname = f.id if isinstance(f, ast.Name) else (f.attr if isinstance(f, ast.Attribute) else "")
            if fname == "sum" and any(재정속성(a) for a in node.args):
                hits.append(f"{node.lineno}: sum")
        # 누적 더하기(`total += e.amount`) — BinOp 가 아니라 따로 본다(검토 [C])
        if isinstance(node, ast.AugAssign) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            if 재정속성(node.value):
                hits.append(f"{node.lineno}: 누적")
        if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            # 목록 · 글자 잇기(`[..] + [..]` · f"..")는 셈이 아니다 — 엑셀의 줄 만들기가 그 모양이다
            잇기 = (ast.List, ast.Tuple, ast.ListComp, ast.JoinedStr)
            if isinstance(node.left, 잇기) or isinstance(node.right, 잇기):
                continue
            if isinstance(node.left, ast.BinOp) and isinstance(node.left.left, 잇기):
                continue
            if 재정속성(node.left) or 재정속성(node.right):
                hits.append(f"{node.lineno}: 셈")
    return hits


_식 = re.compile(r"\{\{(.*?)\}\}|\{%(.*?)%\}", re.S)


def _템플릿에서_센다(text: str, names: set[str]) -> list[str]:
    hits = []
    for m in _식.finditer(text):
        inner = (m.group(1) or m.group(2) or "")
        inner = re.sub(r"'[^']*'|\"[^\"]*\"", "''", inner)        # 문자열 안은 안 본다
        inner = inner.strip().strip("-").strip()                   # Jinja 공백 다듬기 표시 `{%-` · `-%}`
        inner = re.sub(r"^-\s*(?=[A-Za-z_])", "", inner)          # 앞의 부호 하나(`-row.remaining`)
        if "|sum" in inner.replace(" ", ""):
            hits.append(inner.strip())
            continue
        uses = any(re.search(rf"\.{n}\b", inner) for n in names)
        # 빼기는 붙여 써도 잡는다(`a.amount-b.amount` — 검토 [C]). 문자열은 위에서 지웠다
        if uses and re.search(r"[-+*/]", inner):
            hits.append(inner.strip())
    return hits


def test45_f01_밖에서_세는_코드가_없다():
    names = _재정이름()
    assert {"amount", "planned_amount", "settlement_amount"} <= names, "이름을 못 끌어냈다"

    파이썬 = sorted((ROOT / "app" / "routers").glob("*.py")) + [ROOT / "app/domain/budget_xlsx.py"]
    걸림 = {p.name: h for p in 파이썬
            if (h := _파이썬에서_센다(p.read_text(encoding="utf-8"), names))}
    assert not 걸림, f"domain/budget.py 밖에서 재정 숫자를 센다: {걸림}"

    화면 = sorted((ROOT / "app" / "templates").rglob("*.html"))
    assert len(화면) > 10, "템플릿을 못 찾았다"
    걸림 = {p.name: h for p in 화면
            if (h := _템플릿에서_센다(p.read_text(encoding="utf-8"), names))}
    assert not 걸림, f"화면에서 재정 숫자를 센다: {걸림}"


def test45_f02_검사가_실제로_잡는다():
    """③ 검사가 볼 것을 보는가 — 옛 코드의 모양을 심으면 걸려야 한다."""
    names = _재정이름()
    assert _파이썬에서_센다("x = sum(e.amount for e in live)", names)
    assert _파이썬에서_센다("v = row.unit_price * row.headcount", names)
    assert not _파이썬에서_센다("ws.append([e.amount] + ([e.payer_account] if x else []))", names)
    assert _파이썬에서_센다("total += e.amount", names)
    assert _파이썬에서_센다("v = row.planned - row.spent", names)
    assert _템플릿에서_센다("{{ e.amount-e.subsidy_amount }}", names)
    assert _템플릿에서_센다("{{ row.planned - row.spent }}", names)
    assert _파이썬에서_센다("t = func.sum(E.amount)", names)   # 옛 settings 의 SQL 합도 잡힌다
    assert _템플릿에서_센다("{{ (e.subsidy_amount if e.is_meal_expense else e.amount) + 1 }}", names)
    assert _템플릿에서_센다("{{ rows|sum(attribute='amount') }}", names)
    assert not _템플릿에서_센다("{{ e.settlement_amount|num }}", names)
    assert not _템플릿에서_센다("{% if row.remaining < 0 %}", names)
    assert not _템플릿에서_센다("{%- if row.spent -%}", names)       # 다시 보기 [C'] ①
    assert not _템플릿에서_센다("{{ -row.remaining }}", names)
