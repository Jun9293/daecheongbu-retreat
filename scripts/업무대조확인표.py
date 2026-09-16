# -*- coding: utf-8 -*-
"""업무 대조 **확인용 표** — 총무님이 표시할 자리를 만들고 되읽습니다 (2026-09-16 · 사람이 정함 ④).

정본은 `.md` 입니다(다음 판이 그것을 읽습니다). 다만 메모장으로 열면 표가 줄글로 보여
총무님이 표시하기 어렵습니다 — 그래서 **같은 내용을 엑셀로도 함께** 냅니다. 엑셀 쪽은
표시 칸이 **고르기(드롭다운)** 이고, 열쇠 칸(노션 줄 · run id)은 **숨깁니다** — 총무님이
건드릴 자리가 아닙니다.

- **표시가 적히는 자리는 엑셀입니다**(총무님이 여는 것이 그 파일입니다) — `.md` 는 **표의 정본**(어떤 줄이 어떤 열쇠로 서는가)이고,
  표시는 엑셀이 유일한 출처입니다. `견준다` 는 줄과 열쇠가 같은지를 보는 자리이고, 표시가 엑셀에만 있으면 그 차이를 그대로 말합니다
- **열쇠는 번호가 아니라 노션 줄 번호(와 run id)** 입니다 — 규칙을 손봐 번호가 밀려도 되읽습니다
- **등급 A 는 표시할 자리가 아닙니다**(사람이 정함 ①) — 참고용 장으로 따로 냅니다
- 두 파일이 같은 것을 말하는지는 `견준다` 가 봅니다 — 다음 판이 엑셀을 읽어도 `.md` 와 같은 답이어야 합니다

**이 파일은 값을 스스로 읽지 않습니다** — 부르는 쪽(`data/` 의 실행기)이 준 줄만 다룹니다.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

짝표시 = ("맞다", "아니다", "모르겠다")
노션만표시 = ("넣음", "안 넣음", "버림")

짝머리 = ["번호", "등급", "어긋난 칸", "노션 줄", "노션 쪽", "run id", "앱 쪽", "맞다/아니다/모르겠다"]
노션만머리 = ["번호", "노션 줄", "노션 제목", "구분 · 담당 · 시작일", "넣음/안 넣음/버림", "(안 넣음이면) 어느 업무에 묶였나"]
A머리 = ["번호", "노션 줄", "노션 쪽", "run id", "앱 쪽"]

숨길칸 = {"짝": ("노션 줄", "run id"), "노션만": ("노션 줄",), "자동확정": ("노션 줄", "run id")}
# **고르는 칸은 이름으로 짚는다** — 「맨 끝 칸」 으로 짚으면 노션만 표의 끝이 메모 칸이라
# 총무님이 결정을 메모 칸에 넣게 된다(커밋 전 검토 [H] 1).
고를칸 = {"짝": 짝머리[7], "노션만": 노션만머리[4]}


@dataclass
class 표:
    짝: list[list] = field(default_factory=list)        # 짝머리 순서
    노션만: list[list] = field(default_factory=list)     # 노션만머리 순서
    자동확정: list[list] = field(default_factory=list)   # A머리 순서 — 표시 칸이 없다


def _칸(값) -> str:
    """`|` · 줄바꿈 · 탭을 표 칸에 그대로 넣지 않는다(CLAUDE.md 14장 `칸글` 과 같은 까닭)."""
    return str("" if 값 is None else 값).replace("|", "¦").replace("\t", " ").replace("\r", "").replace("\n", " ⏎ ")


def md로(그표: 표) -> str:
    """정본 `.md` 를 만든다 — 다음 판이 읽는 것은 이 글이다."""
    줄 = ["# 총무님 확인용 — 업무 대조 (2026 여름)", "",
         "**표시만 해 주세요.** 칸에 글자를 적으시면 됩니다. 이 파일은 저장소에 올라가지 않습니다.",
         "**같은 내용을 엑셀로도 냈습니다** — 엑셀에서는 칸을 누르면 고르기만 하면 됩니다(`data/업무대조-확인용.real.xlsx`).",
         "**번호 말고 「노션 줄」 과 「run id」 가 열쇠입니다** — 규칙을 손봐 순서가 바뀌어도 그 둘은 그대로입니다.", "",
         "## 1. 이 둘이 같은 업무인가요? (등급 B · C)", "",
         f"「{' / '.join(짝표시)}」 중 하나를 적어 주세요. 비워 두시면 다음 판이 「모르겠다」 로 봅니다.",
         "「어긋난 칸」 은 제목 말고 구분 · 담당 · 시작일 중 양쪽 값이 다른 칸입니다.", "",
         "| " + " | ".join(짝머리) + " |", "|" + "---|" * len(짝머리)]
    줄 += ["| " + " | ".join(_칸(c) for c in r) + " |" for r in 그표.짝]
    줄 += ["", "## 2. 노션에만 있는 업무 — 앱에 넣을까요?", "",
           f"「{' / '.join(노션만표시)}」 중 하나를 적어 주세요 — 안 넣음은 「앱의 다른 업무에 이미 묶여 있다」 는 뜻입니다.", "",
           "| " + " | ".join(노션만머리) + " |", "|" + "---|" * len(노션만머리)]
    줄 += ["| " + " | ".join(_칸(c) for c in r) + " |" for r in 그표.노션만]
    줄 += ["", "## 3. 자동 확정 — 참고용, 표시하지 않으셔도 됩니다", "",
           "제목이 사실상 같고 구분 · 담당 · 시작일도 안 어긋나는 짝입니다(등급 A). 다음 판이 **확인 없이** 잇습니다.", "",
           "| " + " | ".join(A머리) + " |", "|" + "---|" * len(A머리)]
    줄 += ["| " + " | ".join(_칸(c) for c in r) + " |" for r in 그표.자동확정]
    return "\n".join(줄) + "\n"


def _표줄(글: str) -> list[list[str]]:
    난것 = []
    for 한줄 in 글.splitlines():
        칸 = [c.strip() for c in 한줄.strip().strip("|").split("|")]
        if len(칸) >= 4 and 칸[0].isdigit():
            난것.append(칸)
    return 난것


def md읽기(글: str) -> dict[str, dict[str, tuple[str, ...]]]:
    """`.md` 에서 **표시만** 꺼낸다 — 열쇠는 노션 줄 번호다.

    칸 수로 어느 표인지 가른다: 짝 표는 `짝머리` 만큼, 노션만 표는 `노션만머리` 만큼.
    자동 확정 장(표시 칸이 없다)은 읽지 않는다.
    """
    나온것: dict[str, dict[str, tuple[str, ...]]] = {"짝": {}, "노션만": {}}
    for 칸 in _표줄(글):
        if len(칸) == len(짝머리):
            나온것["짝"][칸[3]] = (칸[7],)
        elif len(칸) == len(노션만머리):
            나온것["노션만"][칸[1]] = (칸[4], 칸[5])
    return 나온것


def 엑셀로(경로, 그표: 표) -> None:
    """같은 내용을 엑셀로 — 표시 칸은 고르기, 열쇠 칸은 숨김."""
    wb = Workbook()
    wb.remove(wb.active)
    장들 = (("1. 짝 확인", 짝머리, 그표.짝, "짝", DataValidation(type="list", formula1='"' + ",".join(짝표시) + '"', allow_blank=True)),
          ("2. 노션에만", 노션만머리, 그표.노션만, "노션만", DataValidation(type="list", formula1='"' + ",".join(노션만표시) + '"', allow_blank=True)),
          ("3. 자동 확정(참고)", A머리, 그표.자동확정, "자동확정", None))
    for 이름, 머리, 줄들, 갈래, 고르기 in 장들:
        ws = wb.create_sheet(이름)
        ws.append(머리)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="EFEFED")
        for r in 줄들:
            ws.append([_칸(v) for v in r])
        for i, 칸이름 in enumerate(머리, start=1):
            글자 = get_column_letter(i)
            ws.column_dimensions[글자].width = 14 if 칸이름 in ("번호", "노션 줄", "run id", "등급", "어긋난 칸") else 46
            if 칸이름 in 숨길칸[갈래]:
                ws.column_dimensions[글자].hidden = True        # 총무님이 건드릴 칸이 아니다
        if 고르기 is not None and 줄들:
            ws.add_data_validation(고르기)
            고름 = get_column_letter(머리.index(고를칸[갈래]) + 1)
            고르기.add(f"{고름}2:{고름}{len(줄들) + 1}")
            for 행 in range(2, len(줄들) + 2):
                ws[f"{고름}{행}"].fill = PatternFill("solid", fgColor="FBF3E4")   # 여기서 고르시면 됩니다
        ws.freeze_panes = "A2"
        for 행 in ws.iter_rows(min_row=2):
            for c in 행:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(경로)


def 엑셀읽기(경로) -> dict[str, dict[str, tuple[str, ...]]]:
    """엑셀에서 표시를 꺼낸다 — `md읽기` 와 **같은 꼴**이어야 한다(`견준다` 가 그것을 본다)."""
    wb = load_workbook(경로, data_only=True)
    나온것: dict[str, dict[str, tuple[str, ...]]] = {"짝": {}, "노션만": {}}
    for 이름, 갈래 in (("1. 짝 확인", "짝"), ("2. 노션에만", "노션만")):
        if 이름 not in wb.sheetnames:
            continue
        ws = wb[이름]
        머리 = [str(c.value or "").strip() for c in ws[1]]
        열쇠칸 = 머리.index("노션 줄")
        표시칸 = [i for i, h in enumerate(머리) if h in (짝머리[7], 노션만머리[4], 노션만머리[5])]
        for 행 in ws.iter_rows(min_row=2, values_only=True):
            if 행[0] is None or not str(행[0]).strip():
                continue
            나온것[갈래][str(행[열쇠칸]).strip()] = tuple(str(행[i] or "").strip() for i in 표시칸)
    return 나온것


def 표시있나(읽은것: dict[str, dict[str, tuple[str, ...]]]) -> int:
    """표시가 하나라도 적힌 줄이 몇인가 — **덮기 전에 이것을 본다**(커밋 전 검토 [H] 2).

    총무님이 적는 자리는 엑셀이라, `.md` 만 보던 가드는 엑셀에만 적은 표시를 못 보고 덮었다.
    """
    return sum(1 for 갈래 in ("짝", "노션만") for 표시 in 읽은것.get(갈래, {}).values() if any(표시))


def 견준다(md글: str, 엑셀경로) -> list[str]:
    """`.md` 와 엑셀이 같은 표시를 말하는가 — 다른 자리를 글로 돌려준다(값은 안 담는다)."""
    ㅁ, ㅇ = md읽기(md글), 엑셀읽기(엑셀경로)
    다름 = []
    for 갈래 in ("짝", "노션만"):
        빠진것 = set(ㅁ[갈래]) ^ set(ㅇ[갈래])
        if 빠진것:
            다름.append(f"{갈래}: 한쪽에만 있는 줄 {len(빠진것)}")
        다른표시 = [열쇠 for 열쇠 in ㅁ[갈래] if 열쇠 in ㅇ[갈래] and ㅁ[갈래][열쇠] != ㅇ[갈래][열쇠]]
        if 다른표시:
            다름.append(f"{갈래}: 표시가 다른 줄 {len(다른표시)}")
    return 다름
