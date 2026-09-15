# -*- coding: utf-8 -*-
"""「Belong 예산」 시트를 앱의 재정으로 들여옵니다 (CLAUDE.md 7-5 · 2026-09-14 재정 차례 6).

## 무엇을 어떻게 — 정본은 봐둘것 BA-f · BB-a 입니다

| 시트 | 어디로 | 정한 것 |
|---|---|---|
| 「예산(실시간)」 의 예산 줄 | `BudgetCategory` | 시트-바 ㄷ(이 탭만) · 예산금액 0 인 줄도 그대로 |
| 같은 탭의 수입 묶음 | `IncomeItem` | 시트-사 ㄱ(수입 줄만 · 납부내역은 안 들임) |
| 「지출 상세내역」 의 금액 있는 줄 | `ExpenseEntry` | 취소한 지출은 없다 — 전부 실제로 나간 돈 |
| 지출 → 예산 항목 | `budget_category_id` | 시트 자신의 결산 식으로(BB-i · 글자 비교는 확인용 · 못 이은 줄은 짝 표) |
| 지출자 계좌 | 계좌 셋 | 시트-라 ㄱ — `budget.split_account` 로 가르고 `budget.account_problem` 을 지난 것만 넣는다(규칙을 여기 다시 적지 않는다 · 걸리면 계좌 없이 들이고 짝 표에) |
| 식대 줄 | 인원 · 지원금액 | 시트-다 ㄴ — 인원은 식에서, 없으면 비고의 「숫자 + 명」 에서. 비고는 **숫자만** 읽는다 |
| 영수증번호 | `ExpenseReceipt.original_no` + 잇기 표 | 시트-가 ㄷ · 마-ㄷ — 한 칸에 번호 둘이면 영수증 둘, 같은 번호는 한 장에 여러 지출 |
| 영수증 그림 | 안 들인다 | 시트-아 ㄱ |
| 부서 | `department_id` | 시트-마 ㄱ — 세부항목 칸의 부서 이름을 키로, 못 맞춘 줄은 비움 |
| 지급여부 · 지급일 | `paid` · `paid_date` | 시트 값 그대로(사람이 정함) |
| 구매목록 · 예산 판 탭 여럿 | 안 들인다 | 시트-자 ㄱ · 시트-바 ㄷ |
| seed 로 선 이 회차의 재정 | 취소 표시 | 가-ㄷ — 들여올 때 내린다. 지우지 않는다(0장) |

**비고는 안 들입니다.** 식대 줄의 비고는 참석자 명단이고, 시트-다 ㄴ 이 명단을 DB 에 옮기는 갈래(ㄷ)를
고르지 않았습니다. 다른 줄의 비고에도 무엇이 적혔는지 모릅니다 — 미리보기가 몇 줄인지 셉니다.

## 값을 찍지 않습니다

미리보기와 실행은 **건수만** 찍습니다 — 이름·계좌·금액·항목 이름을 안 찍습니다(이 출력은 보고와 채팅으로
옮겨집니다). **값이 있는 짝 표는 `data/재정짝표.real.md` 에** 씁니다 — `data/*.real.*` 라 저장소로 안 갑니다.
사람이 그 표를 보고 짝을 적으려면 `data/재정짝표.real.json` 에 `{"예산": {"지출 시트 줄": 예산 시트 줄}, "부서":
{"지출 시트 줄": "부서 키"}}` 를 적고 다시 돌립니다. 줄 번호는 시트의 줄 번호입니다.

## 기본은 미리보기

`--실행` 일 때만 바꿉니다. 바꾸기 직전에 `VACUUM INTO` 사본을 뜨고, 세션이 여는 파일과 사본을 뜰 파일이
다르면 아무것도 안 합니다(영수증잇기옮기기.py 와 같은 꼴). **한 회차에 두 번 들이지 않습니다** — 들여온
기록(활동 기록 `재정_들여오기`)이 있으면 미리보기부터 「이미 들여왔습니다」 로 멈춥니다.

## 쓰는 법

    .venv\\Scripts\\python.exe scripts/재정들여오기.py data/Belong예산.real.xlsx --회차 1
    .venv\\Scripts\\python.exe scripts/재정들여오기.py data/Belong예산.real.xlsx --회차 1 --실행

`--회차` 를 안 주면 열려 있는 회차가 하나일 때만 그것을 씁니다(비품들여오기의 `회차머리` 를 부릅니다).
"""

from __future__ import annotations

import sys as _sys

for _stream in (_sys.stdout, _sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

import argparse
import datetime as dt
import json
import pathlib
import re
import sys
import warnings
from collections import Counter
from dataclasses import dataclass, field

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from sqlalchemy import func, select                                  # noqa: E402
from sqlalchemy.engine import make_url                               # noqa: E402
from sqlalchemy.orm import Session                                   # noqa: E402

from app import config                                               # noqa: E402
from app.db import SessionLocal                                      # noqa: E402
from app.domain.budget import account_problem, next_receipt_number, split_account  # noqa: E402
from app.domain.meal import calculate_meal_settlement                # noqa: E402
from app.models import (                                             # noqa: E402
    ActivityLog, BudgetCategory, Department, ExpenseEntry, ExpenseReceipt, ExpenseReceiptLink, IncomeItem,
    Retreat,
)
from scripts import backup                                           # noqa: E402
from scripts.비품들여오기 import 멈춤, 회차머리                           # noqa: E402

예산탭, 지출탭, 영수증탭 = "예산(실시간)", "지출 상세내역", "영수증"
# 예산 줄과 수입 줄을 가르는 표지 줄 — 칸 이름 수준의 낱말이다. 표지 칸에 괄호 설명 같은 것이 더 붙어
# 있어 **낱말이 다 들어 있는가**로 찾는다(글자 그대로 견주면 못 찾는다 · 2026-09-14 실측)
지출끝, 수입끝 = "총 예상 지출", "총 예상 수입"
들여오기행위 = "재정_들여오기"
# 「결산 영수증 파일에 따로 첨부」 라는 뜻의 초록 채움(사람이 답함 · 봐둘것 BB-a)
따로첨부 = "결산 영수증 파일에 따로 첨부"


# ── 시트를 읽는다 (DB 를 안 본다) ─────────────────────────────────────


def _글(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _정수(v) -> int | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    return None


def _날짜(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _병합값(ws) -> dict[tuple[int, int], object]:
    """병합칸의 빈 자리 → 병합 첫 칸의 값. 비품 들여오기의 「합친 칸은 위의 값을 물려받는다」 와 같은 규칙."""
    out = {}
    for m in ws.merged_cells.ranges:
        v = ws.cell(m.min_row, m.min_col).value
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                out[(r, c)] = v
    return out


def _칸(ws, 병합, r, c):
    v = ws.cell(r, c).value
    return 병합.get((r, c)) if v is None else v


def _머리칸(ws, 머리줄: int, 이름들: list[str]) -> dict[str, int]:
    """머리글을 **이름으로** 찾는다 — 칸이 하나 끼면 자리로 찾은 것은 조용히 어긋난다(4-18)."""
    있는 = {}
    for c in range(1, ws.max_column + 1):
        v = _글(ws.cell(머리줄, c).value)
        if v and v not in 있는:
            있는[v] = c
    없는 = [n for n in 이름들 if n not in 있는]
    if 없는:
        raise 멈춤(f"「{ws.title}」 {머리줄}행에 머리글이 없습니다: {', '.join(없는)}")
    return {n: 있는[n] for n in 이름들}


@dataclass
class 시트:
    예산: list[dict] = field(default_factory=list)
    수입: list[dict] = field(default_factory=list)
    지출: list[dict] = field(default_factory=list)
    이름표: set[int] = field(default_factory=set)
    결산식: dict[int, set[int]] = field(default_factory=dict)   # 지출 시트 줄 → 받는 예산 시트 줄들


def 읽는다(경로: pathlib.Path) -> 시트:
    """시트를 읽어 줄마다 dict 로 — **값은 여기 안에만 있고 찍지 않는다.**"""
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wf = load_workbook(경로, data_only=False)   # 식 그대로
        wv = load_workbook(경로, data_only=True)    # 계산된 값
    for 탭 in (예산탭, 지출탭, 영수증탭):
        if 탭 not in wf.sheetnames:
            raise 멈춤(f"탭 「{탭}」 이 없습니다.")
    out = 시트()
    out.결산식 = _결산식(wf)

    # 예산(실시간) — 머리글 줄은 앞의 몇 줄 안에서 이름으로 찾는다
    bf, bv = wf[예산탭], wv[예산탭]
    머리줄 = next((r for r in range(1, 6) if _글(bf.cell(r, 2).value) == "구분"
                 or any(_글(bf.cell(r, c).value) == "예산금액" for c in range(1, 20))), None)
    if 머리줄 is None:
        raise 멈춤(f"「{예산탭}」 의 머리글 줄을 못 찾았습니다.")
    k = _머리칸(bf, 머리줄, ["구분", "항목", "세부항목", "단가", "명수", "횟수", "예산금액", "결산금액"])
    bm = _병합값(bf)
    가운데 = k["항목"] + 1 if k["항목"] + 1 < k["세부항목"] else None
    끝줄 = {}
    for r in range(머리줄 + 1, bf.max_row + 1):
        v = _글(bf.cell(r, k["구분"]).value) or ""
        for 표지 in (지출끝, 수입끝):
            if 표지 not in 끝줄 and all(낱말 in v for 낱말 in 표지.split()):
                끝줄[표지] = r
    if 지출끝 not in 끝줄 or 수입끝 not in 끝줄:
        raise 멈춤(f"「{예산탭}」 에서 「{지출끝}」 · 「{수입끝}」 줄을 못 찾았습니다.")
    for r in range(머리줄 + 1, 끝줄[지출끝]):
        결산 = str(bf.cell(r, k["결산금액"]).value or "")
        # 예산 줄 = 결산금액이 지출 탭의 칸을 가리키는 줄(시트 자신이 예산 줄로 세는 것)
        if not (결산.startswith("=") and "!" in 결산):
            continue
        항목 = _글(_칸(bf, bm, r, k["항목"]))
        중간 = _글(bf.cell(r, 가운데).value) if 가운데 and (r, 가운데) not in bm else None
        out.예산.append({
            "줄": r, "구분": _글(_칸(bf, bm, r, k["구분"])),
            "항목": " · ".join(x for x in (항목, 중간) if x) or None,
            "세부": _글(_칸(bf, bm, r, k["세부항목"])),
            "단가": _정수(bv.cell(r, k["단가"]).value), "명수": _정수(bv.cell(r, k["명수"]).value),
            "횟수": _정수(bv.cell(r, k["횟수"]).value),
            "예산금액": _정수(bv.cell(r, k["예산금액"]).value) or 0,
        })
    for r in range(끝줄[지출끝] + 1, 끝줄[수입끝]):
        이름 = bf.cell(r, k["구분"]).value
        if not isinstance(이름, str) or 이름.startswith("=") or not 이름.strip():
            continue      # 식으로 옮겨 적은 줄은 같은 수입의 사본이다
        out.수입.append({
            "줄": r, "이름": 이름.strip(),
            "단가": _정수(bv.cell(r, k["단가"]).value), "명수": _정수(bv.cell(r, k["명수"]).value),
            "횟수": _정수(bv.cell(r, k["횟수"]).value),
            "금액": _정수(bv.cell(r, k["예산금액"]).value) or 0,
        })

    # 지출 상세내역 — 1행이 머리글
    ef, ev = wf[지출탭], wv[지출탭]
    k = _머리칸(ef, 1, ["구분", "항목", "세부항목-1", "세부항목-2", "세부항목-3", "영수증번호",
                       "지출일자", "금액", "지원금액", "비고", "지급여부", "지급일", "지출자", "지출자 계좌"])
    em = _병합값(ef)
    가운데 = k["항목"] + 1 if k["항목"] + 1 < k["세부항목-1"] else None
    for r in range(2, ef.max_row + 1):
        금액 = _정수(ev.cell(r, k["금액"]).value)
        if 금액 is None:
            continue
        번호칸 = ev.cell(r, k["영수증번호"])
        채움 = 번호칸.fill.fgColor.rgb if 번호칸.fill and 번호칸.fill.fill_type == "solid" else None
        항목 = _글(_칸(ef, em, r, k["항목"]))
        중간 = _글(ef.cell(r, 가운데).value) if 가운데 and (r, 가운데) not in em else None
        out.지출.append({
            "줄": r, "구분": _글(_칸(ef, em, r, k["구분"])),
            "항목": " · ".join(x for x in (항목, 중간) if x) or None,
            "세부1": _글(_칸(ef, em, r, k["세부항목-1"])), "세부2": _글(_칸(ef, em, r, k["세부항목-2"])),
            "세부3": _글(_칸(ef, em, r, k["세부항목-3"])),
            "번호": ev.cell(r, k["영수증번호"]).value, "초록": _초록인가(채움),
            "일자": _날짜(ev.cell(r, k["지출일자"]).value), "금액": 금액,
            "지원": _정수(ev.cell(r, k["지원금액"]).value) if ev.cell(r, k["지원금액"]).value is not None
                    else _곱식(ef.cell(r, k["지원금액"]).value),
            "지원식": _글(ef.cell(r, k["지원금액"]).value),
            "비고": _글(ev.cell(r, k["비고"]).value), "지급": bool(ev.cell(r, k["지급여부"]).value),
            "지급일": _날짜(ev.cell(r, k["지급일"]).value), "지출자": _글(ev.cell(r, k["지출자"]).value),
            "계좌": ev.cell(r, k["지출자 계좌"]).value,
        })

    for row in wv[영수증탭].iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                m = re.match(r"\s*\((\d+)\)", cell.value)
                if m:
                    out.이름표.add(int(m.group(1)))
    return out


_주소 = r"\$?[A-Z]{1,3}\$?(\d+)"


def _식의줄들(식: str, 탭: str | None = None) -> set[int]:
    """**아는 꼴의 식만** 줄 번호로 푼다 — 모르는 꼴이면 빈 집합(→ 짝 표로 간다).

    아는 꼴: `=주소` · `=주소:주소` · `=SUM(범위, …)`(시트에 소문자 sum 도 있다) 과 그것들을 `+` 로 이은 것.
    **탭을 주면 범위 앞에 탭 이름이 있어야 한다** — Excel 이 쓰는 기본 꼴 `'탭'!H2:H4`(앞에 한 번)도, 두 끝에 다 붙인 꼴도
    받는다(사람이 정함 · 2026-09-15 · 봐둘것 BB-i). 뒤 끝에만 붙은 꼴은 모른다. `탭` 을 주면 주소마다 그 탭 이름이
    붙어 있어야 하고(결산 식 → 지출 탭), 안 주면 탭 이름이 없어야 한다(합계 식 → 같은 탭). 범위는 펼치고 거꾸로 쓴
    범위도 받는다. **칸 주소 모양만 모으면** 숫자가 붙은 함수 이름 · 따옴표 안 글자 · 조건 칸(SUMIF)까지 줄로 잡혀
    틀린 한 줄로 조용히 채울 수 있다(2026-09-15 커밋 전 검토) — 그래서 꼴을 먼저 본다.
    """
    머리 = "" if 탭 is None else rf"(?:'{re.escape(탭)}'|{re.escape(탭)})!"
    칸 = r"\$?[A-Z]{1,3}\$?\d+"
    범위 = rf"{머리}{칸}(?:\s*:\s*(?:{머리})?{칸})?"
    항 = rf"(?:(?i:SUM)\(\s*{범위}(?:\s*,\s*{범위})*\s*\)|{범위})"
    본문 = 식.strip()
    if not 본문.startswith("=") or not re.fullmatch(rf"{항}(?:\s*\+\s*{항})*", 본문[1:].strip()):
        return set()
    줄 = set()
    for m in re.finditer(rf"{머리}{_주소}(?:\s*:\s*(?:{머리})?{_주소})?", 본문):
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        줄 |= set(range(min(a, b), max(a, b) + 1))
    return 줄


def 결산식으로(경로: pathlib.Path) -> dict[int, set[int]]:
    """시트 **자신의 연결** — 지출 시트 줄 → 그 줄을 합계로 받는 예산 시트 줄들 (2026-09-15 · 봐둘것 BB-h).

    예산 탭의 결산금액 식은 지출 탭의 합계금액 칸을 주소로 가리키고(7-1), 그 합계금액 칸은 지출 줄 범위를
    더하는 식이다. 그 두 식을 따라가 지출 줄마다 어느 예산 줄이 받는지 구한다. **들여오기(`고른다`)가 이것으로 잇는다** —
    두 탭의 이름이 달라도(한 예산 줄이 여러 품목을 묶은 이름) 시트가 적어 둔 답이다. 합계금액 칸이 식이 아니면
    그 줄 하나다. 값을 찍지 않는다.
    """
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return _결산식(load_workbook(경로, data_only=False))


def _결산식(wf) -> dict[int, set[int]]:
    bf, ef = wf[예산탭], wf[지출탭]
    머리줄 = next((r for r in range(1, 6) if _글(bf.cell(r, 2).value) == "구분"
                 or any(_글(bf.cell(r, c).value) == "예산금액" for c in range(1, 20))), None)
    if 머리줄 is None:
        raise 멈춤(f"「{예산탭}」 의 머리글 줄을 못 찾았습니다.")
    bk = _머리칸(bf, 머리줄, ["구분", "결산금액"])
    ek = _머리칸(ef, 1, ["합계금액"])
    끝 = next((r for r in range(머리줄 + 1, bf.max_row + 1)
              if all(낱말 in (_글(bf.cell(r, bk["구분"]).value) or "") for 낱말 in 지출끝.split())), bf.max_row + 1)
    out: dict[int, set[int]] = {}
    for r in range(머리줄 + 1, 끝):
        결산 = str(bf.cell(r, bk["결산금액"]).value or "")
        if not (결산.startswith("=") and "!" in 결산):     # 읽는다 의 「예산 줄」 과 같은 판정
            continue
        for 합계줄 in _식의줄들(결산, 지출탭):
            합계 = str(ef.cell(합계줄, ek["합계금액"]).value or "")
            for 지출줄 in (_식의줄들(합계) if 합계.startswith("=") else {합계줄}):
                out.setdefault(지출줄, set()).add(r)
    return out


def _곱식(v) -> int | None:
    """계산된 값이 안 딸려 온 파일(식만 저장된 사본)에서 「=숫자*숫자」 식의 값을 센다. 그 밖의 식은 None."""
    if isinstance(v, str) and re.fullmatch(r"=\s*\d+(\s*\*\s*\d+)+\s*", v):
        n = 1
        for x in re.findall(r"\d+", v):
            n *= int(x)
        return n
    return None


def _초록인가(rgb) -> bool:
    """채움이 초록 계열인가 — G 가 R·B 보다 뚜렷이 크다."""
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    try:
        r, g, b = (int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    return g > r + 20 and g > b + 20


def 번호들(v) -> list[int] | None:
    """영수증번호 칸 → 번호 목록. 비면 [] · 번호가 아닌 글이면 None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return []
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return [int(v)]
    if isinstance(v, str) and re.fullmatch(r"\s*\d+(\s*,\s*\d+)*\s*", v):
        # 같은 번호를 한 칸에 두 번 적어도 한 번만 — 잇기 줄이 겹치지 않게
        return list(dict.fromkeys(int(x) for x in re.findall(r"\d+", v)))
    return None


def 인원(줄: dict, 상한: int) -> tuple[bool, int | None, str]:
    """(식대인가, 인원, 어디서). 지원금액이 곱셈 식이면 식에서 상한이 아닌 인수, 아니면 비고의 「숫자 + 명」.
    비고는 **숫자만** 읽는다(시트-다 ㄴ)."""
    식 = 줄["지원식"] or ""
    비고명 = re.findall(r"(\d+)\s*명", 줄["비고"] or "")
    if 식.startswith("=") and "*" in 식:
        인수 = [int(x) for x in re.findall(r"\d+", 식)]
        남은 = [x for x in 인수 if x != 상한]
        if len(남은) == 1:
            return True, 남은[0], "식"
        return True, (int(비고명[0]) if len(비고명) == 1 else None), ("비고" if len(비고명) == 1 else "못뽑음")
    if 비고명:
        return True, (int(비고명[0]) if len(비고명) == 1 else None), ("비고" if len(비고명) == 1 else "못뽑음")
    return False, None, ""


# ── 들일 것을 정한다 (아무것도 안 바꾼다) ────────────────────────────


def _부서키(db: Session, retreat: Retreat) -> dict[str, str]:
    """부서 이름(앞 번호와 빈칸을 뗀 것) → 키. 키 없는 부서는 짝지을 수 없어 뺀다."""
    return {re.sub(r"^\d+\s*", "", d.name).replace(" ", ""): d.key
            for d in db.scalars(select(Department).where(Department.retreat_id == retreat.id)) if d.key}


def 짝표읽기(경로: pathlib.Path) -> dict:
    if not 경로.exists():
        return {"예산": {}, "부서": {}}
    d = json.loads(경로.read_text(encoding="utf-8"))
    return {"예산": {int(a): int(b) for a, b in d.get("예산", {}).items()},
            "부서": {int(a): str(b) for a, b in d.get("부서", {}).items()}}


@dataclass
class 계획:
    수: Counter
    예산: list[dict]
    수입: list[dict]
    지출: list[dict]
    영수증: dict[int, dict]            # 원본 번호 → {"번호": 자동 번호, "초록": bool}
    글영수증: list[tuple[int, int]]     # (번호 대신 글이 적힌 지출 줄, 자동 번호)
    취소할예산: list[int]
    취소할지출: list[int]
    짝: dict[str, list]                  # 짝 표 — 값이 든다(data/ 에만 쓴다)


def 고른다(db: Session, retreat: Retreat, 판: 시트, 짝표: dict) -> 계획:
    """미리보기와 실행이 **같은 계획**을 쓴다 — 보여 준 수와 넣는 수가 갈리지 않게."""
    if db.scalar(select(func.count()).select_from(ActivityLog).where(
            ActivityLog.retreat_id == retreat.id, ActivityLog.action == 들여오기행위)):
        raise 멈춤("이미 들여왔습니다 — 이 회차에는 재정을 들여온 기록이 있습니다. 같은 시트를 두 번 들이지 않습니다.\n"
                  "  다시 들이려면 들여오기 직전 사본으로 되돌린 뒤 돌리세요(docs/배포-안내.md 12장 · CLAUDE.md 11-2).")
    수 = Counter()
    짝 = {"예산": [], "글자와 식": [], "부서": [], "영수증": [], "식대": [], "계좌": []}
    상한 = retreat.meal_subsidy_per_person

    수["seed 예산 항목 취소"] = len(취소할예산 := list(db.scalars(select(BudgetCategory.id).where(
        BudgetCategory.retreat_id == retreat.id, BudgetCategory.canceled_at.is_(None)))))
    수["seed 지출 취소"] = len(취소할지출 := list(db.scalars(select(ExpenseEntry.id).where(
        ExpenseEntry.retreat_id == retreat.id, ExpenseEntry.canceled_at.is_(None)))))

    # 예산 — 키(구분 · 항목 · 세부)가 둘 이상인 줄은 글자로 못 잇는다
    열쇠 = Counter((b["구분"], b["항목"], b["세부"]) for b in 판.예산)
    줄로 = {b["줄"]: b for b in 판.예산}
    for b in 판.예산:
        곱 = (b["단가"] * b["명수"] * b["횟수"]) if None not in (b["단가"], b["명수"], b["횟수"]) else None
        b["셈맞음"] = 곱 == b["예산금액"]
        수["예산 항목"] += 1
        수["예산 항목 · 예산금액 0"] += b["예산금액"] == 0
        수["예산 항목 · 단가×명수×횟수 가 예산금액과 다름(셈 칸은 비워 둠)"] += 곱 is not None and not b["셈맞음"]

    for i in 판.수입:
        곱 = i["단가"] * i["명수"] if None not in (i["단가"], i["명수"]) else None
        # IncomeItem 에는 횟수 칸이 없다 — 단가×명수가 금액과 같을 때만 셈 칸을 넣는다(고치기가 같은 식을 쓴다)
        i["셈맞음"] = 곱 is not None and 곱 == i["금액"]
        수["수입"] += 1
        수["수입 · 단가·명수로"] += i["셈맞음"]
        수["수입 · 금액만"] += not i["셈맞음"]

    부서 = _부서키(db, retreat)
    원본번호: dict[int, dict] = {}
    글영수증 = []
    for e in 판.지출:
        수["지출"] += 1
        # 예산 항목 잇기 — **시트 자신의 결산 식으로** 잇는다(사람이 정함 · 2026-09-15 · 봐둘것 BB-i).
        # 글자 비교는 확인용이다: 둘이 어긋나면 결산 식을 따르고 세어 짝 표에 남긴다. 짝 표 json 이 이긴다
        key = (e["구분"], e["항목"], e["세부1"])
        글줄 = next(b["줄"] for b in 판.예산 if (b["구분"], b["항목"], b["세부"]) == key) if 열쇠.get(key) == 1 else None
        식줄 = 판.결산식.get(e["줄"], set())
        if e["줄"] in 짝표["예산"]:
            e["예산줄"] = 짝표["예산"][e["줄"]]
            수["지출 · 예산 항목 짝 표로 이음"] += 1
        elif len(식줄) == 1:
            e["예산줄"] = next(iter(식줄))
            수["지출 · 예산 항목 결산 식으로 이음"] += 1
            if 글줄 is None:
                수["지출 · 결산 식으로 이음 · 글자로는 못 잇는 줄"] += 1
            elif 글줄 != e["예산줄"]:
                수["지출 · 결산 식으로 이음 · 글자로는 다른 예산 줄"] += 1
                짝["글자와 식"].append((e["줄"], "글자로는 다른 예산 줄(결산 식을 따름)", 글줄))
        else:
            e["예산줄"] = None
            사유 = f"결산 식이 가리키는 예산 줄이 {len(식줄)}개"
            수["지출 · 예산 항목 못 이음(결산 식이 가리키는 예산 줄이 하나가 아님)"] += 1
            수["지출 · 못 이음 · 글자로는 하나"] += 글줄 is not None
            짝["예산"].append((e["줄"], 사유, key))
        if e["예산줄"] is not None and e["예산줄"] not in 줄로:
            raise 멈춤(f"짝 표나 결산 식의 예산 줄 {e['예산줄']} 이 시트의 예산 줄이 아닙니다.")

        # 부서 (시트-마 ㄱ)
        if e["줄"] in 짝표["부서"]:
            e["부서키"] = 짝표["부서"][e["줄"]]
            수["지출 · 부서 짝 표로"] += 1
        else:
            글 = "".join(x for x in (e["세부1"], e["세부2"], e["세부3"]) if x).replace(" ", "")
            걸린 = {키 for 이름, 키 in 부서.items() if 이름 and 이름 in 글}
            # 「총무M」 과 「총무팀」 처럼 한 이름이 다른 이름 안에 들면 긴 쪽만 남긴다
            걸린 = {키 for 키 in 걸린 if not any(키 != 딴 and _이름(부서, 키) in _이름(부서, 딴) for 딴 in 걸린)}
            e["부서키"] = next(iter(걸린)) if len(걸린) == 1 else None
            if len(걸린) == 1:
                수["지출 · 부서 붙음"] += 1
            elif 걸린:
                수["지출 · 부서 비움(이름이 둘 이상)"] += 1
                짝["부서"].append((e["줄"], "부서 이름이 둘 이상", sorted(걸린)))
            else:
                수["지출 · 부서 없음(부서 이름이 안 적힘)"] += 1

        # 계좌 (시트-라 ㄱ) — 규칙은 budget.split_account 하나
        원 = e["계좌"] if isinstance(e["계좌"], str) else None
        if 원 and 원.strip():
            e["은행"], e["계좌번호"], e["예금주"] = split_account(원, e["지출자"])
            수["지출 · 계좌 셋으로 가름"] += 1
            수["지출 · 계좌 끝 빈칸 뗌"] += 원 != 원.rstrip()
            수["지출 · 계좌에 빈칸이 없어 은행 비움"] += e["은행"] is None
            # 꼴 게이트 (봐둘것 BB-c) — 걸린 줄은 멈추지 않고 계좌 없이 들이고 짝 표에 남긴다
            문제 = account_problem(e["은행"], e["계좌번호"], e["예금주"])
            if 문제:
                수["지출 · 계좌 꼴에 걸려 계좌 없이 들임"] += 1
                짝["계좌"].append((e["줄"], 문제, None))
                e["은행"] = e["계좌번호"] = e["예금주"] = None
        else:
            e["은행"] = e["계좌번호"] = e["예금주"] = None

        # 식대 (시트-다 ㄴ)
        식대, 명, 어디 = 인원(e, 상한)
        e["식대"], e["인원"] = 식대, 명
        if 식대:
            수[f"지출 · 식대 · 인원 {어디}"] += 1
            if 명 is not None:
                s = calculate_meal_settlement(amount=e["금액"], headcount=명, per_person_cap=상한)
                e["새지원"], e["새부담"] = s.subsidy_amount, s.personal_burden_amount
                if e["지원"] != s.subsidy_amount:
                    수["지출 · 식대 · 앱이 다시 센 지원금액이 시트와 다름"] += 1
                    짝["식대"].append((e["줄"], "앱이 다시 센 지원금액이 시트와 다름", None))
            else:
                # 인원을 못 뽑은 줄은 시트의 손 셈을 그대로 둔다 — 고칠 때 다시 세지 않는다(7-4)
                e["새지원"] = e["지원"] if e["지원"] is not None else e["금액"]
                e["새부담"] = e["금액"] - e["새지원"]
        else:
            e["새지원"], e["새부담"] = e["금액"], 0
            수["지출 · 식대 아님인데 시트 지원금액이 금액과 다름"] += e["지원"] is not None and e["지원"] != e["금액"]
        수["지출 · 지급됨"] += e["지급"]
        수["지출 · 비고 안 들임"] += e["비고"] is not None

        # 영수증
        ns = 번호들(e["번호"])
        e["번호들"] = ns or []
        if ns is None or (not ns and e["초록"]):
            # 번호 칸이 초록이거나 번호 대신 글이 적힌 줄 — 영수증이 결산 파일에 따로 있다(사람이 답함).
            # 원본 번호 없는 메모 영수증 한 장으로 잇는다(2026-09-14 실측: 두 경우가 겹친다)
            글영수증.append(e["줄"])
            수["영수증 · 따로 첨부(번호 없이 메모로)"] += 1
            if ns is None and not e["초록"]:
                짝["영수증"].append((e["줄"], "번호 대신 글 · 초록 아님", None))
        elif not ns:
            수["지출 · 영수증 번호 없음"] += 1
            짝["영수증"].append((e["줄"], "영수증 번호 없음", None))
        else:
            수["지출 · 번호 둘 이상 적힌 줄"] += len(ns) > 1
            for n in ns:
                원본번호.setdefault(n, {"초록": False, "지출줄": []})
                원본번호[n]["초록"] |= e["초록"]
                원본번호[n]["지출줄"].append(e["줄"])
                if n not in 판.이름표:
                    짝["영수증"].append((e["줄"], "그림 이름표에 없는 번호", n))
    끝 = next_receipt_number(db, retreat) - 1
    for n in sorted(원본번호):
        끝 += 1
        원본번호[n]["번호"] = 끝
    글영수증 = [(줄, 끝 + i) for i, 줄 in enumerate(글영수증, start=1)]
    수["영수증 · 원본 번호로"] = len(원본번호)
    수["영수증 · 한 장에 지출 여럿"] = sum(1 for v in 원본번호.values() if len(set(v["지출줄"])) > 1)
    수["영수증 · 그림 이름표에 없는 번호"] = sum(1 for n in 원본번호 if n not in 판.이름표)
    수["잇기"] = sum(len(v["지출줄"]) for v in 원본번호.values()) + len(글영수증)
    수["영수증"] = len(원본번호) + len(글영수증)
    return 계획(수, 판.예산, 판.수입, 판.지출, 원본번호, 글영수증, 취소할예산, 취소할지출, 짝)


def _이름(부서: dict[str, str], 키: str) -> str:
    return next(이름 for 이름, k in 부서.items() if k == 키)


# ── 넣는다 ───────────────────────────────────────────────────────────


def 넣는다(db: Session, retreat: Retreat, 판: 계획) -> None:
    """한 트랜잭션의 안쪽 — commit 은 부르는 쪽이 DB 를 다시 세어 계획과 맞을 때만 한다."""
    지금 = dt.datetime.now()
    for cid in 판.취소할예산:
        db.get(BudgetCategory, cid).canceled_at = 지금
    for eid in 판.취소할지출:
        db.get(ExpenseEntry, eid).canceled_at = 지금
    끝순서 = db.scalar(select(func.coalesce(func.max(BudgetCategory.sort_order), 0))
                    .where(BudgetCategory.retreat_id == retreat.id)) or 0
    예산 = {}
    for i, b in enumerate(판.예산, start=1):
        c = BudgetCategory(retreat_id=retreat.id, level1=b["구분"] or "", level2=b["항목"] or "",
                           level3=b["세부"], planned_amount=b["예산금액"], sort_order=끝순서 + i)
        if b["셈맞음"]:
            c.unit_price, c.headcount, c.times = b["단가"], b["명수"], b["횟수"]
        db.add(c)
        예산[b["줄"]] = c
    수입순서 = db.scalar(select(func.coalesce(func.max(IncomeItem.sort_order), 0))
                    .where(IncomeItem.retreat_id == retreat.id)) or 0
    for i, s in enumerate(판.수입, start=1):
        db.add(IncomeItem(retreat_id=retreat.id, name=s["이름"], amount=s["금액"], sort_order=수입순서 + i,
                          unit_price=s["단가"] if s["셈맞음"] else None,
                          headcount=s["명수"] if s["셈맞음"] else None))
    db.flush()
    부서 = {d.key: d.id for d in db.scalars(select(Department).where(Department.retreat_id == retreat.id)) if d.key}
    영수증 = {}
    지출 = {}
    for e in 판.지출:
        c = 예산.get(e["예산줄"]) if e["예산줄"] is not None else None
        entry = ExpenseEntry(
            retreat_id=retreat.id, budget_category_id=c.id if c else None,
            # 이은 줄은 예산 항목의 글자를 복사한다(앱의 등록과 같다 · 봐둘것 BB-g) · 못 이은 줄은 시트 글자
            level1=c.level1 if c else e["구분"], level2=c.level2 if c else e["항목"],
            level3a=c.level3 if c else e["세부1"], level3b=e["세부2"], level3c=e["세부3"],
            expense_date=e["일자"], amount=e["금액"], department_id=부서.get(e["부서키"]),
            payer_name=e["지출자"], payer_bank=e["은행"], payer_account_number=e["계좌번호"],
            payer_account_holder=e["예금주"], paid=e["지급"], paid_date=e["지급일"],
            is_meal_expense=e["식대"], meal_headcount=e["인원"],
            subsidy_amount=e["새지원"], personal_burden_amount=e["새부담"],
        )
        db.add(entry)
        지출[e["줄"]] = entry
        for n in e["번호들"]:
            if n not in 영수증:
                v = 판.영수증[n]
                영수증[n] = ExpenseReceipt(number=v["번호"], original_no=str(n),
                                          memo=따로첨부 if v["초록"] else None)
            entry.attach_receipt(영수증[n])
    db.flush()
    for 줄, 번호 in 판.글영수증:
        지출[줄].attach_receipt(ExpenseReceipt(number=번호, memo=따로첨부))
    db.add(ActivityLog(retreat_id=retreat.id, actor_type="system", actor_name="재정들여오기",
                       action=들여오기행위, target_type="retreat", target_id=retreat.id,
                       summary="시트에서 재정을 들여옴 — " + " · ".join(
                           f"{k} {판.수[k]}" for k in ("예산 항목", "수입", "지출", "영수증", "잇기")),
                       after_value={k: v for k, v in 판.수.items()}))
    db.flush()


def 센다(db: Session, retreat: Retreat) -> dict[str, int]:
    """DB 의 지금 수 — 실행 전후를 견준다."""
    산 = lambda m: db.scalar(select(func.count()).select_from(m).where(  # noqa: E731
        m.retreat_id == retreat.id, m.canceled_at.is_(None))) or 0
    return {"산 예산 항목": 산(BudgetCategory), "산 지출": 산(ExpenseEntry), "산 수입": 산(IncomeItem),
            "영수증": db.scalar(select(func.count()).select_from(ExpenseReceipt).join(
                ExpenseEntry, ExpenseEntry.id == ExpenseReceipt.expense_id).where(
                ExpenseEntry.retreat_id == retreat.id)) or 0,
            "산 잇기": db.scalar(select(func.count()).select_from(ExpenseReceiptLink).join(
                ExpenseEntry, ExpenseEntry.id == ExpenseReceiptLink.expense_id).where(
                ExpenseEntry.retreat_id == retreat.id, ExpenseReceiptLink.detached_at.is_(None))) or 0}


def 짝표쓰기(경로: pathlib.Path, 판: 계획) -> None:
    """값이 든 짝 표 — data/*.real.* 라 저장소로 안 간다. 사람이 보고 짝 표 json 을 적는다."""
    줄 = ["# 재정 들여오기 짝 표 (값이 든다 · 저장소에 올리지 않는다)", ""]
    for 이름, 목록 in 판.짝.items():
        줄 += [f"## {이름} — {len(목록)}건", "", "| 지출 시트 줄 | 사유 | 참고 |", "|---|---|---|"]
        줄 += [f"| {r} | {사유} | {참고} |" for r, 사유, 참고 in 목록]
        줄.append("")
    경로.write_text("\n".join(줄), encoding="utf-8")


def 미리보기찍기(수: Counter) -> None:
    for k in sorted(수):
        print(f"  {k}: {수[k]}")


def 사본을_뜬다() -> pathlib.Path:
    원본 = pathlib.Path(config.DATA_DIR) / "app.db"
    여는것 = make_url(config.DATABASE_URL).database
    if not 여는것 or pathlib.Path(여는것).resolve() != 원본.resolve():
        raise 멈춤(f"세션이 여는 DB 와 사본을 뜰 파일이 다릅니다 — 아무것도 안 했습니다.")
    사본 = backup.snapshot(원본, pathlib.Path(config.DATA_DIR) / "backups")
    print(f"  사본을 떴습니다: {사본.name}  ({사본.stat().st_size:,} 바이트)")
    return 사본


def 돌린다(db: Session, 경로: pathlib.Path, 회차id: int | None, 실행: bool, *, 사본: bool = True) -> 계획:
    retreat = 회차머리(db, 회차id)
    판 = 고른다(db, retreat, 읽는다(경로), 짝표읽기(pathlib.Path(config.DATA_DIR) / "재정짝표.real.json"))
    짝표쓰기(pathlib.Path(config.DATA_DIR) / "재정짝표.real.md", 판)
    print("들일 것(건수만):")
    미리보기찍기(판.수)
    print(f"  짝 표: data/재정짝표.real.md (예산 {len(판.짝['예산'])} · 부서 {len(판.짝['부서'])} · 영수증 {len(판.짝['영수증'])} · 식대 {len(판.짝['식대'])} · 계좌 {len(판.짝['계좌'])})")
    if not 실행:
        print("바꾸지 않았습니다. 실제로 들이려면 --실행 을 붙이세요.")
        return 판
    전 = 센다(db, retreat)
    if 사본:
        사본을_뜬다()
    # 넣은 뒤 commit 전에 DB 를 다시 세어 계획과 견준다 — 계획에서 되짚은 수로 견주면 절대 안 어긋난다
    기대 = {"산 예산 항목": 전["산 예산 항목"] - len(판.취소할예산) + 판.수["예산 항목"],
            "산 지출": 전["산 지출"] - len(판.취소할지출) + 판.수["지출"],
            "산 수입": 전["산 수입"] + 판.수["수입"],
            "영수증": 전["영수증"] + 판.수["영수증"],
            "산 잇기": 전["산 잇기"] + 판.수["잇기"]}
    try:
        넣는다(db, retreat, 판)
        잰 = 센다(db, retreat)
        for k, v in 기대.items():
            if 잰[k] != v:
                raise 멈춤(f"{k}: 계획대로면 {v}인데 DB 에서 {잰[k]} — 되돌렸습니다.")
        db.commit()
    except Exception:
        db.rollback()
        raise
    후 = 센다(db, retreat)
    print("들였습니다. 다시 센 수:")
    for k in 후:
        print(f"  {k}: {전[k]} → {후[k]}")
    return 판


def main() -> int:
    ap = argparse.ArgumentParser(description="「Belong 예산」 시트를 재정으로 들여옵니다 (CLAUDE.md 7-5).",
                                 epilog="기본은 미리보기입니다. 실제로 들이려면 --실행 을 붙이세요.")
    ap.add_argument("시트", type=pathlib.Path)
    ap.add_argument("--회차", type=int, default=None)
    ap.add_argument("--실행", action="store_true")
    args = ap.parse_args()
    with SessionLocal() as db:
        try:
            돌린다(db, args.시트, getattr(args, "회차"), getattr(args, "실행"))
            return 0
        except 멈춤 as e:
            print(f"!! {e}")
            return 1


if __name__ == "__main__":
    sys.exit(main())
